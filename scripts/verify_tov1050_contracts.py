"""Reproducible subprocess detector/report comparison and real CSV coverage.

No files are written to the reference TOV640 repository. Canonical synthetic
cases are explicitly separated from real input acceptance.
"""
from __future__ import annotations
import argparse
import hashlib
import gzip
import io
import json
import math
import os
import re
from pathlib import Path
import subprocess
import sys
from itertools import zip_longest

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PIN = "cbca528f090152ec2c776ee04a4c30f0d371810f"
sys.path.insert(0, str(ROOT / "backend"))


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def dump(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def compare(left, right, path="root"):
    """Strict structure/order/strings; numerical absolute tolerance 1e-9, rtol 0."""
    if isinstance(left, dict) and isinstance(right, dict):
        if list(left) != list(right):
            raise AssertionError(f"{path}: keys/order differ")
        for key in left:
            compare(left[key], right[key], f"{path}.{key}")
    elif isinstance(left, list) and isinstance(right, list):
        if len(left) != len(right):
            raise AssertionError(f"{path}: lengths {len(left)} != {len(right)}")
        for index, (a, b) in enumerate(zip(left, right)):
            compare(a, b, f"{path}[{index}]")
    elif isinstance(left, (int, float)) and isinstance(right, (int, float)):
        if not math.isclose(left, right, rel_tol=0, abs_tol=1e-9):
            raise AssertionError(f"{path}: {left} != {right}")
    elif left != right:
        raise AssertionError(f"{path}: {left!r} != {right!r}")


def excel_equal(actual,expected):
    if expected is None or pd.isna(expected) or expected == "":
        return actual is None or pd.isna(actual) or actual == ""
    if isinstance(expected,(int,float)):
        return isinstance(actual,(int,float)) and math.isclose(actual,expected,rel_tol=0,abs_tol=1e-9)
    return str(actual)==str(expected)


def subprocess_pair(reference, spec, output, name):
    source = output / f"{name}-input.json"
    dump(source, spec)
    payloads = []
    for label, repo in (("tov640", reference), ("tov1050", ROOT)):
        target = output / f"{name}-{label}.json.gz"
        subprocess.run([sys.executable, str(ROOT / "scripts/tov1050_contract_worker.py"),
                        "--repo", str(repo), "--input", str(source), "--output", str(target)],
                       check=True, cwd=output, env={**os.environ, "PYTHONDONTWRITEBYTECODE": "1"},
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
        with gzip.open(target, "rt", encoding="utf-8") as stream:
            payloads.append(json.load(stream))
    compare(payloads[0], payloads[1])
    return dict(case=name, status="pass", input_sha256=sha(source),
                input_kind="real_datac" if spec.get("datac") else "synthetic_canonical_contract_only",
                counts={k: len(v["rows"]) for k, v in payloads[0]["results"].items()},
                levels=sorted({r["level"] for v in payloads[0]["results"].values() for r in v["rows"]}),
                raw_rows=payloads[0]["raw_rows"], report_sheets=list(payloads[0]["report"]))


def canonical_cases():
    # Four-wire extrema, equal peaks, exact threshold values, NaN wires,
    # separated groups, ascending/descending and both ID modes.
    heights = [5300]*4 + [4400, 4450, 4500, 4550, 4599, 4600]*2 + [5300]*4 + [6000, 5900, 5850, 5800]*2 + [5300]*4
    n = len(heights)
    frame = pd.DataFrame({"Chainage": [100000 + i * .2 for i in range(n)]})
    for wire in range(1, 5):
        frame[f"height{wire}"] = heights if wire == 1 else [5300]*n
        frame[f"wear{wire}"] = ([8.5, 9.5, 10.2, 11.0, 12.0, 13.0] * (n//6+1))[:n]
        frame[f"stagger{wire}"] = ([0, 550, 550, 450, 350, 300, 0, -550, -450, -350] * (n//10+1))[:n]
    frame.loc[5, "height4"] = None
    frame.loc[8, "wear4"] = None
    for reverse in (False, True):
        context = dict(line="EAL", section="Mainline", track="UP", date_str="20261023")
        if reverse:
            context.update(task_no="U1", station_start="AAA", station_end="BBB")
        selected = frame.iloc[::-1] if reverse else frame
        yield ("canonical-desc-custom" if reverse else "canonical-asc-default"), json.loads(selected.to_json(orient="records")), context
    isolated = pd.DataFrame({"Chainage": [100000 + i * .2 for i in range(36)]})
    stagger = [0]*3 + [550]*3 + [0]*3 + [450]*3 + [0]*3 + [350]*3 + [0]*3 + [-550]*3 + [0]*3 + [-450]*3 + [0]*3 + [-350]*3
    for wire in range(1,5):
        isolated[f"height{wire}"] = [4550]*3 + [5300]*6 + [5850]*3 + [5300]*24
        isolated[f"wear{wire}"] = [13]*9 + [9.5]*3 + [13]*24
        isolated[f"stagger{wire}"] = stagger
    yield "canonical-isolated-levels", json.loads(isolated.to_json(orient="records")), dict(line="EAL",section="Mainline",track="UP",date_str="20261023")


def run_real_csv(output):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from app.api.endpoints import analysis
    from app.core.tov1050_contract import LINE_SESSIONS, parse_input_filename
    from app.core.tov1050_data_ingestion import TOV1050DataLoader
    app = FastAPI()
    app.include_router(analysis.router, prefix="/api")
    client = TestClient(app)
    fixture_roots = [ROOT / "docs/TOV1050_Migration/Raw data", ROOT / "00 Reference Document/Raw data"]
    inventory = []
    seen = set()
    context_path = ROOT / "config/tov1050-regression-fixture-contexts.json"
    contexts = json.loads(context_path.read_text(encoding="utf-8")) if context_path.exists() else {}
    from types import SimpleNamespace
    for path in sorted(p for root in fixture_roots for p in root.rglob("*.csv")):
        digest = sha(path)
        if digest in seen:
            continue
        seen.add(digest)
        try:
            context = parse_input_filename(path)
            inventory.append((path, context))
        except ValueError:
            short = re.fullmatch(r"(\d{8})_(KTL|TKL)_(UT|DT)",path.stem)
            if short:
                date,line,track=short.groups()
                inventory.append((path,SimpleNamespace(line=line,track=track,date_str=date,
                                  station_start="NOT_PROVIDED",station_end="NOT_PROVIDED",session="Mainline")))
            elif str(path.relative_to(ROOT)).replace("\\","/") in contexts:
                inventory.append((path,SimpleNamespace(**contexts[str(path.relative_to(ROOT)).replace("\\","/")])))
            else:
                inventory.append((path, None))
    rows, evidence = [], []
    for line, sessions in LINE_SESSIONS.items():
        for session in sorted(sessions):
            for direction in ("UT", "DT"):
                matches = [(p, c) for p, c in inventory if c and c.line == line and c.track == direction and session == getattr(c,"session","Mainline")]
                row = dict(line=line, session=session, direction=direction, fixture_count=len(matches),
                           status="blocked_missing_real_fixture" if not matches else "pass", fixtures=";".join(str(p) for p, _ in matches))
                rows.append(row)
                for index, (path, context) in enumerate(matches):
                    session_id = f"regression-{line}-{session}-{direction}-{index}"
                    params = dict(file_path=str(path), line=line, track=direction, session=session,
                                  section=session, date_str=context.date_str, station_start=context.station_start,
                                  station_end=context.station_end, client_session_id=session_id)
                    response = client.post("/api/analyze", json=params)
                    if response.status_code != 200:
                        raise AssertionError(f"{path}: {response.status_code} {response.text}")
                    payload = response.json()
                    print(f"Analyzed {line}/{session}/{direction}: {payload['cleaning_summary']['retained_row_count']} rows",flush=True)
                    frame = analysis.ANALYSIS_SESSIONS[session_id]["raw_df"]
                    tiny_chart = {k: v[:1] for k, v in payload["chart_data"].items()}
                    report = client.post("/api/export/report/generate", json=dict(params=payload["params"], exceptions=payload["exceptions"], chart_data=tiny_chart, client_session_id=session_id))
                    raw = client.post("/api/export/raw/generate", json=dict(params=payload["params"], chart_data=tiny_chart, client_session_id=session_id))
                    assert report.status_code == raw.status_code == 200
                    print(f"Exported {line}/{session}/{direction}: report={len(report.content)} raw={len(raw.content)} bytes",flush=True)
                    exported = pd.read_csv(io.BytesIO(raw.content))
                    # Read ChartData sequentially so the verifier does not build another 25M-cell tree.
                    workbook=load_workbook(io.BytesIO(report.content),read_only=True,data_only=True)
                    chart_rows=workbook["ChartData"].iter_rows(values_only=True)
                    chart_columns=list(next(chart_rows))
                    chart_count=0
                    for actual,expected in zip_longest(chart_rows,frame[chart_columns].itertuples(index=False,name=None)):
                        assert actual is not None and expected is not None,"ChartData row count differs"
                        actual=tuple(actual)+(None,)*(len(expected)-len(actual))
                        assert all(excel_equal(a,b) for a,b in zip(actual,expected)),f"{line} ChartData row {chart_count+2} differs"
                        chart_count+=1
                    workbook.close()
                    assert len(exported) == chart_count == len(frame)
                    summary=pd.read_excel(io.BytesIO(report.content),sheet_name="Summary")
                    from app.core.exporter import ExcelExporter
                    expected_summary=pd.DataFrame([record for records in payload["exceptions"].values() for record in records])
                    expected_summary=ExcelExporter._reorder_columns(expected_summary,ExcelExporter.REPORT_COLUMNS).rename(columns=ExcelExporter.REPORT_HEADERS)
                    assert list(summary.columns)==list(expected_summary.columns) and len(summary)==len(expected_summary)
                    for actual,expected in zip(summary.itertuples(index=False,name=None),expected_summary.itertuples(index=False,name=None)):
                        assert all(excel_equal(a,b) for a,b in zip(actual,expected)),f"{line} Summary differs"
                    checked = []
                    for column in frame.columns.intersection(exported.columns):
                        if pd.api.types.is_numeric_dtype(frame[column]):
                            pd.testing.assert_series_equal(exported[column], frame[column], check_dtype=False, check_names=False, atol=1e-9, rtol=0)
                            checked.append(column)
                    prefix = output / f"real-{line}-{session}-{direction}-{index}"
                    dump(prefix.with_name(prefix.name+"-summary.json"),json.loads(summary.to_json(orient="records",double_precision=15)))
                    # Large artifacts are validated in full, stored by digest, not committed as huge binaries.
                    if len(frame)<10000:
                        prefix.with_suffix(".xlsx").write_bytes(report.content)
                        prefix.with_suffix(".csv").write_bytes(raw.content)
                    evidence.append(dict(path=str(path), sha256=sha(path), context=params,
                                         cleaning=payload["cleaning_summary"], exception_counts={k:len(v) for k,v in payload["exceptions"].items()},
                                         api_status=200, report_status=200, raw_status=200, raw_rows=len(exported), verified_columns=checked,
                                         report_sha256=hashlib.sha256(report.content).hexdigest(),raw_sha256=hashlib.sha256(raw.content).hexdigest(),
                                         report_bytes=len(report.content),raw_bytes=len(raw.content),
                                         chart_rows_verified=chart_count,chart_columns_verified=chart_columns,summary_rows_verified=len(summary),
                                         station_context="not supplied in source filename; NOT_PROVIDED is an export test label" if context.station_start=="NOT_PROVIDED" else "source filename",
                                         metadata_capabilities=payload["metadata_capabilities"],
                                         limitation="Zero-exception fixture cannot establish real peak/level/ID acceptance" if not any(payload["exceptions"].values()) else "Actual exceptions exercised"))
                    dump(output / "real-run-checkpoint.json", evidence)
                    analysis.ANALYSIS_SESSIONS.clear()
                    analysis.LAST_RAW_DF=pd.DataFrame()
                    del frame,workbook,exported,report,raw,response,payload
    pd.DataFrame(rows).to_csv(output / "real-fixture-coverage.csv", index=False, encoding="utf-8-sig")
    return dict(matrix=rows, real_runs=evidence, unassigned_fixtures=[str(p) for p,c in inventory if c is None])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--reference", type=Path, default=ROOT.parent / "TOV640_Analyzer")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    reference = args.reference.resolve()
    head = subprocess.check_output(["git", "-C", str(reference), "rev-parse", "HEAD"], text=True).strip()
    if head != PIN:
        raise SystemExit(f"Reference HEAD drift: {head}; expected {PIN}")
    # Reject local modifications to imported core; unrelated user edits remain untouched.
    for name in ("analyzers.py", "metadata.py", "exporter.py", "data_ingestion.py"):
        relative = f"backend/app/core/{name}"
        pinned = subprocess.check_output(["git", "-C", str(reference), "show", f"{PIN}:{relative}"])
        current = (reference / relative).read_bytes()
        if pinned.replace(b"\r\n", b"\n") != current.replace(b"\r\n", b"\n"):
            raise SystemExit(f"Reference core has uncommitted changes: {relative}")
    args.output.mkdir(parents=True, exist_ok=False)
    output = args.output.resolve()
    # Synthetic metadata is created by a test helper using the existing product exporter conventions.
    meta = output / "canonical-metadata.xlsx"
    boundary = pd.DataFrame(dict(Line=["EAL"], Class=["Mainline"], **{"UP Track FromM":[99990],"UP Track ToM":[100100],"DN Track FromM":[99990],"DN Track ToM":[100100]}))
    direction = pd.DataFrame({"Track Type":["Tangent"],"Track Type FromM":[99990],"Track Type ToM":[100100],"Overlap FromM":[99990],"Overlap ToM":[100100],"Overlap":["Y"],"Tension Length":["A, B"],"Landmark FromM":[99990],"Landmark ToM":[100100],"Landmark":["Contract fixture"]})
    threshold_rows = []
    for name, limits in (("Low Height",[4500,4600]),("High Height",[5900,5800]),("Wire Wear",[9,10.2]),("Stagger Left",[500,400,300]),("Stagger Right",[500,400,300])):
        row = {"Class":"Mainline","Track Type":"both","Exc Type":name}
        prefix = "Stagger" if name.startswith("Stagger") else name
        row.update({f"{prefix} L{i}":v for i,v in enumerate(limits,1)})
        threshold_rows.append(row)
    with pd.ExcelWriter(meta) as writer:
        boundary.to_excel(writer,sheet_name="Exception Boundarys",index=False)
        direction.to_excel(writer,sheet_name="EAL UP",index=False)
        pd.DataFrame(threshold_rows).to_excel(writer,sheet_name="threshold",index=False)
    comparisons = []
    for name, measurements, context in canonical_cases():
        comparisons.append(subprocess_pair(reference, dict(measurements=measurements,metadata=str(meta),context=context), output,name))
    assert all(all(v > 0 for v in case["counts"].values()) for case in comparisons)
    assert {"L1","L2","L3"}.issubset(set().union(*(set(c["levels"]) for c in comparisons)))
    legacy = []
    for line in ("EAL","TML"):
        files = sorted((reference / "docs/test data" / line).rglob("*.datac"))
        if not files:
            legacy.append(dict(line=line,status="blocked_missing_real_datac"))
            continue
        path = files[0]
        track = "DN" if "_D" in path.name else "UP"
        spec = dict(measurements=[],datac=str(path),metadata=str(reference / "config" / f"{line} metadata.xlsx"),context=dict(line=line,section="Mainline",track=track,date_str="20261023"))
        case = subprocess_pair(reference,spec,output,f"legacy-{line}")
        legacy.append(dict(**case, source=str(path),source_sha256=sha(path)))
    real = run_real_csv(output)
    report = dict(reference_commit=head, tolerance=dict(atol=1e-9,rtol=0,strings_ids_columns="exact"),canonical=comparisons,legacy=legacy,**real,
                  limitations=["Canonical fixtures are synthetic; not field acceptance", "TL/wear physical parity blocked until approved intervals/scope", "Only inventoried real fixtures are exercised; no portable/performance claim"])
    dump(output / "contract-results.json", report)
    print(json.dumps(dict(output=str(output),canonical=len(comparisons),real_runs=len(real["real_runs"]),blocked=sum(r["status"]!="pass" for r in real["matrix"])),ensure_ascii=False))


if __name__ == "__main__":
    main()
