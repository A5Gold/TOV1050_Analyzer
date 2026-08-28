from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.core.calculation import wear_historical_import as historical_import
from app.core.calculation.wear_cycle_repository import current_data_version
from app.core.calculation.wear_cycle_types import MetadataInterval
from app.core.calculation.wear_historical_adapters import ADAPTERS
from app.core.calculation.wear_historical_import import (
    HistoricalWorkbookError,
    StaleHistoricalPreviewError,
    discover_historical_workbook,
    preview_historical_workbook,
)
from app.core.database import DatabaseManager


def _workbook_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in sheets.items():
        worksheet = workbook.create_sheet(sheet_name)
        for row in rows:
            worksheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _matrix_workbook() -> bytes:
    return _workbook_bytes({
        "EAL": [
            ["Track", "Tension Length", 202402, 202401],
            [" Up ", "X1", 11.8, 12.1],
        ],
        "LMC": [
            ["Track", "Tension Length", "202401"],
            ["DN", "L01", 12.4],
        ],
        "TML": [
            ["Track", "Tension Length", 202403, 202401, 202402],
            ["DOWN", "M01", 12.7, 12.9, 12.8],
        ],
        "LRL": [["Tension Length", "2024"], [1, 12.5]],
        "Notes": [["ignored"]],
    })


def _metadata() -> tuple[MetadataInterval, ...]:
    return (
        MetadataInterval("X1", "UP", Decimal("0"), Decimal("10"), "EAL UP"),
        MetadataInterval("L01", "DN", Decimal("10"), Decimal("20"), "LMC DN"),
        MetadataInterval("M01", "DN", Decimal("20"), Decimal("30"), "TML DN"),
        MetadataInterval("001A", "UP", Decimal("30"), Decimal("40"), "EAL UP"),
    )


@pytest.fixture
def database(tmp_path):
    manager = DatabaseManager(str(tmp_path / "historical-import.db"))
    try:
        with manager.get_connection() as connection:
            yield connection
    finally:
        manager.close()
        DatabaseManager.reset_instance()


def _insert_record(
    connection,
    *,
    cycle_date: str,
    average: float,
    updated_at: str,
) -> None:
    cycle_id = connection.execute(
        """
        INSERT INTO wire_wear_cycles (
            line_group, line_class, cycle_date, source_type, completeness_state,
            created_at, updated_at
        ) VALUES ('EAL', 'EAL', ?, 'manual', 'complete', ?, ?)
        """,
        (cycle_date, updated_at, updated_at),
    ).lastrowid
    connection.execute(
        """
        INSERT INTO wire_wear_cycle_records (
            cycle_id, line_group, line_class, cycle_date, tension_length,
            track, from_m, to_m, avg_wear_min, wear_percentage,
            created_at, updated_at
        ) VALUES (?, 'EAL', 'EAL', ?, 'X1', 'UP', 0, 10, ?, 0, ?, ?)
        """,
        (cycle_id, cycle_date, average, updated_at, updated_at),
    )
    connection.commit()


def test_discovery_selects_supported_sheets_and_keeps_lrl_disabled():
    discovery = discover_historical_workbook(_matrix_workbook())
    by_name = {sheet.sheet_name: sheet for sheet in discovery.sheets}

    assert [sheet.sheet_name for sheet in discovery.sheets] == [
        "EAL", "LMC", "TML", "LRL", "Notes"
    ]
    assert by_name["EAL"].selected_by_default is True
    assert by_name["LMC"].selected_by_default is True
    assert by_name["TML"].selected_by_default is True
    assert (by_name["LRL"].support, by_name["LRL"].enabled) == (
        "reserved", False
    )
    assert by_name["Notes"].support == "ignored"


def test_discovery_rejects_empty_and_corrupt_workbooks():
    with pytest.raises(HistoricalWorkbookError, match="XLSX bytes"):
        discover_historical_workbook(b"")
    with pytest.raises(HistoricalWorkbookError, match="readable XLSX"):
        discover_historical_workbook(b"not an xlsx file")


def test_preview_classifies_eal_lmc_tml_and_unsorted_tml_months(database):
    _insert_record(
        database,
        cycle_date="2024-01-01",
        average=12.1,
        updated_at="2026-08-07 10:00:00",
    )
    _insert_record(
        database,
        cycle_date="2024-02-01",
        average=12.0,
        updated_at="2026-08-07 11:00:00",
    )
    version = current_data_version(database)

    preview = preview_historical_workbook(
        database,
        _matrix_workbook(),
        ("EAL", "LMC", "TML"),
        _metadata(),
        version,
    )
    result = preview.to_dict()
    candidates = result["candidates"]

    assert result["wire_wear_data_version"] == version
    assert result["totals"] == {
        "skipped": 0,
        "new": 4,
        "update": 1,
        "no_change": 1,
        "duplicate": 0,
        "error": 0,
        "total": 6,
    }
    assert [
        item["key"]["cycle_date"]
        for item in candidates
        if item["source_sheet"] == "TML"
    ] == ["2024-03-01", "2024-01-01", "2024-02-01"]
    lmc = next(item for item in candidates if item["source_sheet"] == "LMC")
    assert lmc["track"] == "DOWN"
    assert lmc["key"]["line_group"] == "EAL"
    assert lmc["key"]["line_class"] == "LMC"
    updated = next(item for item in candidates if item["status"] == "update")
    assert updated["existing_avg_wear_min"] == 12.0
    assert updated["expected_updated_at"] == "2026-08-07 11:00:00"


def test_preview_reports_heading_track_tl_and_cell_diagnostics(database):
    workbook = _workbook_bytes({
        "EAL": [
            [
                "Track", "Tension Length", 202401, "202401", 202402,
                "202413", "bad",
            ],
            ["sideways", "X1", 12.0, None, None, None, None],
            ["UP", None, 12.1, None, None, None, None],
            ["UP", "001A", ".", None, "text", None, None],
        ],
    })

    preview = preview_historical_workbook(
        database, workbook, ("EAL",), _metadata(), current_data_version(database)
    )
    result = preview.to_dict()
    diagnostics = {item["code"]: item for item in result["diagnostics"]}

    assert diagnostics["duplicate_month_heading"]["cell"] == "D1"
    assert diagnostics["invalid_month_heading"]["cell"] in {"F1", "G1"}
    assert diagnostics["invalid_track"]["cell"] == "A2"
    assert diagnostics["empty_tension_length"]["cell"] == "B3"
    assert diagnostics["non_numeric_wear"]["cell"] == "E4"
    assert result["totals"]["skipped"] == 1
    assert result["totals"]["error"] == 6


def test_preview_reports_missing_required_columns(database):
    workbook = _workbook_bytes({
        "EAL": [
            ["Direction", "TL", 202401],
            ["UP", "X1", 12.0],
        ],
    })

    preview = preview_historical_workbook(
        database, workbook, ("EAL",), _metadata(), current_data_version(database)
    )

    assert [(item.code, item.cell) for item in preview.diagnostics] == [
        ("missing_required_column", "A1"),
        ("missing_required_column", "B1"),
    ]
    assert preview.candidates == ()


@pytest.mark.parametrize(
    ("second_value", "expected_status", "expected_issue"),
    [
        (12.0, "duplicate", "duplicate_equal"),
        (11.9, "error", "duplicate_conflict"),
    ],
)
def test_preview_classifies_equal_and_conflicting_duplicate_keys(
    database, second_value, expected_status, expected_issue
):
    workbook = _workbook_bytes({
        "EAL": [
            ["Track", "Tension Length", 202401],
            ["UP", "X1", 12.0],
            ["UP", "X1", second_value],
        ],
    })
    preview = preview_historical_workbook(
        database, workbook, ("EAL",), _metadata(), current_data_version(database)
    )

    assert [candidate.status for candidate in preview.candidates] == [
        "new", expected_status
    ]
    assert preview.candidates[1].issues[0].code == expected_issue


def test_preview_preserves_text_tension_length_and_is_read_only(database):
    workbook = _workbook_bytes({
        "EAL": [
            ["Track", "Tension Length", 202401],
            ["UP", "001A", 12.0],
        ],
    })
    before = {
        "cycles": database.execute("SELECT COUNT(*) FROM wire_wear_cycles").fetchone()[0],
        "records": database.execute(
            "SELECT COUNT(*) FROM wire_wear_cycle_records"
        ).fetchone()[0],
        "version": current_data_version(database),
        "changes": database.total_changes,
    }

    preview = preview_historical_workbook(
        database, workbook, ("EAL",), _metadata(), before["version"]
    )
    after = {
        "cycles": database.execute("SELECT COUNT(*) FROM wire_wear_cycles").fetchone()[0],
        "records": database.execute(
            "SELECT COUNT(*) FROM wire_wear_cycle_records"
        ).fetchone()[0],
        "version": current_data_version(database),
        "changes": database.total_changes,
    }

    assert preview.candidates[0].key["tension_length"] == "001A"
    assert after == before


def test_preview_rejects_unsupported_selection_and_stale_version(database):
    with pytest.raises(HistoricalWorkbookError, match="not supported: LRL"):
        preview_historical_workbook(
            database, _matrix_workbook(), ("LRL",), _metadata(), 0
        )
    with pytest.raises(StaleHistoricalPreviewError, match="data version changed"):
        preview_historical_workbook(
            database, _matrix_workbook(), ("EAL",), _metadata(), 999
        )


def test_preview_batches_classifier_metadata_by_canonical_tension_length(
    database, monkeypatch
):
    tension_lengths = [f"TL{index:02d}" for index in range(40)]
    months = [202401, 202402, 202403, 202404]
    workbook = _workbook_bytes({
        "EAL": [
            ["Track", "Tension Length", *months],
            *[
                ["UP", tension_length, 12.0, 11.9, 11.8, 11.7]
                for tension_length in tension_lengths
            ],
        ],
    })
    metadata = tuple(
        MetadataInterval(
            tension_length,
            "UP",
            Decimal(index * 2),
            Decimal(index * 2 + 1),
            "EAL UP",
        )
        for index, tension_length in enumerate(tension_lengths)
    )
    original_classifier = historical_import.classify_wear_record_candidates
    metadata_sizes: list[int] = []

    def measured_classifier(connection, candidates, candidate_metadata):
        metadata_sizes.append(len(candidate_metadata))
        return original_classifier(connection, candidates, candidate_metadata)

    monkeypatch.setattr(
        historical_import, "classify_wear_record_candidates", measured_classifier
    )
    preview = preview_historical_workbook(
        database,
        workbook,
        ("EAL",),
        metadata,
        current_data_version(database),
    )

    assert len(preview.candidates) == 160
    assert len(metadata_sizes) == 12
    assert max(metadata_sizes) == historical_import._CLASSIFIER_TL_BATCH_SIZE


def test_reference_workbook_has_controlled_adapter_acceptance_entry():
    path = Path(__file__).parents[2] / (
        "docs/Wear Calculator_Database/"
        "EAL TML LRL LMC Historical Avg Wear Min Database .xlsx"
    )
    if not path.exists():
        pytest.skip("reference historical workbook is not available")

    content = path.read_bytes()
    discovery = discover_historical_workbook(content)
    assert [sheet.sheet_name for sheet in discovery.sheets] == [
        "EAL", "LMC", "TML", "LRL"
    ]

    workbook = load_workbook(BytesIO(content), data_only=True, read_only=False)
    try:
        parsed = {name: ADAPTERS[name].parse(workbook[name]) for name in ADAPTERS}
    finally:
        workbook.close()
    assert all(parsed[name].candidates for name in ("EAL", "LMC", "TML"))
    heading_diagnostics = [
        (issue.sheet, issue.cell, issue.original_value)
        for sheet in parsed.values()
        for issue in sheet.diagnostics
        if issue.code == "invalid_month_heading"
    ]
    assert heading_diagnostics == [
        ("EAL", "BT1", "202308A"),
        ("EAL", "BV1", "202309A"),
        ("EAL", "CB1", "202402A"),
        ("EAL", "CF1", "202405A"),
        ("EAL", "CI1", "202407A"),
        ("EAL", "CT1", "202507A"),
    ]
