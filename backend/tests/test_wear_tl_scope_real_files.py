from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.core.calculation.wear_tl_scope import (
    TensionLengthScope,
    classify_tension_length_scope,
)


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
TML_EXPORT = (
    REPOSITORY_ROOT
    / "docs"
    / "Wear Calculator_Database"
    / "2026-05-12_TML_Wear_Cycle.xlsx"
)
EAL_RAC_REPORT = (
    REPOSITORY_ROOT
    / "docs"
    / "test data"
    / "EAL"
    / "Cycle 8"
    / "20260611_EAL_UP_RAC_Exception_Report.xlsx"
)


def _require_fixture(path: Path) -> Path:
    if not path.is_file():
        pytest.skip(f"Real-file fixture unavailable: {path.relative_to(REPOSITORY_ROOT)}")
    return path


@pytest.mark.parametrize(
    ("line_group", "identity", "expected"),
    [
        pytest.param("TML", "M01", TensionLengthScope.MAINLINE, id="tml-m-lower"),
        pytest.param("TML", "M1", TensionLengthScope.MAINLINE, id="tml-m-unpadded"),
        pytest.param("TML", "M33", TensionLengthScope.MAINLINE, id="tml-m-upper"),
        pytest.param("TML", "M34", TensionLengthScope.UNKNOWN, id="tml-m-above"),
        pytest.param("TML", "D01", TensionLengthScope.MAINLINE, id="tml-d-lower"),
        pytest.param("TML", "D60", TensionLengthScope.MAINLINE, id="tml-d-upper"),
        pytest.param("TML", "D61", TensionLengthScope.UNKNOWN, id="tml-d-above"),
        pytest.param("TML", "U01", TensionLengthScope.MAINLINE, id="tml-u-lower"),
        pytest.param("TML", "U55", TensionLengthScope.MAINLINE, id="tml-u-upper"),
        pytest.param("TML", "U56", TensionLengthScope.UNKNOWN, id="tml-u-above"),
        pytest.param("TML", "U07/1", TensionLengthScope.MAINLINE, id="tml-u-slash"),
        pytest.param("TML", "U7/1", TensionLengthScope.MAINLINE, id="tml-u-slash-unpadded"),
        pytest.param("TML", "U07/3", TensionLengthScope.UNKNOWN, id="tml-u-slash-invalid"),
        pytest.param("TML", "K01", TensionLengthScope.MAINLINE, id="tml-k-lower"),
        pytest.param("TML", "K16", TensionLengthScope.MAINLINE, id="tml-k-upper"),
        pytest.param("TML", "K17", TensionLengthScope.UNKNOWN, id="tml-k-above"),
        pytest.param("TML", "01", TensionLengthScope.MAINLINE, id="tml-numeric-padded"),
        pytest.param("TML", "1", TensionLengthScope.MAINLINE, id="tml-numeric-unpadded"),
        pytest.param("TML", "27", TensionLengthScope.MAINLINE, id="tml-27-mainline"),
        pytest.param("TML", "27A", TensionLengthScope.SIDING, id="tml-27a-siding"),
        pytest.param("TML", "74", TensionLengthScope.MAINLINE, id="tml-numeric-upper"),
        pytest.param("TML", "75", TensionLengthScope.UNKNOWN, id="tml-numeric-above"),
        pytest.param("TML", "MD1", TensionLengthScope.SIDING, id="tml-md-lower"),
        pytest.param("TML", "MD22", TensionLengthScope.SIDING, id="tml-md-upper"),
        pytest.param("TML", "MD23", TensionLengthScope.UNKNOWN, id="tml-md-above"),
        pytest.param("TML", "MX4", TensionLengthScope.SIDING, id="tml-mx-unpadded"),
        pytest.param("TML", " mx04 ", TensionLengthScope.SIDING, id="tml-mx-padded-whitespace"),
        pytest.param("TML", "mX09", TensionLengthScope.SIDING, id="tml-mx-case-insensitive"),
        pytest.param("TML", "MX11", TensionLengthScope.SIDING, id="tml-mx-upper"),
        pytest.param("TML", "MX12", TensionLengthScope.UNKNOWN, id="tml-mx-above"),
        pytest.param("TML", "X01", TensionLengthScope.SIDING, id="tml-x-lower"),
        pytest.param("TML", "X29", TensionLengthScope.SIDING, id="tml-x-upper"),
        pytest.param("TML", "X30", TensionLengthScope.UNKNOWN, id="tml-x-above"),
        pytest.param("TML", "KX1", TensionLengthScope.SIDING, id="tml-kx-lower"),
        pytest.param("TML", "KX7", TensionLengthScope.SIDING, id="tml-kx-upper"),
        pytest.param("TML", "KX8", TensionLengthScope.UNKNOWN, id="tml-kx-above"),
        pytest.param("TML", "MP24", TensionLengthScope.SIDING, id="tml-mp24"),
        pytest.param("TML", "MP23", TensionLengthScope.UNKNOWN, id="tml-mp-unknown"),
        pytest.param("TML", "EM1", TensionLengthScope.SIDING, id="tml-em1"),
        pytest.param("TML", "EM2", TensionLengthScope.UNKNOWN, id="tml-em-unknown"),
        pytest.param("TML", "PT1", TensionLengthScope.SIDING, id="tml-pt-lower"),
        pytest.param("TML", "PT2", TensionLengthScope.SIDING, id="tml-pt-upper"),
        pytest.param("TML", "PT3", TensionLengthScope.UNKNOWN, id="tml-pt-above"),
        pytest.param("TML", "CT1", TensionLengthScope.SIDING, id="tml-ct-lower"),
        pytest.param("TML", "CT3", TensionLengthScope.SIDING, id="tml-ct-upper"),
        pytest.param("TML", "CT4", TensionLengthScope.UNKNOWN, id="tml-ct-above"),
        pytest.param("TML", "W1-1", TensionLengthScope.SIDING, id="tml-w1-1"),
        pytest.param("TML", "101", TensionLengthScope.SIDING, id="tml-special-101"),
        pytest.param("TML", "103", TensionLengthScope.SIDING, id="tml-special-103"),
        pytest.param("TML", "104", TensionLengthScope.UNKNOWN, id="tml-special-above"),
        pytest.param(
            "TML",
            "Neutral Section",
            TensionLengthScope.SIDING,
            id="tml-neutral-section",
        ),
        pytest.param("EAL", "H01", TensionLengthScope.MAINLINE, id="eal-h-lower"),
        pytest.param("EAL", "H52", TensionLengthScope.MAINLINE, id="eal-h-upper"),
        pytest.param("EAL", "H53", TensionLengthScope.UNKNOWN, id="eal-h-above"),
        pytest.param("EAL", "H23D", TensionLengthScope.MAINLINE, id="eal-h23d"),
        pytest.param("EAL", "H26D", TensionLengthScope.MAINLINE, id="eal-h26d"),
        pytest.param("EAL", "H24D", TensionLengthScope.UNKNOWN, id="eal-h-suffix-unknown"),
        pytest.param("EAL", "D31", TensionLengthScope.MAINLINE, id="eal-d31-mainline"),
        pytest.param("EAL", "D32", TensionLengthScope.SIDING, id="eal-d32-siding"),
        pytest.param("EAL", "D33", TensionLengthScope.SIDING, id="eal-d33-siding"),
        pytest.param("EAL", "D34", TensionLengthScope.UNKNOWN, id="eal-d34-unknown"),
        pytest.param("EAL", "U01", TensionLengthScope.MAINLINE, id="eal-u-lower"),
        pytest.param("EAL", "U55", TensionLengthScope.MAINLINE, id="eal-u-upper"),
        pytest.param("EAL", "U56", TensionLengthScope.UNKNOWN, id="eal-u-above"),
        pytest.param("EAL", "U31/2", TensionLengthScope.MAINLINE, id="eal-u-slash"),
        pytest.param("EAL", "U31/3", TensionLengthScope.UNKNOWN, id="eal-u-slash-invalid"),
        pytest.param("EAL", "K01", TensionLengthScope.MAINLINE, id="eal-k-lower"),
        pytest.param("EAL", "K16", TensionLengthScope.MAINLINE, id="eal-k-upper"),
        pytest.param("EAL", "K17", TensionLengthScope.UNKNOWN, id="eal-k-above"),
        pytest.param("EAL", "1", TensionLengthScope.MAINLINE, id="eal-numeric-lower"),
        pytest.param("EAL", "76", TensionLengthScope.MAINLINE, id="eal-numeric-upper"),
        pytest.param("EAL", "77", TensionLengthScope.UNKNOWN, id="eal-numeric-above"),
        pytest.param("EAL", "69B", TensionLengthScope.MAINLINE, id="eal-69b"),
        pytest.param("EAL", "70A", TensionLengthScope.MAINLINE, id="eal-70a"),
        pytest.param("EAL", "27A", TensionLengthScope.UNKNOWN, id="eal-27a-cross-line"),
        pytest.param("EAL", "L01", TensionLengthScope.MAINLINE, id="eal-l-lower"),
        pytest.param("EAL", "L20", TensionLengthScope.MAINLINE, id="eal-l-upper"),
        pytest.param("EAL", "L21", TensionLengthScope.UNKNOWN, id="eal-l-above"),
        pytest.param("EAL", "T2", TensionLengthScope.MAINLINE, id="eal-t2"),
        pytest.param("EAL", "T3", TensionLengthScope.MAINLINE, id="eal-t3"),
        pytest.param("EAL", "T4", TensionLengthScope.UNKNOWN, id="eal-t-unknown"),
        pytest.param("EAL", "X01", TensionLengthScope.SIDING, id="eal-x-lower"),
        pytest.param("EAL", "X36", TensionLengthScope.MAINLINE, id="eal-x36-mainline"),
        pytest.param("EAL", "X37", TensionLengthScope.MAINLINE, id="eal-x37-mainline"),
        pytest.param("EAL", "X38", TensionLengthScope.SIDING, id="eal-x38-siding"),
        pytest.param("EAL", "X39", TensionLengthScope.SIDING, id="eal-x39-siding"),
        pytest.param("EAL", "X40", TensionLengthScope.UNKNOWN, id="eal-x40-unknown"),
        pytest.param("EAL", "X50", TensionLengthScope.MAINLINE, id="eal-x50-mainline"),
        pytest.param("EAL", "X53", TensionLengthScope.MAINLINE, id="eal-x53-mainline"),
        pytest.param("EAL", "X54", TensionLengthScope.MAINLINE, id="eal-x54-mainline"),
        pytest.param("EAL", "X51", TensionLengthScope.UNKNOWN, id="eal-x51-unknown"),
        pytest.param("EAL", "HX02", TensionLengthScope.SIDING, id="eal-hx02-siding"),
        pytest.param("EAL", "HX2", TensionLengthScope.SIDING, id="eal-hx2-unpadded"),
        pytest.param("EAL", "HX1", TensionLengthScope.UNKNOWN, id="eal-hx-unknown"),
        pytest.param("EAL", "LX1", TensionLengthScope.SIDING, id="eal-lx-lower"),
        pytest.param("EAL", "LX001", TensionLengthScope.SIDING, id="eal-lx-padded"),
        pytest.param("EAL", "LX0", TensionLengthScope.UNKNOWN, id="eal-lx-zero"),
        pytest.param("EAL", "EM1", TensionLengthScope.SIDING, id="eal-em1"),
        pytest.param("EAL", "EM2", TensionLengthScope.UNKNOWN, id="eal-em-unknown"),
        pytest.param(
            "EAL",
            "Neutral Section",
            TensionLengthScope.SIDING,
            id="eal-neutral-section",
        ),
        pytest.param("TML", "69B", TensionLengthScope.UNKNOWN, id="tml-69b-cross-line"),
        pytest.param("TML", "X36", TensionLengthScope.UNKNOWN, id="tml-x36-cross-line"),
        pytest.param("EAL", "69B-extra", TensionLengthScope.UNKNOWN, id="anchored-suffix"),
        pytest.param("TML", "pre-MD1", TensionLengthScope.UNKNOWN, id="anchored-prefix"),
        pytest.param("EAL", "", TensionLengthScope.UNKNOWN, id="empty-unknown"),
        pytest.param("TML", "new-yard-tl", TensionLengthScope.UNKNOWN, id="drift-unknown"),
    ],
)
def test_classifier_collision_and_boundary_matrix(
    line_group: str,
    identity: str,
    expected: TensionLengthScope,
) -> None:
    assert classify_tension_length_scope(line_group, identity) is expected


def test_supplied_tml_export_records_are_all_mainline() -> None:
    workbook = load_workbook(_require_fixture(TML_EXPORT), read_only=True, data_only=True)
    try:
        worksheet = workbook["Wear Records"]
        rows = worksheet.iter_rows(values_only=True)
        headers = {name: index for index, name in enumerate(next(rows))}
        records = list(rows)
    finally:
        workbook.close()

    assert len(records) == 164
    assert {row[headers["Line"]] for row in records} == {"TML"}
    assert {row[headers["Track"]] for row in records} == {"UP", "DN"}
    assert {
        classify_tension_length_scope("TML", row[headers["Tension Length"]])
        for row in records
    } == {TensionLengthScope.MAINLINE}


def test_supplied_tml_export_diagnostic_gaps_are_known_siding_identities() -> None:
    workbook = load_workbook(_require_fixture(TML_EXPORT), read_only=True, data_only=True)
    try:
        worksheet = workbook["Cycle Coverage"]
        rows = worksheet.iter_rows(values_only=True)
        headers = {name: index for index, name in enumerate(next(rows))}
        diagnostic_gaps = {
            identity.strip()
            for row in rows
            for identity in str(row[headers["Diagnostic Gaps"]] or "").split(";")
            if identity.strip()
        }
    finally:
        workbook.close()

    assert {"27A", "MD16", "MX04", "MX09", "KX1", "X01", "103"} <= diagnostic_gaps
    assert {
        classify_tension_length_scope("TML", identity)
        for identity in diagnostic_gaps
    } == {TensionLengthScope.SIDING}


def test_eal_rac_approach_zone_preserves_exact_composite_signature() -> None:
    workbook = load_workbook(_require_fixture(EAL_RAC_REPORT), read_only=True, data_only=True)
    try:
        worksheet = workbook["ChartData"]
        rows = worksheet.iter_rows(min_row=1, max_row=217, values_only=True)
        headers = {name: index for index, name in enumerate(next(rows))}
        approach_rows = list(rows)
    finally:
        workbook.close()

    assert len(approach_rows) == 216
    assert {row[headers["line"]] for row in approach_rows} == {"EAL"}
    assert {row[headers["track"]] for row in approach_rows} == {"UP"}
    assert {row[headers["Section"]] for row in approach_rows} == {"RAC"}
    assert approach_rows[0][headers["Chainage"]] == 111301.0
    assert approach_rows[-1][headers["Chainage"]] == 111354.75
    assert {row[headers["Tension Length"]] for row in approach_rows} == {"23, 25"}

    ordered_signature = [
        identity.strip()
        for identity in approach_rows[0][headers["Tension Length"]].split(",")
    ]
    assert ordered_signature == ["23", "25"]
    assert [
        classify_tension_length_scope("EAL", identity) for identity in ordered_signature
    ] == [TensionLengthScope.MAINLINE, TensionLengthScope.MAINLINE]
