from __future__ import annotations

import json
import shutil
import sqlite3
import sys
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
import pytest

sys.path.insert(0, str(Path(__file__).parents[1]))

from app.core.database import DatabaseManager, check_1_year_records
from app.core.database_record_seed import (
    EXPECTED_HEADERS,
    EXPECTED_RECORD_COUNTS,
    EXPECTED_SECTION_COUNTS,
    LEGACY_EXPECTED_RECORD_COUNTS,
    LEGACY_SEED_VERSION,
    SEED_VERSION,
    WORKBOOKS,
    load_database_record_seed,
    seed_database_records,
)


REPO_ROOT = Path(__file__).parents[2]
CONFIG_DIR = REPO_ROOT / "config"


def _metadata(conn: sqlite3.Connection) -> dict[str, str]:
    return dict(conn.execute(
        "SELECT key, value FROM system_metadata WHERE key LIKE 'database_record_seed_%'"
    ).fetchall())


def _record_counts(conn: sqlite3.Connection) -> tuple[int, dict[str, int]]:
    total = conn.execute(
        "SELECT COUNT(*) FROM saved_repeated_exceptions"
    ).fetchone()[0]
    by_line = dict(conn.execute(
        "SELECT line, COUNT(*) FROM saved_repeated_exceptions GROUP BY line"
    ).fetchall())
    return total, by_line


def _insert_legacy_v1_baseline(
    conn: sqlite3.Connection,
    *,
    actual_counts: dict[str, int] | None = None,
    metadata_count: int = 2706,
    metadata_counts: dict[str, int] | None = None,
    version: str = LEGACY_SEED_VERSION,
) -> None:
    actual_counts = actual_counts or dict(LEGACY_EXPECTED_RECORD_COUNTS)
    metadata_counts = metadata_counts or dict(LEGACY_EXPECTED_RECORD_COUNTS)
    rows = []
    for line, count in actual_counts.items():
        rows.extend(
            (
                f"legacy-{line.lower()}-{index}",
                "Wire Wear",
                "L1",
                float(index),
                float(index + 1),
                line,
                "UP",
                "20250801",
                "built-in-seed",
            )
            for index in range(count)
        )
    conn.executemany(
        """
        INSERT INTO saved_repeated_exceptions (
            exception_id, exception_type, level, from_m, to_m,
            line, track, date_str, saved_by
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        rows,
    )
    conn.executemany(
        """
        INSERT INTO system_metadata (key, value, description)
        VALUES (?, ?, 'test v1 seed metadata')
        ON CONFLICT(key) DO UPDATE SET value = excluded.value
        """,
        [
            ("database_record_seed_version", version),
            ("database_record_seed_record_count", str(metadata_count)),
            (
                "database_record_seed_counts_by_line",
                json.dumps(metadata_counts, sort_keys=True),
            ),
        ],
    )
    conn.commit()


def _create_empty_database(path: Path) -> None:
    DatabaseManager.reset_instance()
    manager = DatabaseManager(str(path), seed_database_records=False)
    manager.close()
    DatabaseManager.reset_instance()


def _connect(path: Path) -> sqlite3.Connection:
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def _copy_seed_assets(target_config: Path) -> Path:
    source_dir = CONFIG_DIR / "database-records"
    target_dir = target_config / "database-records"
    target_dir.mkdir(parents=True)
    for filename in (*WORKBOOKS.values(), "database-record-seed-quality.json"):
        shutil.copy2(source_dir / filename, target_dir / filename)
    return target_dir


def test_real_seed_loader_has_expected_counts_and_hashes():
    bundle = load_database_record_seed(CONFIG_DIR)
    quality = json.loads(
        (CONFIG_DIR / "database-records" / "database-record-seed-quality.json").read_text(
            encoding="utf-8"
        )
    )

    assert len(bundle.records) == 2800
    assert bundle.counts_by_line == EXPECTED_RECORD_COUNTS
    assert bundle.counts_by_section == EXPECTED_SECTION_COUNTS
    assert bundle.seed_version == SEED_VERSION
    expected_hashes = {
        "EAL-1-year-database-record.xlsx": (
            "ddc3d385c65e5a2475e8a1070fe8eff401f7967866cef11fc370c0a975f80fad"
        ),
        "TML-1-year-database-record.xlsx": (
            "1d9d989663f9ee67bd60c4bc7d9496a823d263fd3be67e341838d780d1a94165"
        ),
    }
    assert bundle.output_sha256 == expected_hashes
    assert bundle.source_sha256 == expected_hashes
    assert quality["schema"] == "database-record-seed-quality-v2"
    assert quality["source_record_count"] == 2800
    assert quality["excluded_record_count"] == 0
    assert quality["provenance"]["base_seed_version"] == LEGACY_SEED_VERSION
    rac_duplicate = next(
        exclusion
        for exclusion in quality["historical_exclusions"]
        if exclusion["key"].startswith("20251009_EAL_RAC_DN_SL9|")
    )
    assert rac_duplicate["rule"] == "exact_duplicate"


def test_real_seed_records_match_import_format_and_have_unique_keys():
    bundle = load_database_record_seed(CONFIG_DIR)
    keys = [
        (record["exception_id"], record["line"], record["track"], record["date_str"])
        for record in bundle.records
    ]

    assert len(keys) == len(set(keys))
    assert all(record["exception_id"] for record in bundle.records)
    assert all(record["max_location"] is not None for record in bundle.records)
    assert all(
        isinstance(record["date_str"], str)
        and len(record["date_str"]) == 8
        and record["date_str"].isdigit()
        and record["task_run_date"] == record["date_str"]
        for record in bundle.records
    )
    assert {record["line"] for record in bundle.records} == {"EAL", "TML"}


def test_empty_database_is_seeded_with_metadata_and_reopen_is_idempotent(tmp_path):
    db_path = tmp_path / "analysis.db"
    _create_empty_database(db_path)

    with _connect(db_path) as conn:
        first = seed_database_records(conn, CONFIG_DIR)
        second = seed_database_records(conn, CONFIG_DIR)
        count = conn.execute(
            "SELECT COUNT(*) FROM saved_repeated_exceptions"
        ).fetchone()[0]
        metadata = _metadata(conn)
        conn.execute(
            """DELETE FROM saved_repeated_exceptions
               WHERE record_id = (SELECT MIN(record_id) FROM saved_repeated_exceptions)"""
        )
        third = seed_database_records(conn, CONFIG_DIR)
        modified_count = conn.execute(
            "SELECT COUNT(*) FROM saved_repeated_exceptions"
        ).fetchone()[0]

    assert first.status == "seeded"
    assert first.inserted_count == 2800
    assert first.counts_by_line == EXPECTED_RECORD_COUNTS
    assert second.status == "skipped_up_to_date"
    assert third.status == "skipped_up_to_date"
    assert second.inserted_count == 0
    assert count == 2800
    assert modified_count == 2799
    assert metadata["database_record_seed_version"] == SEED_VERSION
    assert metadata["database_record_seed_record_count"] == "2800"
    assert json.loads(metadata["database_record_seed_counts_by_line"]) == EXPECTED_RECORD_COUNTS
    assert metadata["database_record_seed_status"] == "seeded"
    assert "database_record_seed_replacement_source_version" not in metadata
    assert set(json.loads(metadata["database_record_seed_output_sha256"])) == set(
        WORKBOOKS.values()
    )
    assert len(json.loads(metadata["database_record_seed_source_sha256"])) == 2


def test_exact_v1_seed_is_replaced_once_without_touching_other_tables(tmp_path):
    db_path = tmp_path / "analysis.db"
    _create_empty_database(db_path)

    with _connect(db_path) as conn:
        _insert_legacy_v1_baseline(conn)
        conn.execute(
            """
            INSERT INTO analysis_sessions (id, line, section, track, date_str)
            VALUES ('preserve-me', 'EAL', 'Mainline', 'UP', '20260819')
            """
        )
        conn.commit()

        first = seed_database_records(conn, CONFIG_DIR)
        second = seed_database_records(conn, CONFIG_DIR)
        total, by_line = _record_counts(conn)
        metadata = _metadata(conn)
        preserved = conn.execute(
            "SELECT id, line, section FROM analysis_sessions WHERE id = 'preserve-me'"
        ).fetchone()
        foreign_key_errors = conn.execute("PRAGMA foreign_key_check").fetchall()

    assert first.status == "replaced"
    assert first.inserted_count == 2800
    assert second.status == "skipped_up_to_date"
    assert total == 2800
    assert by_line == EXPECTED_RECORD_COUNTS
    assert metadata["database_record_seed_version"] == SEED_VERSION
    assert metadata["database_record_seed_status"] == "replaced"
    assert (
        metadata["database_record_seed_replacement_source_version"]
        == LEGACY_SEED_VERSION
    )
    assert tuple(preserved) == ("preserve-me", "EAL", "Mainline")
    assert foreign_key_errors == []


def test_database_manager_startup_replaces_exact_v1_seed(tmp_path):
    db_path = tmp_path / "analysis.db"
    _create_empty_database(db_path)
    with _connect(db_path) as conn:
        _insert_legacy_v1_baseline(conn)

    DatabaseManager.reset_instance()
    manager = DatabaseManager(str(db_path), seed_database_records=True)
    try:
        with manager.get_connection() as conn:
            total, by_line = _record_counts(conn)
            metadata = _metadata(conn)
    finally:
        manager.close()
        DatabaseManager.reset_instance()

    assert total == 2800
    assert by_line == EXPECTED_RECORD_COUNTS
    assert metadata["database_record_seed_version"] == SEED_VERSION
    assert metadata["database_record_seed_status"] == "replaced"


@pytest.mark.parametrize(
    ("metadata_count", "metadata_counts", "actual_counts"),
    [
        (2705, dict(LEGACY_EXPECTED_RECORD_COUNTS), dict(LEGACY_EXPECTED_RECORD_COUNTS)),
        (2706, {"EAL": 1731, "TML": 974}, dict(LEGACY_EXPECTED_RECORD_COUNTS)),
        (2706, dict(LEGACY_EXPECTED_RECORD_COUNTS), {"EAL": 1731, "TML": 974}),
    ],
)
def test_v1_count_drift_is_not_replaced(
    tmp_path, metadata_count, metadata_counts, actual_counts
):
    db_path = tmp_path / "analysis.db"
    _create_empty_database(db_path)

    with _connect(db_path) as conn:
        _insert_legacy_v1_baseline(
            conn,
            actual_counts=actual_counts,
            metadata_count=metadata_count,
            metadata_counts=metadata_counts,
        )
        before = _record_counts(conn)
        before_metadata = _metadata(conn)
        result = seed_database_records(conn, CONFIG_DIR)
        after = _record_counts(conn)
        after_metadata = _metadata(conn)

    assert result.status == "skipped_non_empty"
    assert after == before
    assert after_metadata == before_metadata


def test_non_empty_database_is_never_modified(tmp_path):
    db_path = tmp_path / "analysis.db"
    _create_empty_database(db_path)

    with _connect(db_path) as conn:
        conn.execute(
            """
            INSERT INTO saved_repeated_exceptions (
                exception_id, exception_type, level, from_m, to_m,
                line, track, date_str, action
            ) VALUES ('local-only', 'Wire Wear', 'L1', 10, 20,
                      'EAL', 'UP', '20260801', 'Keep monitoring')
            """
        )
        conn.commit()
        before = tuple(conn.execute(
            "SELECT exception_id, action FROM saved_repeated_exceptions"
        ).fetchone())
        result = seed_database_records(conn, CONFIG_DIR)
        rows = conn.execute(
            "SELECT exception_id, action FROM saved_repeated_exceptions"
        ).fetchall()
        seed_metadata = conn.execute(
            "SELECT COUNT(*) FROM system_metadata WHERE key LIKE 'database_record_seed_%'"
        ).fetchone()[0]

    assert result.status == "skipped_non_empty"
    assert [tuple(row) for row in rows] == [before]
    assert seed_metadata == 0


def test_missing_seed_assets_fail_without_breaking_database_startup(tmp_path):
    db_path = tmp_path / "analysis.db"
    _create_empty_database(db_path)

    with _connect(db_path) as conn:
        _insert_legacy_v1_baseline(conn)
        before = _record_counts(conn)
        before_metadata = _metadata(conn)
        result = seed_database_records(conn, tmp_path / "missing-config")
        after = _record_counts(conn)
        after_metadata = _metadata(conn)
        schema_version = conn.execute(
            "SELECT value FROM system_metadata WHERE key = 'schema_version'"
        ).fetchone()[0]

    assert result.status == "failed"
    assert result.error
    assert after == before
    assert after_metadata == before_metadata
    assert schema_version


def test_unknown_workbook_header_is_rejected_before_insert(tmp_path):
    config_dir = tmp_path / "config"
    asset_dir = _copy_seed_assets(config_dir)
    workbook_path = asset_dir / WORKBOOKS["EAL"]
    workbook = load_workbook(workbook_path)
    workbook["Database Records"].cell(row=1, column=1, value="Unknown Header")
    workbook.save(workbook_path)
    workbook.close()
    (asset_dir / "database-record-seed-quality.json").unlink()

    db_path = tmp_path / "analysis.db"
    _create_empty_database(db_path)
    with _connect(db_path) as conn:
        _insert_legacy_v1_baseline(conn)
        before = _record_counts(conn)
        before_metadata = _metadata(conn)
        result = seed_database_records(conn, config_dir)
        after = _record_counts(conn)
        after_metadata = _metadata(conn)

    assert result.status == "failed"
    assert "unexpected header" in result.error
    assert after == before
    assert after_metadata == before_metadata


def test_corrupt_workbook_is_rejected_before_insert(tmp_path):
    config_dir = tmp_path / "config"
    asset_dir = config_dir / "database-records"
    asset_dir.mkdir(parents=True)
    (asset_dir / WORKBOOKS["EAL"]).write_bytes(b"not-an-xlsx")

    db_path = tmp_path / "analysis.db"
    _create_empty_database(db_path)
    with _connect(db_path) as conn:
        _insert_legacy_v1_baseline(conn)
        before = _record_counts(conn)
        before_metadata = _metadata(conn)
        result = seed_database_records(conn, config_dir)
        after = _record_counts(conn)
        after_metadata = _metadata(conn)

    assert result.status == "failed"
    assert "cannot read workbook" in result.error
    assert after == before
    assert after_metadata == before_metadata


def test_hash_mismatch_is_rejected_before_v1_delete(tmp_path):
    config_dir = tmp_path / "config"
    asset_dir = _copy_seed_assets(config_dir)
    manifest_path = asset_dir / "database-record-seed-quality.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["output_sha256"][WORKBOOKS["EAL"]] = "0" * 64
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")

    db_path = tmp_path / "analysis.db"
    _create_empty_database(db_path)
    with _connect(db_path) as conn:
        _insert_legacy_v1_baseline(conn)
        before = _record_counts(conn)
        before_metadata = _metadata(conn)
        result = seed_database_records(conn, config_dir)
        after = _record_counts(conn)
        after_metadata = _metadata(conn)

    assert result.status == "failed"
    assert "hash" in result.error
    assert after == before
    assert after_metadata == before_metadata


class _FailAfterFirstInsertConnection:
    def __init__(self, conn: sqlite3.Connection):
        self._conn = conn

    def execute(self, *args, **kwargs):
        return self._conn.execute(*args, **kwargs)

    def executemany(self, query, parameters):
        rows = list(parameters)
        self._conn.execute(query, rows[0])
        raise sqlite3.OperationalError("simulated bulk insert failure")


class _FailOnSeedMetadataConnection:
    def __init__(self, conn: sqlite3.Connection):
        self._conn = conn

    def execute(self, query, parameters=()):
        if (
            "INSERT INTO system_metadata" in query
            and parameters
            and parameters[0] == "database_record_seed_record_count"
        ):
            raise sqlite3.OperationalError("simulated metadata failure")
        return self._conn.execute(query, parameters)

    def executemany(self, *args, **kwargs):
        return self._conn.executemany(*args, **kwargs)


def test_mid_insert_failure_rolls_back_complete_v1_baseline(tmp_path):
    db_path = tmp_path / "analysis.db"
    _create_empty_database(db_path)

    with _connect(db_path) as conn:
        _insert_legacy_v1_baseline(conn)
        before = _record_counts(conn)
        before_metadata = _metadata(conn)
        result = seed_database_records(_FailAfterFirstInsertConnection(conn), CONFIG_DIR)
        after = _record_counts(conn)
        after_metadata = _metadata(conn)

    assert result.status == "failed"
    assert "simulated bulk insert failure" in result.error
    assert after == before
    assert after_metadata == before_metadata


def test_metadata_failure_rolls_back_complete_v1_baseline(tmp_path):
    db_path = tmp_path / "analysis.db"
    _create_empty_database(db_path)

    with _connect(db_path) as conn:
        _insert_legacy_v1_baseline(conn)
        before = _record_counts(conn)
        before_metadata = _metadata(conn)
        result = seed_database_records(_FailOnSeedMetadataConnection(conn), CONFIG_DIR)
        after = _record_counts(conn)
        after_metadata = _metadata(conn)

    assert result.status == "failed"
    assert "simulated metadata failure" in result.error
    assert after == before
    assert after_metadata == before_metadata


def test_check_1_year_records_matches_a_real_seed_record(tmp_path):
    db_path = tmp_path / "analysis.db"
    _create_empty_database(db_path)

    with _connect(db_path) as conn:
        assert seed_database_records(conn, CONFIG_DIR).status == "seeded"
        seeded = conn.execute(
            """
            SELECT record_id, exception_id, exception_type, section, line, track,
                   from_m, to_m, max_location, task_run_date
            FROM saved_repeated_exceptions
            WHERE max_location BETWEEN from_m AND to_m
              AND task_run_date IS NOT NULL
            ORDER BY record_id
            LIMIT 1
            """
        ).fetchone()
        assert seeded is not None
        current_date = (
            datetime.strptime(seeded["task_run_date"], "%Y%m%d") + timedelta(days=30)
        ).strftime("%Y%m%d")
        current_id = "seed-integration-current"
        result = check_1_year_records(
            conn,
            [{
                "id": current_id,
                "exception type": seeded["exception_type"],
                "level": "L1",
                "FromM": seeded["from_m"],
                "ToM": seeded["to_m"],
                "maxLocation": seeded["max_location"],
                "section": seeded["section"],
                "task_run_date": current_date,
                "action": "Pending",
            }],
            seeded["line"],
            seeded["track"],
        )
        stored_reoccurrence = conn.execute(
            "SELECT reoccurrence_id FROM saved_repeated_exceptions WHERE record_id = ?",
            (seeded["record_id"],),
        ).fetchone()[0]

    assert result[0]["action"] == "No action required (Verified within 1 year)"
    assert result[0]["reoccurrence_id"] == seeded["exception_id"]
    assert current_id in stored_reoccurrence
