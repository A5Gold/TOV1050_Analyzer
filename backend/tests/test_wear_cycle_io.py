import json
from contextlib import closing
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import sqlite3

import pytest
from openpyxl import load_workbook

from app.core.calculation.wear_cycle_repository import (
    StaleDataVersionError,
    current_data_version,
    save_analysis_cycle,
)
from app.core.calculation.wear_cycle_types import (
    AggregatedWearRecord,
    BusinessKey,
    ConflictPreview,
    CyclePreview,
    MetadataInterval,
    PhysicalInterval,
    SegmentCoverage,
)
from app.core.database import DatabaseManager


@pytest.fixture()
def conn(tmp_path):
    manager = DatabaseManager(str(tmp_path / "wear-cycle-io.db"))
    try:
        with manager.get_connection() as connection:
            yield connection
    finally:
        DatabaseManager.reset_instance()


def _metadata(*, to_m=200):
    return (
        MetadataInterval("28", "UP", Decimal("100"), Decimal(str(to_m)), "EAL UP"),
        MetadataInterval("29", "UP", Decimal("201"), Decimal("300"), "EAL UP"),
    )


def _seed_cycle(conn):
    records = (
        AggregatedWearRecord(
            BusinessKey("EAL", date(2026, 5, 28), "29"), "UP",
            Decimal("201"), Decimal("300"), 11.123456789, 12.3456789,
            0.123456789, False, (), ("EAL_U2.xlsx",),
        ),
        AggregatedWearRecord(
            BusinessKey("EAL", date(2026, 5, 28), "28"), "UP",
            Decimal("100"), Decimal("200"), 11.987654321, 6.7890123,
            None, True, ("conflict-28",), ("EAL_U1-a.xlsx", "EAL_U1-b.xlsx"),
        ),
    )
    segments = tuple(
        SegmentCoverage(
            name, name == "U1", 100.0 if name == "U1" else 0.0,
            () if name == "U1" else ("missing",),
            ("EAL_U1-a.xlsx",) if name == "U1" else (),
            (date(2026, 5, 27),) if name == "U1" else (),
        )
        for name in ("U1", "U2", "U3", "D1", "D2", "D3", "LOW S1", "RAC UP", "RAC DN", "LMC UP", "LMC DN")
    )
    conflict = ConflictPreview(
        "conflict-28", "measurement-28",
        (("EAL_U1-a.xlsx", 11.99), ("EAL_U1-b.xlsx", 11.98)),
        11.98, True,
    )
    preview = CyclePreview(
        "EAL", date(2026, 5, 28), records, segments, (conflict,), (), (), True,
        datetime(2026, 5, 28, tzinfo=timezone.utc),
    )
    save_analysis_cycle(conn, preview, expected_data_version=0)


def _package_for_record(*, updated_at="2026-07-10T12:00:00Z", avg=10.5):
    return {
        "schema": "wear-cycle-v1",
        "package_id": "package-1",
        "exported_at": "2026-07-10T12:00:00Z",
        "source_workstation": "TOV640-01",
        "metadata_fingerprint": {"EAL": "same", "TML": "same"},
        "cycles": [{
            "line_group": "EAL", "line_class": "EAL",
            "cycle_date": "2026-05-28", "source_type": "sync",
            "acquisition_date_from": None, "acquisition_date_to": None,
            "completeness_state": "incomplete", "source_lineage": [],
            "created_at": updated_at, "updated_at": updated_at,
        }],
        "records": [{
            "line_group": "EAL", "line_class": "EAL",
            "cycle_date": "2026-05-28", "tension_length": "28.0",
            "track": "UP", "from_m": 100.0, "to_m": 200.0,
            "avg_wear_min": avg, "wear_percentage": 999.0, "measurement_sd": 0.25,
            "source_lineage": ["remote.xlsx"], "created_at": updated_at, "updated_at": updated_at,
        }],
        "segments": [],
        "conflict_decisions": [],
        "tombstones": [],
    }


def test_excel_report_contains_exact_approved_sheets_and_precise_sorted_records(conn):
    from app.core.calculation.wear_cycle_io import build_excel_report

    _seed_cycle(conn)
    workbook = load_workbook(
        BytesIO(build_excel_report(
            conn, line_group="EAL", cycle_date="2026-05-28",
            metadata_fingerprint={"EAL": "fingerprint"},
            exported_at=datetime(2026, 7, 10, 12, tzinfo=timezone.utc),
        )),
        data_only=True,
    )

    assert workbook.sheetnames == ["Wear Records", "Cycle Coverage", "Conflict Audit", "Workbook Info"]
    sheet = workbook["Wear Records"]
    assert [cell.value for cell in sheet[1]] == [
        "Cycle Date", "Line", "Class", "Track", "Tension Length", "From (m)", "To (m)",
        "Interval Count", "Physical Intervals", "Avg Wear Min", "Wear %", "Measurement SD",
    ]
    assert [sheet.cell(row, 6).value for row in (2, 3)] == [100, 201]
    assert sheet.cell(2, 8).value == 1
    assert sheet.cell(2, 10).value == pytest.approx(11.987654321)
    assert sheet.cell(3, 11).value == pytest.approx(12.3456789)
    assert sheet.max_row == 3

    coverage = workbook["Cycle Coverage"]
    assert coverage.max_row == 12
    audit = list(workbook["Conflict Audit"].iter_rows(min_row=2, values_only=True))
    assert [(row[3], row[4]) for row in audit] == [
        ("EAL_U1-a.xlsx", 11.99), ("EAL_U1-b.xlsx", 11.98),
    ]
    info = {row[0]: row[1] for row in workbook["Workbook Info"].iter_rows(min_row=2, values_only=True)}
    assert info["Schema"] == "wear-cycle-report-v1"
    assert info["App Version"] == "2.0.0"
    assert json.loads(info["Metadata Fingerprint"]) == {"EAL": "fingerprint"}


def test_excel_report_filters_legacy_siding_records_and_diagnostic_gaps(conn):
    from app.core.calculation.wear_cycle_io import build_excel_report

    _seed_cycle(conn)
    conn.execute("""
        INSERT INTO wire_wear_cycle_records (
            cycle_id, line_group, line_class, cycle_date, tension_length, track,
            from_m, to_m, avg_wear_min, wear_percentage, measurement_sd,
            physical_intervals, source_lineage
        )
        SELECT cycle_id, line_group, line_class, cycle_date, 'X1', 'UP',
            301, 400, 11.5, 8.0, NULL, '[]', '[]'
        FROM wire_wear_cycle_records
        WHERE tension_length = '28'
    """)
    conn.execute(
        """
        UPDATE wire_wear_cycle_segments
        SET diagnostic_gaps = '["29", "X1"]'
        WHERE segment_name = 'U1'
        """
    )

    workbook = load_workbook(
        BytesIO(build_excel_report(conn, line_group="EAL", cycle_date="2026-05-28")),
        data_only=True,
    )

    wear_rows = list(workbook["Wear Records"].iter_rows(min_row=2, values_only=True))
    assert {row[4] for row in wear_rows} == {"28", "29"}
    coverage_rows = list(
        workbook["Cycle Coverage"].iter_rows(min_row=2, values_only=True)
    )
    u1_row = next(row for row in coverage_rows if row[3] == "U1")
    assert u1_row[6] == "29"


def test_build_sync_package_has_exact_schema_and_normalized_tables(conn):
    from app.core.calculation.wear_cycle_io import build_sync_package

    _seed_cycle(conn)
    package = build_sync_package(
        conn,
        source_workstation="TOV640-01",
        metadata_fingerprint={"EAL": "eal-hash", "TML": "tml-hash"},
        package_id="fixed-id",
        exported_at=datetime(2026, 7, 10, 12, tzinfo=timezone.utc),
    )

    assert set(package) == {
        "schema", "package_id", "exported_at", "source_workstation",
        "metadata_fingerprint", "cycles", "records", "segments",
        "conflict_decisions", "tombstones",
    }
    assert package["schema"] == "wear-cycle-v1"
    assert package["package_id"] == "fixed-id"
    assert package["records"][0]["tension_length"] == "28"
    assert package["records"][0]["line_class"] == "EAL"
    assert package["records"][0]["physical_intervals"]
    assert package["records"][0]["interval_count"] == 1
    assert len(package["segments"]) == 11
    assert {
        (segment["line_group"], segment["line_class"], segment["cycle_date"])
        for segment in package["segments"]
    } == {("EAL", "EAL", "2026-05-28")}
    assert package["conflict_decisions"][0]["measurement_identity"] == "measurement-28"
    assert package["conflict_decisions"][0]["line_group"] == "EAL"
    assert package["conflict_decisions"][0]["line_class"] == "EAL"
    assert package["conflict_decisions"][0]["cycle_date"] == "2026-05-28"


def test_preview_rejects_duplicate_package_keys_and_unknown_tl(conn):
    from app.core.calculation.wear_cycle_io import preview_sync_import

    package = _package_for_record()
    package["records"].append(dict(package["records"][0]))
    duplicate = preview_sync_import(
        conn, package, _metadata(), metadata_fingerprint={"EAL": "same", "TML": "same"}
    )
    assert any(action.action == "error" and "duplicate" in action.detail.lower() for action in duplicate.actions)

    package = _package_for_record()
    package["records"][0]["tension_length"] = "UNKNOWN"
    unknown = preview_sync_import(
        conn, package, _metadata(), metadata_fingerprint={"EAL": "same", "TML": "same"}
    )
    assert any(action.action == "error" and "unknown" in action.detail.lower() for action in unknown.actions)


@pytest.mark.parametrize(
    ("collection", "row"),
    [
        ("cycles", {"line_group": "EAL", "line_class": "EAL", "cycle_date": "not-a-date"}),
        ("segments", {"line_group": "EAL", "line_class": "EAL", "cycle_date": "2026-05-28", "segment_name": "U1"}),
        ("conflict_decisions", {"line_group": "EAL", "line_class": "EAL", "cycle_date": "2026-05-28"}),
    ],
)
def test_preview_rejects_malformed_and_duplicate_cycle_audit_rows(conn, collection, row):
    from app.core.calculation.wear_cycle_io import SyncValidationError, preview_sync_import

    package = _package_for_record()
    package[collection] = [row, dict(row)]
    with pytest.raises(SyncValidationError):
        preview_sync_import(
            conn, package, _metadata(), metadata_fingerprint=package["metadata_fingerprint"]
        )


def test_metadata_fingerprint_difference_is_advisory_only_when_geometry_matches(conn):
    from app.core.calculation.wear_cycle_io import preview_sync_import

    package = _package_for_record()
    package["metadata_fingerprint"]["EAL"] = "remote"
    advisory = preview_sync_import(
        conn, package, _metadata(), metadata_fingerprint={"EAL": "local", "TML": "same"}
    )
    assert advisory.warnings
    assert advisory.actions[0].action == "create"

    blocked = preview_sync_import(
        conn, package, _metadata(to_m=250), metadata_fingerprint={"EAL": "local", "TML": "same"}
    )
    assert blocked.actions[0].action == "error"
    assert "metadata" in blocked.actions[0].detail.lower()


def test_sync_metadata_isolated_by_record_line_for_same_tension_length(conn):
    from app.core.calculation.wear_cycle_io import SyncValidationError, preview_sync_import

    package = _package_for_record()
    metadata_by_line = {
        "EAL": (
            MetadataInterval("28", "UP", Decimal("0"), Decimal("10"), "EAL UP"),
        ),
        "TML": (
            MetadataInterval("28", "UP", Decimal("100"), Decimal("110"), "TML UP"),
        ),
    }
    package["records"][0].update({"from_m": 0.0, "to_m": 10.0})

    preview = preview_sync_import(
        conn,
        package,
        metadata_by_line,
        metadata_fingerprint=package["metadata_fingerprint"],
    )

    assert preview.actions[0].action == "create"
    assert preview.actions[0].record["from_m"] == 0.0
    assert preview.actions[0].record["to_m"] == 10.0
    assert preview.actions[0].record["physical_intervals"] == [
        {"track": "UP", "from_m": 0.0, "to_m": 10.0},
    ]

    missing = preview_sync_import(
        conn,
        package,
        {"TML": metadata_by_line["TML"]},
        metadata_fingerprint=package["metadata_fingerprint"],
    )
    assert missing.actions[0].action == "error"
    assert "EAL" in missing.actions[0].detail
    assert "metadata" in missing.actions[0].detail.lower()

    invalid_package = _package_for_record()
    invalid_package["records"][0]["line_group"] = "INVALID"
    with pytest.raises(SyncValidationError, match="line group must be EAL or TML"):
        preview_sync_import(
            conn,
            invalid_package,
            metadata_by_line,
            metadata_fingerprint=invalid_package["metadata_fingerprint"],
        )


def test_sync_round_trip_isolates_eal_and_lmc_same_date_and_tension_length(conn):
    from app.core.calculation.wear_cycle_io import apply_sync_import, preview_sync_import

    package = _package_for_record()
    lmc_cycle = dict(package["cycles"][0])
    lmc_cycle["line_class"] = "LMC"
    package["cycles"].append(lmc_cycle)
    lmc_record = dict(package["records"][0])
    lmc_record.update({"line_class": "LMC", "avg_wear_min": 9.75})
    package["records"].append(lmc_record)
    metadata = {
        ("EAL", "EAL"): (
            MetadataInterval("28", "UP", Decimal("100"), Decimal("200"), "EAL UP"),
        ),
        ("EAL", "LMC"): (
            MetadataInterval("28", "UP", Decimal("100"), Decimal("200"), "LMC UP"),
        ),
    }

    preview = preview_sync_import(
        conn,
        package,
        metadata,
        metadata_fingerprint=package["metadata_fingerprint"],
    )
    assert [(action.key[1], action.action) for action in preview.actions] == [
        ("EAL", "create"),
        ("LMC", "create"),
    ]

    apply_sync_import(
        conn,
        source_package=preview.source_package,
        preview_digest=preview.preview_digest,
        expected_data_version=preview.expected_data_version,
        metadata=metadata,
        metadata_fingerprint=package["metadata_fingerprint"],
    )

    rows = conn.execute(
        """
        SELECT line_group, line_class, cycle_date, tension_length, avg_wear_min
        FROM wire_wear_cycle_records
        ORDER BY line_class
        """
    ).fetchall()
    parents = conn.execute(
        """
        SELECT line_group, line_class, cycle_date
        FROM wire_wear_cycles
        ORDER BY line_class
        """
    ).fetchall()
    assert [tuple(row) for row in rows] == [
        ("EAL", "EAL", "2026-05-28", "28", 10.5),
        ("EAL", "LMC", "2026-05-28", "28", 9.75),
    ]
    assert [tuple(row) for row in parents] == [
        ("EAL", "EAL", "2026-05-28"),
        ("EAL", "LMC", "2026-05-28"),
    ]


def test_legacy_single_interval_sync_package_is_idempotent_at_equal_timestamp(conn):
    from app.core.calculation.wear_cycle_io import apply_sync_import, preview_sync_import

    package = _package_for_record()
    assert "physical_intervals" not in package["records"][0]
    assert "interval_count" not in package["records"][0]

    first = preview_sync_import(
        conn, package, _metadata(), metadata_fingerprint=package["metadata_fingerprint"]
    )
    assert first.actions[0].action == "create"
    apply_sync_import(
        conn,
        source_package=first.source_package,
        preview_digest=first.preview_digest,
        expected_data_version=first.expected_data_version,
        metadata=_metadata(),
        metadata_fingerprint=package["metadata_fingerprint"],
    )

    second = preview_sync_import(
        conn, package, _metadata(), metadata_fingerprint=package["metadata_fingerprint"]
    )
    assert second.actions[0].action == "unchanged"


def test_metadata_fingerprint_difference_blocks_changed_physical_intervals_with_same_bounds(conn):
    from app.core.calculation.wear_cycle_io import preview_sync_import

    package = _package_for_record()
    package["metadata_fingerprint"]["EAL"] = "remote"
    package["records"][0]["physical_intervals"] = [
        {"track": "UP", "from_m": 100.0, "to_m": 140.0},
        {"track": "UP", "from_m": 160.0, "to_m": 200.0},
    ]
    package["records"][0]["interval_count"] = 2
    local_metadata = (
        MetadataInterval("28", "UP", Decimal("100"), Decimal("130"), "EAL UP"),
        MetadataInterval("28", "UP", Decimal("170"), Decimal("200"), "EAL UP"),
    )

    preview = preview_sync_import(
        conn,
        package,
        local_metadata,
        metadata_fingerprint={"EAL": "local", "TML": "same"},
    )

    assert preview.actions[0].action == "error"
    assert "metadata" in preview.actions[0].detail.lower()


def test_sync_apply_keeps_source_lineage_and_physical_intervals_in_their_own_columns(conn):
    from app.core.calculation.wear_cycle_io import apply_sync_import, preview_sync_import

    package = _package_for_record()
    metadata = (
        MetadataInterval("28", "UP", Decimal("100"), Decimal("140"), "EAL UP"),
        MetadataInterval("28", "UP", Decimal("160"), Decimal("200"), "EAL UP"),
    )
    preview = preview_sync_import(
        conn, package, metadata, metadata_fingerprint=package["metadata_fingerprint"]
    )

    apply_sync_import(
        conn,
        source_package=preview.source_package,
        preview_digest=preview.preview_digest,
        expected_data_version=preview.expected_data_version,
        metadata=metadata,
        metadata_fingerprint=package["metadata_fingerprint"],
    )

    row = conn.execute(
        "SELECT source_lineage, physical_intervals FROM wire_wear_cycle_records"
    ).fetchone()
    assert json.loads(row["source_lineage"]) == ["remote.xlsx"]
    assert json.loads(row["physical_intervals"]) == [
        {"track": "UP", "from_m": 100.0, "to_m": 140.0},
        {"track": "UP", "from_m": 160.0, "to_m": 200.0},
    ]


def test_equal_timestamp_tombstone_wins_and_newer_record_recreates(conn):
    from app.core.calculation.wear_cycle_io import apply_sync_import, preview_sync_import

    timestamp = "2026-07-10T12:00:00Z"
    package = _package_for_record(updated_at=timestamp)
    package["tombstones"] = [{
        "line_group": "EAL", "line_class": "EAL",
        "cycle_date": "2026-05-28", "tension_length": "28",
        "deleted_at": timestamp, "source_package_id": "package-1",
    }]
    preview = preview_sync_import(conn, package, _metadata(), metadata_fingerprint=package["metadata_fingerprint"])
    assert preview.actions[0].action == "delete"
    applied = apply_sync_import(
        conn, source_package=preview.source_package, preview_digest=preview.preview_digest,
        expected_data_version=preview.expected_data_version, metadata=_metadata(),
        metadata_fingerprint=package["metadata_fingerprint"],
    )
    assert applied["deleted"] == 1
    assert conn.execute("SELECT COUNT(*) FROM wire_wear_cycle_records").fetchone()[0] == 0

    recreated_package = _package_for_record(updated_at="2026-07-10T12:00:01Z")
    recreated = preview_sync_import(
        conn, recreated_package, _metadata(), metadata_fingerprint=recreated_package["metadata_fingerprint"]
    )
    assert recreated.actions[0].action == "create"
    apply_sync_import(
        conn, source_package=recreated.source_package, preview_digest=recreated.preview_digest,
        expected_data_version=recreated.expected_data_version, metadata=_metadata(),
        metadata_fingerprint=recreated_package["metadata_fingerprint"],
    )
    row = conn.execute("SELECT track, from_m, to_m, wear_percentage FROM wire_wear_cycle_records").fetchone()
    assert tuple(row[:3]) == ("UP", 100.0, 200.0)
    assert row[3] != 999.0


def test_apply_round_trips_cycle_coverage_and_conflict_audit(conn):
    from app.core.calculation.wear_cycle_io import apply_sync_import, preview_sync_import

    package = _package_for_record()
    package["segments"] = [{
        "line_group": "EAL", "line_class": "EAL",
        "cycle_date": "2026-05-28", "segment_name": "U1",
        "is_present": 1, "coverage_percentage": 100.0, "diagnostic_gaps": [],
        "source_file_names": ["remote.xlsx"], "acquisition_date_from": "2026-05-27",
        "acquisition_date_to": "2026-05-27",
    }]
    package["conflict_decisions"] = [{
        "line_group": "EAL", "line_class": "EAL", "cycle_date": "2026-05-28",
        "measurement_identity": "measurement-28",
        "source_values": [["remote-a.xlsx", 10.6], ["remote-b.xlsx", 10.5]],
        "selected_wear_min": 10.5, "accepted_at": "2026-07-10T12:00:00Z",
    }]
    preview = preview_sync_import(
        conn, package, _metadata(), metadata_fingerprint=package["metadata_fingerprint"]
    )
    apply_sync_import(
        conn, source_package=preview.source_package, preview_digest=preview.preview_digest,
        expected_data_version=preview.expected_data_version, metadata=_metadata(),
        metadata_fingerprint=package["metadata_fingerprint"],
    )

    segment = conn.execute(
        "SELECT segment_name, source_file_names FROM wire_wear_cycle_segments"
    ).fetchone()
    decision = conn.execute(
        "SELECT measurement_identity, selected_wear_min FROM wire_wear_conflict_decisions"
    ).fetchone()
    assert tuple(segment) == ("U1", '["remote.xlsx"]')
    assert tuple(decision) == ("measurement-28", 10.5)


def test_apply_rejects_stale_version_changed_digest_and_rolls_back_atomically(conn, monkeypatch):
    import app.core.calculation.wear_cycle_io as sync

    package = _package_for_record()
    preview = sync.preview_sync_import(
        conn, package, _metadata(), metadata_fingerprint=package["metadata_fingerprint"]
    )
    with pytest.raises(sync.SyncPreviewMismatchError):
        sync.apply_sync_import(
            conn, source_package=preview.source_package, preview_digest="changed",
            expected_data_version=preview.expected_data_version, metadata=_metadata(),
            metadata_fingerprint=package["metadata_fingerprint"],
        )

    conn.execute("INSERT OR IGNORE INTO system_metadata (key, value) VALUES ('wire_wear_data_version', '0')")
    conn.execute("UPDATE system_metadata SET value = '1' WHERE key = 'wire_wear_data_version'")
    conn.commit()
    with pytest.raises(StaleDataVersionError):
        sync.apply_sync_import(
            conn, source_package=preview.source_package, preview_digest=preview.preview_digest,
            expected_data_version=preview.expected_data_version, metadata=_metadata(),
            metadata_fingerprint=package["metadata_fingerprint"],
        )

    conn.execute("UPDATE system_metadata SET value = '0' WHERE key = 'wire_wear_data_version'")
    conn.commit()
    original_execute = sync._apply_action
    monkeypatch.setattr(sync, "_apply_action", lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("boom")))
    with pytest.raises(RuntimeError, match="boom"):
        sync.apply_sync_import(
            conn, source_package=preview.source_package, preview_digest=preview.preview_digest,
            expected_data_version=0, metadata=_metadata(),
            metadata_fingerprint=package["metadata_fingerprint"],
        )
    monkeypatch.setattr(sync, "_apply_action", original_execute)
    assert conn.execute("SELECT COUNT(*) FROM wire_wear_cycle_records").fetchone()[0] == 0
    assert current_data_version(conn) == 0


def test_apply_rolls_back_first_action_when_second_action_fails(conn, monkeypatch):
    import app.core.calculation.wear_cycle_io as sync

    package = _package_for_record()
    second = dict(package["records"][0])
    second.update({"tension_length": "29", "from_m": 201.0, "to_m": 300.0})
    package["records"].append(second)
    preview = sync.preview_sync_import(
        conn, package, _metadata(), metadata_fingerprint=package["metadata_fingerprint"]
    )
    original_apply = sync._apply_action
    calls = 0

    def fail_second(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError("second action failed")
        return original_apply(*args, **kwargs)

    monkeypatch.setattr(sync, "_apply_action", fail_second)
    with pytest.raises(RuntimeError, match="second action failed"):
        sync.apply_sync_import(
            conn, source_package=preview.source_package, preview_digest=preview.preview_digest,
            expected_data_version=preview.expected_data_version, metadata=_metadata(),
            metadata_fingerprint=package["metadata_fingerprint"],
        )
    assert conn.execute("SELECT COUNT(*) FROM wire_wear_cycle_records").fetchone()[0] == 0
    assert current_data_version(conn) == 0


def test_preview_exposes_normalized_issue_10_status_matrix(conn):
    from app.core.calculation.wear_cycle_io import apply_sync_import, preview_sync_import

    base = _package_for_record()
    initial = preview_sync_import(
        conn, base, _metadata(), metadata_fingerprint=base["metadata_fingerprint"]
    )
    apply_sync_import(
        conn,
        source_package=initial.source_package,
        preview_digest=initial.preview_digest,
        expected_data_version=initial.expected_data_version,
        metadata=_metadata(),
        metadata_fingerprint=base["metadata_fingerprint"],
    )

    new_package = _package_for_record()
    new_package["records"][0].update({
        "tension_length": "29",
        "from_m": 201.0,
        "to_m": 300.0,
    })
    newer = _package_for_record(updated_at="2026-07-10T12:00:01Z", avg=10.4)
    older = _package_for_record(updated_at="2026-07-10T11:59:59Z", avg=10.6)
    conflict = _package_for_record(avg=10.4)
    invalid = _package_for_record()
    invalid["records"][0]["tension_length"] = "unknown"
    deleted = _package_for_record()
    deleted["tombstones"] = [{
        "line_group": "EAL",
        "line_class": "EAL",
        "cycle_date": "2026-05-28",
        "tension_length": "28",
        "deleted_at": "2026-07-10T12:00:01Z",
        "source_package_id": "remote-delete",
    }]

    cases = (
        (new_package, "new"),
        (newer, "update"),
        (older, "keep_local"),
        (base, "no_change"),
        (conflict, "conflict"),
        (invalid, "error"),
        (deleted, "delete"),
    )
    statuses = []
    for package, expected in cases:
        preview = preview_sync_import(
            conn,
            package,
            _metadata(),
            metadata_fingerprint=package["metadata_fingerprint"],
        )
        statuses.append(preview.actions[0].status)
        assert preview.actions[0].status == expected

    assert statuses == [
        "new", "update", "keep_local", "no_change", "conflict", "error", "delete",
    ]


@pytest.mark.parametrize(
    ("decision", "expected_average", "expected_updated", "expected_keep_local"),
    [
        ("incoming", 9.75, 1, 0),
        ("local", 10.5, 0, 1),
    ],
)
def test_conflict_requires_explicit_decision_and_applies_selected_side(
    conn,
    decision,
    expected_average,
    expected_updated,
    expected_keep_local,
):
    from app.core.calculation.wear_cycle_io import (
        SyncValidationError,
        apply_sync_import,
        preview_sync_import,
    )

    base = _package_for_record()
    initial = preview_sync_import(
        conn, base, _metadata(), metadata_fingerprint=base["metadata_fingerprint"]
    )
    apply_sync_import(
        conn,
        source_package=initial.source_package,
        preview_digest=initial.preview_digest,
        expected_data_version=initial.expected_data_version,
        metadata=_metadata(),
        metadata_fingerprint=base["metadata_fingerprint"],
    )

    incoming = _package_for_record(avg=9.75)
    preview = preview_sync_import(
        conn, incoming, _metadata(), metadata_fingerprint=incoming["metadata_fingerprint"]
    )
    assert preview.actions[0].status == "conflict"
    with pytest.raises(SyncValidationError, match="unresolved conflicts"):
        apply_sync_import(
            conn,
            source_package=preview.source_package,
            preview_digest=preview.preview_digest,
            expected_data_version=preview.expected_data_version,
            metadata=_metadata(),
            metadata_fingerprint=incoming["metadata_fingerprint"],
        )

    result = apply_sync_import(
        conn,
        source_package=preview.source_package,
        preview_digest=preview.preview_digest,
        expected_data_version=preview.expected_data_version,
        metadata=_metadata(),
        metadata_fingerprint=incoming["metadata_fingerprint"],
        conflict_decisions=[{
            "key": {
                "line_group": "EAL",
                "line_class": "EAL",
                "cycle_date": "2026-05-28",
                "tension_length": "28",
            },
            "resolution": decision,
        }],
    )

    average = conn.execute(
        "SELECT avg_wear_min FROM wire_wear_cycle_records"
    ).fetchone()[0]
    assert average == expected_average
    assert result["updated"] == expected_updated
    assert result["keep_local"] == expected_keep_local
    assert result["conflicts_resolved"] == 1
    assert bool(result["backup_path"]) is (decision == "incoming")


def test_conflict_rejects_unknown_and_invalid_decisions(conn):
    from app.core.calculation.wear_cycle_io import (
        SyncValidationError,
        apply_sync_import,
        preview_sync_import,
    )

    base = _package_for_record()
    initial = preview_sync_import(
        conn, base, _metadata(), metadata_fingerprint=base["metadata_fingerprint"]
    )
    apply_sync_import(
        conn,
        source_package=initial.source_package,
        preview_digest=initial.preview_digest,
        expected_data_version=initial.expected_data_version,
        metadata=_metadata(),
        metadata_fingerprint=base["metadata_fingerprint"],
    )
    incoming = _package_for_record(avg=9.75)
    preview = preview_sync_import(
        conn, incoming, _metadata(), metadata_fingerprint=incoming["metadata_fingerprint"]
    )

    with pytest.raises(SyncValidationError, match="does not match"):
        apply_sync_import(
            conn,
            source_package=preview.source_package,
            preview_digest=preview.preview_digest,
            expected_data_version=preview.expected_data_version,
            metadata=_metadata(),
            metadata_fingerprint=incoming["metadata_fingerprint"],
            conflict_decisions={"EAL|EAL|2026-05-28|29": "local"},
        )
    with pytest.raises(SyncValidationError, match="incoming or local"):
        apply_sync_import(
            conn,
            source_package=preview.source_package,
            preview_digest=preview.preview_digest,
            expected_data_version=preview.expected_data_version,
            metadata=_metadata(),
            metadata_fingerprint=incoming["metadata_fingerprint"],
            conflict_decisions={"EAL|EAL|2026-05-28|28": "merge"},
        )


def test_backup_failure_blocks_sync_without_mutation(conn, monkeypatch):
    import app.core.calculation.wear_cycle_io as sync
    from app.core.calculation.wear_cycle_application import BackupCreationError

    package = _package_for_record()
    preview = sync.preview_sync_import(
        conn, package, _metadata(), metadata_fingerprint=package["metadata_fingerprint"]
    )

    def fail_backup(*_args, **_kwargs):
        raise BackupCreationError("backup unavailable")

    monkeypatch.setattr(sync, "create_sqlite_backup", fail_backup)
    with pytest.raises(BackupCreationError, match="backup unavailable"):
        sync.apply_sync_import(
            conn,
            source_package=preview.source_package,
            preview_digest=preview.preview_digest,
            expected_data_version=preview.expected_data_version,
            metadata=_metadata(),
            metadata_fingerprint=package["metadata_fingerprint"],
        )

    assert conn.execute("SELECT COUNT(*) FROM wire_wear_cycle_records").fetchone()[0] == 0
    assert current_data_version(conn) == 0


def test_successful_sync_returns_readable_pre_apply_backup(conn):
    from app.core.calculation.wear_cycle_io import apply_sync_import, preview_sync_import

    package = _package_for_record()
    preview = preview_sync_import(
        conn, package, _metadata(), metadata_fingerprint=package["metadata_fingerprint"]
    )
    result = apply_sync_import(
        conn,
        source_package=preview.source_package,
        preview_digest=preview.preview_digest,
        expected_data_version=preview.expected_data_version,
        metadata=_metadata(),
        metadata_fingerprint=package["metadata_fingerprint"],
    )

    backup_path = Path(result["backup_path"])
    assert result == {
        "created": 1,
        "updated": 0,
        "deleted": 0,
        "keep_local": 0,
        "no_change": 0,
        "unchanged": 0,
        "conflicts_resolved": 0,
        "backup_path": str(backup_path),
        "data_version": 1,
    }
    assert backup_path.exists()
    with closing(sqlite3.connect(backup_path)) as backup_conn:
        assert backup_conn.execute(
            "SELECT COUNT(*) FROM wire_wear_cycle_records"
        ).fetchone()[0] == 0
        assert current_data_version(backup_conn) == 0


def test_two_database_sync_round_trip_preserves_normalized_state(tmp_path):
    from app.core.calculation.wear_cycle_io import (
        apply_sync_import,
        build_sync_package,
        preview_sync_import,
    )

    source_path = tmp_path / "source" / "wear.db"
    target_path = tmp_path / "target" / "wear.db"
    source_manager = DatabaseManager(str(source_path))
    try:
        with source_manager.get_connection() as source_conn:
            package = _package_for_record()
            preview = preview_sync_import(
                source_conn,
                package,
                _metadata(),
                metadata_fingerprint=package["metadata_fingerprint"],
            )
            apply_sync_import(
                source_conn,
                source_package=preview.source_package,
                preview_digest=preview.preview_digest,
                expected_data_version=preview.expected_data_version,
                metadata=_metadata(),
                metadata_fingerprint=package["metadata_fingerprint"],
            )
            exported = build_sync_package(
                source_conn,
                source_workstation="source-machine",
                package_id="round-trip",
                metadata_fingerprint=package["metadata_fingerprint"],
            )
    finally:
        source_manager.close()
        DatabaseManager.reset_instance()

    target_manager = DatabaseManager(str(target_path))
    try:
        with target_manager.get_connection() as target_conn:
            preview = preview_sync_import(
                target_conn,
                exported,
                _metadata(),
                metadata_fingerprint=exported["metadata_fingerprint"],
            )
            assert [(action.status, action.key) for action in preview.actions] == [
                ("new", ("EAL", "EAL", "2026-05-28", "28")),
            ]
            result = apply_sync_import(
                target_conn,
                source_package=preview.source_package,
                preview_digest=preview.preview_digest,
                expected_data_version=preview.expected_data_version,
                metadata=_metadata(),
                metadata_fingerprint=exported["metadata_fingerprint"],
            )
            target_export = build_sync_package(
                target_conn,
                source_workstation="target-machine",
                package_id="target-copy",
                metadata_fingerprint=exported["metadata_fingerprint"],
            )
    finally:
        target_manager.close()
        DatabaseManager.reset_instance()

    assert result["created"] == 1
    assert Path(result["backup_path"]).parent == target_path.parent / "backups"
    for collection in ("cycles", "records", "segments", "conflict_decisions", "tombstones"):
        assert target_export[collection] == exported[collection]
