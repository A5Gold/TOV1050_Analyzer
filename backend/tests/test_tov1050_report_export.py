import pytest
from openpyxl import load_workbook
import pandas as pd
from app.core.exporter import ExcelExporter
from app.core.tov1050_report_export import export_tov1050_report,export_analysis_report


def test_streaming_report_preserves_every_sheet_cell_order_and_metadata():
    result={"Wire Wear":pd.DataFrame([{"id":"ID0","FromM":98150.2,"ToM":98151,"maxValue":9.1,"level":"L1","Tension Length":"A, B"}]),
            "Low Height":pd.DataFrame()}
    frame=pd.DataFrame({"Chainage":[98150.2,98150.3,98150.3,98150.4],"height1":[4400,None,4550,4600],
                        "wear1":[9.1,10.2,float('nan'),11],"stagger1":[-350,0,500,500],
                        "line":["ISL"]*4,"Overlap":[None,"Y",None,None],"source_row_number":[102,103,104,105]})
    metadata={"enrichment_version":"tov640-first-nonempty-v1","metadata_capabilities":"unknown"}
    legacy=load_workbook(ExcelExporter.export_report(result,frame,metadata),read_only=True)
    streamed=load_workbook(export_tov1050_report(result,frame,metadata),read_only=True)
    assert legacy.sheetnames==streamed.sheetnames
    for name in legacy.sheetnames:
        # Write-only XML omits trailing blank cells; compare the full header width.
        assert list(legacy[name].values)==list(streamed[name].iter_rows(values_only=True,max_col=legacy[name].max_column))
        assert legacy[name].sheet_state==streamed[name].sheet_state
    legacy.close()
    streamed.close()


def test_large_tov1050_chartdata_streams_full_rows_without_downsampling():
    frame=pd.DataFrame({"Chainage":range(50001),"height1":[4500.0]*50001})
    workbook=load_workbook(export_analysis_report({},frame,line="AEL"),read_only=True)
    rows=workbook["ChartData"].values
    assert next(rows)==("Chainage","height1")
    count=0
    for count,row in enumerate(rows,1):
        assert row==(count-1,4500)
    assert count==50001
    workbook.close()


def test_eal_large_report_keeps_legacy_writer(monkeypatch):
    from app.core import tov1050_report_export as module
    sentinel=object()
    monkeypatch.setattr(module.ExcelExporter,"export_report",lambda *args,**kwargs:sentinel)
    monkeypatch.setattr(module,"export_tov1050_report",lambda *args,**kwargs: (_ for _ in ()).throw(AssertionError("EAL must keep legacy writer")))
    assert module.export_analysis_report({},pd.DataFrame({"Chainage":range(50001)}),line="EAL") is sentinel


def test_tov1050_download_chunks_preserve_binary_and_csv_bytes():
    import io
    from app.api.endpoints.analysis import _download_body
    for content in (bytes(range(256))*1025, b"Chainage,height1\n"+b"98150.2,4500\n"*20000):
        chunks=list(_download_body(io.BytesIO(content),"KTL"))
        assert b"".join(chunks)==content
        assert max(map(len,chunks))==65536
        assert len(chunks)==(len(content)+65535)//65536
    for line in ("EAL","TML"):
        stream=io.BytesIO(b"legacy\n")
        assert _download_body(stream,line) is stream


@pytest.mark.parametrize("streamed", [False, True])
def test_report_four_decimal_display_preserves_full_measurement_values(streamed):
    value = 48430.1234567
    frame = pd.DataFrame({"Chainage": [value], "height1": [3500.1234567]})
    results = {"Low Height": pd.DataFrame([{"FromM": value, "maxValue": 3500.1234567}])}
    output = export_tov1050_report(results, frame) if streamed else export_analysis_report(results, frame, line="AEL")
    book = load_workbook(output, read_only=True)
    for sheet_name, header, expected in (("ChartData", "Chainage", value), ("Summary", "FromM", value),
                                         ("Low Height", "MaxValue", 3500.1234567)):
        rows = list(book[sheet_name].rows)
        cell = rows[1][[c.value for c in rows[0]].index(header)]
        assert cell.number_format == "0.####"
        assert cell.value == expected
    book.close()


def test_path_writer_progress_and_compact_statistics_preserve_raw_report(tmp_path):
    from app.core.chainage_aggregation import aggregate_chainage
    raw = pd.DataFrame({"Chainage": [100.125] * 2, "height1": [4000.1234567, 5000.1234567]})
    mean, metadata = aggregate_chainage(raw)
    events = []
    destination = tmp_path / "compact_ChainageMean.xlsx"
    output = export_analysis_report({}, mean, metadata, line="TWL", chart_statistics_df=mean,
                                    output_path=destination, progress=lambda *event: events.append(event))
    assert output == destination
    book = load_workbook(output, read_only=True)
    assert book.sheetnames == ["Summary", "ChartData", "ChainageStatistics", "_Metadata"]
    assert list(book["ChartData"].values)[1] == (100.125, 4500.1234567, 4500.1234567, 4500.1234567)
    stats = list(book["ChainageStatistics"].values)
    assert stats[1][stats[0].index("height1_max")] == 5000.1234567
    assert book["_Metadata"].sheet_state == "hidden"
    book.close()
    assert ("exporting_ChartData", 1, 1, "rows") in events
    assert events[-1] == ("finalizing", None, None, None)
    full = load_workbook(export_analysis_report({}, raw, line="TWL"), read_only=True)
    assert [row[1] for row in list(full["ChartData"].values)[1:]] == raw.height1.tolist()
    full.close()


def test_writer_cancellation_closes_temporary_worksheets(tmp_path):
    from openpyxl.worksheet._writer import ALL_TEMP_FILES
    previous = set(ALL_TEMP_FILES)
    destination = tmp_path / "cancelled.xlsx"
    def cancel(stage, processed=None, total=None, unit=None):
        if stage == "exporting_ChartData" and processed == 1000:
            raise RuntimeError("cancelled")
    with pytest.raises(RuntimeError, match="cancelled"):
        export_tov1050_report({}, pd.DataFrame({"Chainage": range(1500)}),
                              output_path=destination, progress=cancel)
    assert set(ALL_TEMP_FILES) == previous
    assert not destination.exists()


def test_streaming_writer_does_not_mutate_published_detector_results():
    frame = pd.DataFrame([{"FromM": 1.25, "maxValue": 3.5}])
    original = frame.copy(deep=True)
    export_tov1050_report({"Wire Wear": frame})
    pd.testing.assert_frame_equal(frame, original)


def test_compact_statistics_export_native_counter_alias_ranges_and_metadata():
    from app.core.chainage_aggregation import aggregate_chainage
    raw = pd.DataFrame({"Chainage": [1., 1., 1.], "height1": [2., 4., 20.],
                        "SampleCount [n.]": [100, 101, 1]})
    mean, metadata = aggregate_chainage(raw)
    book = load_workbook(export_tov1050_report({}, mean, metadata, chart_statistics_df=mean), read_only=True)
    try:
        rows = list(book["ChainageStatistics"].values)
        first = rows[0].index("SampleCount_first")
        last = rows[0].index("SampleCount_last")
        assert [(row[first], row[last]) for row in rows[1:]] == [(100, 101), (1, 1)]
        exported_metadata = dict(list(book["_Metadata"].values)[1:])
        assert exported_metadata["sample_counter_column"] == "SampleCount [n.]"
    finally:
        book.close()


@pytest.mark.parametrize("streamed", [False, True])
def test_integer_display_has_no_decimal_separator_and_keeps_numeric_values(streamed):
    values = [0, 9142, -200, 9390.4, -226.775, 0.1234567]
    frame = pd.DataFrame({"Chainage": values, "height1": values})
    results = {"Low Height": pd.DataFrame({
        "id": [f"ID{i}" for i in range(len(values))],
        "FromM": values, "ToM": values, "Threshold Value": values,
        "maxValue": values,
    })}
    output = export_tov1050_report(results, frame) if streamed else export_analysis_report(
        results, frame, line="TWL"
    )
    book = load_workbook(output, read_only=True)
    try:
        for title, names in (("Summary", ["FromM", "ToM", "Threshold Value", "MaxValue"]),
                             ("Low Height", ["FromM", "ToM", "Threshold Value", "MaxValue"]),
                             ("ChartData", ["Chainage", "height1"])):
            rows = list(book[title].rows)
            headers = [cell.value for cell in rows[0]]
            for name in names:
                cells = [row[headers.index(name)] for row in rows[1:]]
                for cell, value in zip(cells, values):
                    assert cell.data_type == "n"
                    assert cell.value == value
                    assert cell.number_format == ("0" if float(value).is_integer() else "0.####")
    finally:
        book.close()


@pytest.mark.parametrize("streamed", [False, True])
def test_median_report_preserves_channel_and_exception_source_lineage(streamed):
    frame = pd.DataFrame({
        "Chainage": [9142.0], "height1": [4401.25], "wear1": [10.125],
        "group_id": [1], "pass_id": [0], "segment_id": [2], "sample_count": [10],
        "source_row_first": [12], "source_row_last": [21],
        "height1_source_row": [14], "height1_sample_counter": [103], "height1_valid_count": [10],
        "wear1_source_row": [18], "wear1_sample_counter": [107], "wear1_valid_count": [9],
    })
    result = {"Wire Wear": pd.DataFrame([{
        "id": "WW0", "FromM": 9142.0, "ToM": 9142.0, "maxValue": 10.125,
        "maxLocation": 9142.0, "peak_channel": "wear1", "peak_source_row": 18,
        "peak_sample_counter": 107,
    }])}
    original = result["Wire Wear"].copy(deep=True)
    metadata = {
        "algorithm": "same_chainage_median_sample_v1", "source_sha256": "a" * 64,
        "metadata_sha256": "b" * 64, "raw_points": 10, "grouped_points": 1,
        "even_median_rule": "earliest source sample among middle two values",
    }
    output = export_tov1050_report(result, frame, metadata) if streamed else export_analysis_report(
        result, frame, metadata, line="TWL"
    )
    book = load_workbook(output, read_only=True)
    try:
        assert "ChainageStatistics" not in book.sheetnames
        chart = list(book["ChartData"].values)
        exported = dict(zip(chart[0], chart[1]))
        assert exported == frame.iloc[0].to_dict()
        provenance = list(book["ExceptionProvenance"].values)
        peak = dict(zip(provenance[0], provenance[1]))
        assert peak["id"] == "WW0"
        assert peak["exception type"] == "Wire Wear"
        assert peak["maxValue"] == 10.125
        assert peak["peak_channel"] == "wear1"
        assert peak["peak_source_row"] == exported["wear1_source_row"]
        assert peak["peak_sample_counter"] == exported["wear1_sample_counter"]
        assert dict(list(book["_Metadata"].values)[1:]) == {key: str(value) for key, value in metadata.items()}
        assert book["_Metadata"].sheet_state == "hidden"
        summary = next(book["Summary"].values)
        assert summary == tuple(ExcelExporter.REPORT_HEADERS[column] for column in ExcelExporter.REPORT_COLUMNS)
    finally:
        book.close()
    pd.testing.assert_frame_equal(result["Wire Wear"], original)
