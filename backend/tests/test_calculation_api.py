"""E2E tests for calculation API endpoints using FastAPI TestClient."""
import io
import sys
import os
import zipfile
from dataclasses import replace
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from app.main import app
from app.core.calculation.excel_parser import ChartDataRecord
from app.core.calculation.excel_parser import parse_exception_report
from app.core.calculation.wear_cycle_aggregation import preview_digest
from app.core.calculation.wear_cycle_types import (
    AggregatedWearRecord,
    BusinessKey,
    CyclePreview,
    MetadataInterval,
    PhysicalInterval,
    SegmentCoverage,
)
from app.core.calculation.wear_tl_scope import (
    TensionLengthScope,
    classify_tension_length_scope,
)

client = TestClient(app)


def _complete_cycle_preview(cycle_date="2026-05-28"):
    record = AggregatedWearRecord(
        key=BusinessKey("EAL", date.fromisoformat(cycle_date), "28"),
        track="UP",
        from_m=Decimal("100"),
        to_m=Decimal("200"),
        avg_wear_min=11.4,
        wear_percentage=10.0,
        measurement_sd=None,
        has_data_conflict=False,
        conflict_ids=(),
        source_lineage=("EAL_U1.xlsx",),
        intervals=(
            PhysicalInterval("UP", Decimal("100"), Decimal("140")),
            PhysicalInterval("UP", Decimal("160"), Decimal("200")),
        ),
    )
    segments = tuple(
        SegmentCoverage(name, True, 100.0, (), (f"EAL_{name}.xlsx",), (date(2026, 5, 28),))
        for name in ("U1", "U2", "U3", "D1", "D2", "D3", "LOW S1", "RAC UP", "RAC DN", "LMC UP", "LMC DN")
    )
    return CyclePreview(
        "EAL", date.fromisoformat(cycle_date), (record,), segments, (), (), (), True,
        datetime(2026, 5, 28, tzinfo=timezone.utc),
    )


def _required_workbook_bytes() -> bytes:
    workbook = Workbook()
    workbook.active.title = "Wire Wear"
    workbook.create_sheet("ChartData")
    content = io.BytesIO()
    workbook.save(content)
    return content.getvalue()


def test_upload_wear_returns_complete_cycle_preview(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    preview = _complete_cycle_preview()
    monkeypatch.setattr(
        calculation_endpoint,
        "_build_complete_cycle_preview",
        lambda **_kwargs: preview,
        raising=False,
    )

    response = client.post(
        "/api/calculation/wear",
        files=[("files", ("EAL_U1.xlsx", b"source", "application/octet-stream"))],
        data={"line_group": "EAL", "cycle_date": "2026-05-28", "accepted_conflict_ids": "[]"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["cycle_date"] == "2026-05-28"
    assert set(body) >= {
        "records", "segments", "conflicts", "blocking_reasons", "can_save",
        "preview_digest", "expected_data_version",
    }
    assert body["preview_digest"] == preview_digest(preview)
    assert all(row["line_group"] == "EAL" for row in body["records"])
    assert body["records"][0]["interval_count"] == 2
    assert body["records"][0]["intervals"] == [
        {"track": "UP", "from_m": 100.0, "to_m": 140.0},
        {"track": "UP", "from_m": 160.0, "to_m": 200.0},
    ]


def test_upload_wear_rejects_invalid_complete_cycle_date(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    monkeypatch.setattr(
        calculation_endpoint,
        "_build_complete_cycle_preview",
        lambda **_kwargs: replace(
            _complete_cycle_preview(),
            cycle_date=date.min,
            blocking_reasons=("cycle_date_invalid",),
            can_save=False,
        ),
        raising=False,
    )
    response = client.post(
        "/api/calculation/wear",
        files=[("files", ("EAL_U1.xlsx", b"source", "application/octet-stream"))],
        data={"line_group": "EAL", "cycle_date": "not-a-date", "accepted_conflict_ids": "[]"},
    )

    assert response.status_code == 422


def test_upload_wear_rejects_unresolved_complete_cycle_source(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    monkeypatch.setattr(
        calculation_endpoint,
        "_build_complete_cycle_preview",
        lambda **_kwargs: replace(
            _complete_cycle_preview(),
            unresolved=("unknown tension length X99 for EAL",),
            blocking_reasons=("unresolved_tension_length",),
            can_save=False,
        ),
    )
    response = client.post(
        "/api/calculation/wear",
        files=[("files", ("EAL_U1.xlsx", b"source", "application/octet-stream"))],
        data={"line_group": "EAL", "cycle_date": "2026-05-28", "accepted_conflict_ids": "[]"},
    )

    assert response.status_code == 422
    assert response.json()["detail"]["blocking_reasons"] == ["unresolved_tension_length"]


def test_complete_cycle_source_skips_blank_tl_but_preserves_source_range(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    rows = [
        ChartDataRecord(
            task_run_date="2026-01-28",
            line="EAL",
            track="DN",
            section="LMC",
            task_no="",
            station_start="",
            station_end="",
            tension_length="10",
            chainage=129985.25,
            wear_min=11.4,
            track_type="Straight",
            overlap=None,
        ),
        ChartDataRecord(
            task_run_date="2026-01-28",
            line="EAL",
            track="DN",
            section="LMC",
            task_no="",
            station_start="",
            station_end="",
            tension_length="",
            chainage=130125.75,
            wear_min=11.21,
            track_type="Curve",
            overlap=None,
        ),
    ]
    monkeypatch.setattr(
        calculation_endpoint,
        "parse_exception_report",
        lambda _content: ([], rows),
    )

    source = calculation_endpoint._parse_complete_cycle_source(
        "20260128_EAL_DN_LMC_Exception_Report.xlsx",
        b"source",
        "EAL",
        metadata=(
            MetadataInterval(
                "10", "DN", Decimal("129900"), Decimal("130000"), "LMC DN", "LMC"
            ),
        ),
    )

    assert [measurement.tension_length for measurement in source.measurements] == ["10"]
    assert source.from_m == Decimal("129985.25")
    assert source.to_m == Decimal("130125.75")


def test_complete_cycle_source_preserves_combined_segment_membership_without_copying_rows(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    rows = [
        ChartDataRecord(
            task_run_date="2026-05-11",
            line="TML",
            track="DN",
            section="Mainline",
            task_no="D5-D3",
            station_start="TUM",
            station_end="HUH",
            tension_length="67",
            chainage=134900.0,
            wear_min=11.4,
            track_type="Tangent",
            overlap=None,
        ),
    ]
    monkeypatch.setattr(calculation_endpoint, "parse_exception_report", lambda _content: ([], rows))

    source = calculation_endpoint._parse_complete_cycle_source(
        "20260511_TML_D5-D3_TUM-HUH_Exception_Report.xlsx",
        b"source",
        "TML",
        metadata=(
            MetadataInterval("67", "DN", Decimal("134800"), Decimal("135000"), "TML DN"),
        ),
    )

    assert source.segment_name == "D3,D4,D5"
    assert len(source.measurements) == 1


def test_complete_cycle_source_maps_compound_raw_tl_by_unique_chainage(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    rows = [ChartDataRecord(
        task_run_date="2026-05-11", line="TML", track="DN", section="Mainline",
        task_no="D1", station_start="TAW", station_end="WKS",
        tension_length="M07,M05", chainage=115.0, wear_min=11.4,
        track_type="Tangent", overlap="Y",
    )]
    monkeypatch.setattr(calculation_endpoint, "parse_exception_report", lambda _content: ([], rows))
    metadata = (
        MetadataInterval("M05", "DN", Decimal("100"), Decimal("120"), "TML DN"),
        MetadataInterval("M07", "DN", Decimal("120.1"), Decimal("140"), "TML DN"),
    )

    source = calculation_endpoint._parse_complete_cycle_source(
        "20260511_TML_D1_TAW-WKS_Exception_Report.xlsx", b"source", "TML", metadata=metadata,
    )
    assert source.measurements[0].tension_length == "M05"


@pytest.mark.parametrize(
    "filename, line_group, track, section, task_no, raw_tl, chainage, metadata, expected",
    [
        (
            "20260617_EAL_D3_LOW-TAP_Exception_Report.xlsx",
            "EAL",
            "DN",
            "Mainline",
            "D3",
            "70, X32, L02",
            130170.25,
            (
                MetadataInterval(
                    "70", "DN", Decimal("130100"), Decimal("130200"), "EAL DN", "Mainline", 0
                ),
                MetadataInterval(
                    "X32", "DN", Decimal("130100"), Decimal("130200"), "EAL DN", "Mainline", 1
                ),
                MetadataInterval(
                    "L02", "DN", Decimal("130100"), Decimal("130200"), "EAL DN", "Mainline", 2
                ),
                MetadataInterval(
                    "L02", "DN", Decimal("130100"), Decimal("130200"), "LMC DN", "LMC", 0
                ),
            ),
            "70",
        ),
        (
            "20260617_EAL_DN_LMC_Exception_Report.xlsx",
            "EAL",
            "DN",
            "LMC",
            "LMC",
            "L02, 70, X38",
            130170.25,
            (
                MetadataInterval(
                    "L02", "DN", Decimal("130100"), Decimal("130200"), "LMC DN", "LMC", 0
                ),
                MetadataInterval(
                    "70", "DN", Decimal("130100"), Decimal("130200"), "LMC DN", "LMC", 1
                ),
                MetadataInterval(
                    "X38", "DN", Decimal("130100"), Decimal("130200"), "LMC DN", "LMC", 2
                ),
            ),
            "L02",
        ),
        (
            "20260511_TML_D5-D3_TUM-HUH_Exception_Report.xlsx",
            "TML",
            "DN",
            "Mainline",
            "D5-D3",
            "48,50",
            126335.25,
            (
                MetadataInterval(
                    "48", "DN", Decimal("126300"), Decimal("126400"), "TML DN", "Mainline", 0
                ),
                MetadataInterval(
                    "50", "DN", Decimal("126300"), Decimal("126400"), "TML DN", "Mainline", 1
                ),
            ),
            "48",
        ),
    ],
)
def test_complete_cycle_source_resolves_section_specific_primary_tl(
    monkeypatch,
    filename,
    line_group,
    track,
    section,
    task_no,
    raw_tl,
    chainage,
    metadata,
    expected,
):
    import app.api.endpoints.calculation as calculation_endpoint

    rows = [
        ChartDataRecord(
            task_run_date="2026-06-17",
            line=line_group,
            track=track,
            section=section,
            task_no=task_no,
            station_start="START",
            station_end="END",
            tension_length=raw_tl,
            chainage=chainage,
            wear_min=11.4,
            track_type="Tangent",
            overlap="Y",
        )
    ]
    monkeypatch.setattr(calculation_endpoint, "parse_exception_report", lambda _content: ([], rows))

    source = calculation_endpoint._parse_complete_cycle_source(
        filename, b"source", line_group, metadata=metadata
    )

    assert source.measurements[0].tension_length == expected


def test_complete_cycle_source_falls_back_to_detected_segment_when_section_is_blank(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    rows = [ChartDataRecord(
        task_run_date="2026-06-17", line="EAL", track="DN", section="",
        task_no="D3", station_start="START", station_end="END",
        tension_length="70,L02", chainage=130170.25, wear_min=11.4,
        track_type="Tangent", overlap="Y",
    )]
    monkeypatch.setattr(calculation_endpoint, "parse_exception_report", lambda _content: ([], rows))

    source = calculation_endpoint._parse_complete_cycle_source(
        "20260617_EAL_D3_Exception_Report.xlsx",
        b"source",
        "EAL",
        metadata=(
            MetadataInterval(
                "70", "DN", Decimal("130100"), Decimal("130200"), "EAL DN", "Mainline", 0
            ),
            MetadataInterval(
                "L02", "DN", Decimal("130100"), Decimal("130200"), "EAL DN", "Mainline", 1
            ),
        ),
    )

    assert source.measurements[0].tension_length == "70"


def test_complete_cycle_source_rejects_report_section_that_conflicts_with_segment(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    rows = [ChartDataRecord(
        task_run_date="2026-06-17", line="EAL", track="DN", section="LMC",
        task_no="D3", station_start="START", station_end="END",
        tension_length="L02", chainage=130170.25, wear_min=11.4,
        track_type="Tangent", overlap=None,
    )]
    monkeypatch.setattr(calculation_endpoint, "parse_exception_report", lambda _content: ([], rows))

    with pytest.raises(
        calculation_endpoint.MetadataValidationError,
        match=r"section LMC.*source segments.*D3.*Mainline",
    ):
        calculation_endpoint._parse_complete_cycle_source(
            "20260617_EAL_D3_Exception_Report.xlsx",
            b"source",
            "EAL",
            metadata=(
                MetadataInterval(
                    "L02", "DN", Decimal("130100"), Decimal("130200"), "LMC DN", "LMC"
                ),
            ),
        )


def test_complete_cycle_preview_groups_measurement_metadata_once_for_all_files(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint
    import app.core.calculation.wear_cycle_metadata as metadata_module

    metadata = (
        MetadataInterval("1", "UP", Decimal("100"), Decimal("120"), "EAL UP"),
    )
    original_group = metadata_module._group_intervals_by_tl
    calls = {"group": 0}

    def counted_group(*args, **kwargs):
        calls["group"] += 1
        return original_group(*args, **kwargs)

    def parsed_rows(content):
        task_no = content.decode("ascii")
        rows = [ChartDataRecord(
            task_run_date="2026-06-17", line="EAL", track="UP", section="Mainline",
            task_no=task_no, station_start="START", station_end="END",
            tension_length="1", chainage=110, wear_min=11.4,
            track_type="Tangent", overlap=None,
        )]
        return [], rows

    monkeypatch.setattr(metadata_module, "_group_intervals_by_tl", counted_group)
    monkeypatch.setattr(calculation_endpoint, "_validate_wear_zip_declarations", lambda *_args: None)
    monkeypatch.setattr(calculation_endpoint, "_metadata_manager", lambda _line: object())
    monkeypatch.setattr(calculation_endpoint, "load_line_metadata", lambda *_args: metadata)
    monkeypatch.setattr(calculation_endpoint, "parse_exception_report", parsed_rows)
    monkeypatch.setattr(calculation_endpoint, "build_cycle_preview", lambda **kwargs: kwargs)

    result = calculation_endpoint._build_complete_cycle_preview(
        file_payloads=[("EAL_U1.xlsx", b"U1"), ("EAL_U2.xlsx", b"U2")],
        line_group="EAL",
        cycle_date="2026-06-17",
        accepted_conflict_ids=set(),
    )

    assert len(result["sources"]) == 2
    assert calls == {"group": 1}


def test_complete_cycle_source_excludes_siding_primary_without_losing_mainline_rows(
    monkeypatch,
):
    import app.api.endpoints.calculation as calculation_endpoint

    rows = [
        ChartDataRecord(
            task_run_date="2026-06-17",
            line="EAL",
            track="UP",
            section="Mainline",
            task_no=str(index),
            station_start="START",
            station_end="END",
            tension_length=tension_length,
            chainage=chainage,
            wear_min=11.4,
            track_type="Tangent",
            overlap=None,
        )
        for index, (tension_length, chainage) in enumerate((("1", 110), ("X1", 130)))
    ]
    monkeypatch.setattr(calculation_endpoint, "parse_exception_report", lambda _content: ([], rows))

    source = calculation_endpoint._parse_complete_cycle_source(
        "20260617_EAL_U1_Exception_Report.xlsx",
        b"source",
        "EAL",
        metadata=(
            MetadataInterval("1", "UP", Decimal("100"), Decimal("120"), "EAL UP"),
            MetadataInterval("X1", "UP", Decimal("120"), Decimal("140"), "EAL UP"),
        ),
    )

    assert [measurement.tension_length for measurement in source.measurements] == ["1"]


def test_complete_cycle_source_rejects_unknown_primary_with_actionable_context(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    rows = [ChartDataRecord(
        task_run_date="2026-06-17",
        line="EAL",
        track="UP",
        section="Mainline",
        task_no="U1",
        station_start="START",
        station_end="END",
        tension_length="NEW-TL",
        chainage=110,
        wear_min=11.4,
        track_type="Tangent",
        overlap=None,
    )]
    monkeypatch.setattr(calculation_endpoint, "parse_exception_report", lambda _content: ([], rows))

    with pytest.raises(calculation_endpoint.MetadataValidationError) as exc_info:
        calculation_endpoint._parse_complete_cycle_source(
            "20260617_EAL_U1_Exception_Report.xlsx",
            b"source",
            "EAL",
            metadata=(
                MetadataInterval(
                    "NEW-TL", "UP", Decimal("100"), Decimal("120"), "EAL UP"
                ),
            ),
        )

    detail = str(exc_info.value)
    assert "20260617_EAL_U1_Exception_Report.xlsx" in detail
    assert "EAL" in detail
    assert "NEW-TL" in detail
    assert "section Mainline" in detail
    assert "track UP" in detail
    assert "chainage 110" in detail


def test_complete_cycle_preview_uses_full_source_for_resolution_and_split_canonical_geometry(
    monkeypatch,
):
    import app.api.endpoints.calculation as calculation_endpoint

    class SourceMetadataManager:
        def get_tension_length_source_rows(self, _line, track, section):
            if (track, section) != ("DN", "Mainline"):
                return pd.DataFrame()
            return pd.DataFrame([
                {
                    "from_m": 100,
                    "to_m": 120,
                    "tension_length": "40,44",
                }
            ])

    rows = [ChartDataRecord(
        task_run_date="2026-06-17", line="TML", track="DN", section="Mainline",
        task_no="D1", station_start="START", station_end="END",
        tension_length="40,44", chainage=115, wear_min=11.4,
        track_type="Tangent", overlap="Y",
    )]
    monkeypatch.setattr(calculation_endpoint, "_validate_wear_zip_declarations", lambda *_args: None)
    monkeypatch.setattr(
        calculation_endpoint, "_metadata_manager", lambda _line: SourceMetadataManager()
    )
    monkeypatch.setattr(
        calculation_endpoint, "parse_exception_report", lambda _content: ([], rows)
    )

    preview = calculation_endpoint._build_complete_cycle_preview(
        file_payloads=[("20260617_TML_D1_Exception_Report.xlsx", b"source")],
        line_group="TML",
        cycle_date="2026-06-17",
        accepted_conflict_ids=set(),
    )

    assert len(preview.records) == 1
    record = preview.records[0]
    assert record.key.tension_length == "40"
    assert (record.from_m, record.to_m) == (Decimal("100"), Decimal("110"))


def test_complete_cycle_source_reports_file_raw_tl_track_and_chainage_on_ambiguous_metadata(
    monkeypatch,
):
    import app.api.endpoints.calculation as calculation_endpoint

    rows = [ChartDataRecord(
        task_run_date="2026-05-11", line="TML", track="DN", section="Mainline",
        task_no="D1", station_start="TAW", station_end="WKS",
        tension_length="M21,MX9", chainage=84160.75, wear_min=11.4,
        track_type="Tangent", overlap=None,
    )]
    monkeypatch.setattr(calculation_endpoint, "parse_exception_report", lambda _content: ([], rows))
    metadata = (
        MetadataInterval("M21", "DN", Decimal("84054"), Decimal("84625.9"), "TML DN"),
        MetadataInterval("MX9", "DN", Decimal("84137.9"), Decimal("84160.95"), "TML DN"),
    )

    with pytest.raises(
        calculation_endpoint.MetadataValidationError,
        match=r"D1.xlsx.*M21,MX9.*DN.*84160\.75.*ambiguous metadata",
    ):
        calculation_endpoint._parse_complete_cycle_source(
            "D1.xlsx", b"source", "TML", metadata=metadata
        )


def test_upload_wear_returns_500_for_unexpected_data_version_failure(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    monkeypatch.setattr(
        calculation_endpoint,
        "_build_complete_cycle_preview",
        lambda **_kwargs: _complete_cycle_preview(),
    )
    monkeypatch.setattr(
        calculation_endpoint,
        "get_database",
        lambda: (_ for _ in ()).throw(RuntimeError("database unavailable")),
    )
    with TestClient(app, raise_server_exceptions=False) as error_client:
        response = error_client.post(
            "/api/calculation/wear",
            files=[("files", ("EAL_U1.xlsx", b"source", "application/octet-stream"))],
            data={
                "line_group": "EAL",
                "cycle_date": "2026-05-28",
                "accepted_conflict_ids": "[]",
            },
        )

    assert response.status_code == 500


def test_upload_wear_maps_invalid_parser_exception_to_422_with_filename(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    monkeypatch.setattr(calculation_endpoint, "_metadata_manager", lambda _line: object())
    monkeypatch.setattr(calculation_endpoint, "load_line_metadata", lambda _manager, _line: ())
    monkeypatch.setattr(
        calculation_endpoint,
        "parse_exception_report",
        lambda _content: (_ for _ in ()).throw(zipfile.BadZipFile("invalid archive")),
    )
    with TestClient(app, raise_server_exceptions=False) as error_client:
        response = error_client.post(
            "/api/calculation/wear",
            files=[
                (
                    "files",
                    ("broken-cycle.xlsx", _required_workbook_bytes(), "application/octet-stream"),
                )
            ],
            data={
                "line_group": "EAL",
                "cycle_date": "2026-05-28",
                "accepted_conflict_ids": "[]",
            },
        )

    assert response.status_code == 422
    detail = response.json()["detail"]
    assert "broken-cycle.xlsx" in detail
    assert "invalid archive" not in detail


def test_upload_wear_rejects_non_ooxml_bytes_before_parser():
    with TestClient(app, raise_server_exceptions=False) as error_client:
        response = error_client.post(
            "/api/calculation/wear",
            files=[
                (
                    "files",
                    ("not-ooxml.xlsx", b"arbitrary non-zip bytes", "application/octet-stream"),
                )
            ],
            data={"line_group": "EAL", "cycle_date": "2026-05-28"},
        )

    assert response.status_code == 422
    assert response.json()["detail"] == "Invalid wear upload: not-ooxml.xlsx"


def test_upload_wear_rejects_real_xlsx_missing_required_worksheet_before_parse(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    workbook = Workbook()
    workbook.active.title = "ChartData"
    content = io.BytesIO()
    workbook.save(content)
    parser_calls = []

    monkeypatch.setattr(calculation_endpoint, "_metadata_manager", lambda _line: object())
    monkeypatch.setattr(calculation_endpoint, "load_line_metadata", lambda _manager, _line: ())

    def parser_spy(_content):
        parser_calls.append(True)
        raise RuntimeError("parser must not inspect a structurally invalid workbook")

    monkeypatch.setattr(calculation_endpoint, "parse_exception_report", parser_spy)
    with TestClient(app, raise_server_exceptions=False) as error_client:
        response = error_client.post(
            "/api/calculation/wear",
            files=[("files", ("missing-sheet.xlsx", content.getvalue(), "application/octet-stream"))],
            data={"line_group": "EAL", "cycle_date": "2026-05-28"},
        )

    assert response.status_code == 422
    assert response.json()["detail"] == "Invalid wear upload: missing-sheet.xlsx"
    assert parser_calls == []


def test_upload_wear_leaves_unexpected_parser_exception_as_500(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    monkeypatch.setattr(calculation_endpoint, "_metadata_manager", lambda _line: object())
    monkeypatch.setattr(calculation_endpoint, "load_line_metadata", lambda _manager, _line: ())
    monkeypatch.setattr(
        calculation_endpoint,
        "parse_exception_report",
        lambda _content: (_ for _ in ()).throw(RuntimeError("unexpected parser defect")),
    )
    with TestClient(app, raise_server_exceptions=False) as error_client:
        response = error_client.post(
            "/api/calculation/wear",
            files=[
                ("files", ("runtime.xlsx", _required_workbook_bytes(), "application/octet-stream"))
            ],
            data={"line_group": "EAL", "cycle_date": "2026-05-28"},
        )

    assert response.status_code == 500


@pytest.mark.parametrize("parser_error", [ValueError("parser bug"), TypeError("parser bug")])
def test_upload_wear_leaves_generic_parser_errors_as_500(monkeypatch, parser_error):
    import app.api.endpoints.calculation as calculation_endpoint

    monkeypatch.setattr(calculation_endpoint, "_metadata_manager", lambda _line: object())
    monkeypatch.setattr(calculation_endpoint, "load_line_metadata", lambda _manager, _line: ())

    def fail_parser(_content):
        raise parser_error

    monkeypatch.setattr(calculation_endpoint, "parse_exception_report", fail_parser)
    with TestClient(app, raise_server_exceptions=False) as error_client:
        response = error_client.post(
            "/api/calculation/wear",
            files=[
                ("files", ("runtime.xlsx", _required_workbook_bytes(), "application/octet-stream"))
            ],
            data={"line_group": "EAL", "cycle_date": "2026-05-28"},
        )

    assert response.status_code == 500


def _complete_cycle_upload(client_instance, files):
    return client_instance.post(
        "/api/calculation/wear",
        files=files,
        data={"line_group": "EAL", "cycle_date": "2026-05-28"},
    )


def _declared_zip_bytes(content: bytes) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("payload.txt", content)
    return buffer.getvalue()


def test_complete_cycle_upload_limits_are_desktop_safe():
    import app.api.endpoints.calculation as calculation_endpoint

    assert calculation_endpoint.MAX_WEAR_UPLOAD_FILES == 24
    assert calculation_endpoint.MAX_WEAR_FILE_BYTES == 32 * 1024 * 1024
    assert calculation_endpoint.MAX_WEAR_TOTAL_BYTES == 128 * 1024 * 1024
    assert calculation_endpoint.MAX_WEAR_ZIP_UNCOMPRESSED_BYTES == 512 * 1024 * 1024
    assert calculation_endpoint.MAX_WEAR_ZIP_COMPRESSION_RATIO == 100.0
    assert calculation_endpoint.MAX_WEAR_ARCHIVE_MEMBERS == 2048
    assert calculation_endpoint.MAX_WEAR_WORKBOOK_XML_BYTES == 4 * 1024 * 1024


def test_upload_wear_rejects_too_many_complete_cycle_files(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_UPLOAD_FILES", 1, raising=False)
    response = _complete_cycle_upload(client, [
        ("files", ("one.xlsx", b"a", "application/octet-stream")),
        ("files", ("two.xlsx", b"b", "application/octet-stream")),
    ])

    assert response.status_code == 413
    assert "too many" in response.json()["detail"].lower()


def test_upload_wear_rejects_single_and_total_compressed_limits(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    monkeypatch.setattr(
        calculation_endpoint,
        "_validate_wear_zip_declarations",
        lambda _filename, _content: None,
    )
    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_FILE_BYTES", 3, raising=False)
    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_TOTAL_BYTES", 20, raising=False)
    single = _complete_cycle_upload(
        client,
        [("files", ("large.xlsx", b"1234", "application/octet-stream"))],
    )

    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_FILE_BYTES", 10)
    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_TOTAL_BYTES", 5)
    total = _complete_cycle_upload(client, [
        ("files", ("one.xlsx", b"123", "application/octet-stream")),
        ("files", ("two.xlsx", b"456", "application/octet-stream")),
    ])

    assert single.status_code == 413
    assert "file size" in single.json()["detail"].lower()
    assert total.status_code == 413
    assert "total" in total.json()["detail"].lower()


def test_upload_wear_rejects_zip_declared_size_and_ratio(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    payload = _declared_zip_bytes(b"A" * 10_000)
    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_FILE_BYTES", len(payload) + 10)
    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_TOTAL_BYTES", len(payload) + 10)
    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_ZIP_UNCOMPRESSED_BYTES", 100)
    declared_size = _complete_cycle_upload(
        client,
        [("files", ("bomb.xlsx", payload, "application/octet-stream"))],
    )

    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_ZIP_UNCOMPRESSED_BYTES", 20_000)
    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_ZIP_COMPRESSION_RATIO", 2.0)
    ratio = _complete_cycle_upload(
        client,
        [("files", ("ratio.xlsx", payload, "application/octet-stream"))],
    )

    assert declared_size.status_code == 413
    assert "uncompressed" in declared_size.json()["detail"].lower()
    assert ratio.status_code == 413
    assert "ratio" in ratio.json()["detail"].lower()


def test_upload_wear_rejects_high_ratio_member_hidden_by_zip_padding(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("bomb.txt", b"A" * 1_000_000)
        archive.writestr("padding.bin", os.urandom(20_000))
    payload = buffer.getvalue()
    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_FILE_BYTES", len(payload) + 10)
    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_TOTAL_BYTES", len(payload) + 10)
    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_ZIP_UNCOMPRESSED_BYTES", 2_000_000)

    response = _complete_cycle_upload(
        client,
        [("files", ("padded-bomb.xlsx", payload, "application/octet-stream"))],
    )

    assert response.status_code == 413
    assert "ratio" in response.json()["detail"].lower()


def test_upload_wear_rejects_excessive_archive_members_including_empty_entries(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("directory/", b"")
        archive.writestr("empty.txt", b"")
    payload = buffer.getvalue()
    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_FILE_BYTES", len(payload) + 10)
    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_TOTAL_BYTES", len(payload) + 10)
    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_ARCHIVE_MEMBERS", 1, raising=False)

    response = _complete_cycle_upload(
        client,
        [("files", ("many-members.xlsx", payload, "application/octet-stream"))],
    )

    assert response.status_code == 413
    assert "member" in response.json()["detail"].lower()


def _make_mock_excel_bytes() -> bytes:
    """Build a minimal valid Exception Report Excel in memory."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        pd.DataFrame({
            'Run Date': ['2026-01-01', '2026-02-01'],
            'ID': ['W001', 'W002'],
            'Tension Length': ['H02', 'H02'],
            'FromM': [1000.0, 1000.0],
            'ToM': [1100.0, 1100.0],
            'MaxValue': [10.5, 10.3],
            'Level': ['L2', 'L2'],
            'ACTION': [None, None],
        }).to_excel(writer, sheet_name='Wire Wear', index=False)

        pd.DataFrame({
            'task_run_date': ['2026-01-01', '2026-01-01', '2026-02-01', '2026-02-01'],
            'Tension Length': ['H02', 'H02', 'H02', 'H02'],
            'Chainage': [1000.0, 1050.0, 1000.0, 1050.0],
            'wear_min': [10.5, 10.3, 10.1, 9.9],
            'Track Type': ['Tangent', 'Tangent', 'Tangent', 'Tangent'],
            'Overlap': [None, None, None, None],
        }).to_excel(writer, sheet_name='ChartData', index=False)
    return buf.getvalue()


def _make_legacy_wear_excel_bytes() -> bytes:
    """Build a legacy Exception Report where ChartData has wear1..wear4 instead of wear_min."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        pd.DataFrame({
            'Run Date': ['2026-01-16'],
            'ID': ['20260116_EAL_U3_TAP-LOW_W35'],
            'Tension Length': ['35'],
            'FromM': [114644.0],
            'ToM': [114644.5],
            'MaxValue': [10.1],
            'Level': ['L2'],
            'ACTION': [None],
        }).to_excel(writer, sheet_name='Wire Wear', index=False)

        pd.DataFrame({
            'task_run_date': ['2026-01-16', '2026-01-16'],
            'Tension Length': ['35', '35'],
            'Chainage': [114644.0, 114644.5],
            'wear1': [10.8, 10.2],
            'wear2': [10.4, None],
            'wear3': [10.6, 10.1],
            'wear4': [None, 10.5],
            'Track Type': ['Curve', 'Curve'],
            'Overlap': [None, None],
        }).to_excel(writer, sheet_name='ChartData', index=False)
    return buf.getvalue()


def test_health_endpoint():
    """GET /api/calculation/health should return {"status": "ok"}."""
    resp = client.get('/api/calculation/health')
    assert resp.status_code == 200
    assert resp.json() == {'status': 'ok'}


def test_app_health_endpoint_includes_expected_version():
    resp = client.get('/api/health')
    assert resp.status_code == 200
    assert resp.json() == {'status': 'ok', 'version': '2.0.0'}


def test_upload_endpoint_no_files():
    """POST /api/calculation/upload with no files should return 422."""
    resp = client.post('/api/calculation/upload')
    assert resp.status_code == 422


def test_upload_endpoint_with_mock():
    """POST /api/calculation/upload with mock Excel bytes returns wear and trend results."""
    excel_bytes = _make_mock_excel_bytes()
    resp = client.post(
        '/api/calculation/upload',
        files=[('files', ('mock_report.xlsx', excel_bytes,
                          'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))],
        data={'line': 'EAL'},
    )
    assert resp.status_code == 200

    body = resp.json()
    assert 'wear_results' in body
    assert 'trend_results' in body
    assert isinstance(body['wear_results'], list)
    assert isinstance(body['trend_results'], list)

    # Verify wear result structure
    assert len(body['wear_results']) > 0
    wr = body['wear_results'][0]
    for field in ('tension_length', 'from_m', 'to_m', 'avg_wear_min', 'sd', 'wear_percentage'):
        assert field in wr, f"wear_results missing field: {field}"

    # Verify trend result structure when present
    if body['trend_results']:
        tr = body['trend_results'][0]
        for field in ('tension_length', 'from_m', 'to_m', 'dates', 'record_points',
                      'trend_points', 'trend_next', 'logic_1', 'logic_2', 'recommendation'):
            assert field in tr, f"trend_results missing field: {field}"


def test_upload_wear_accepts_legacy_chartdata_wear_channels():
    excel_bytes = _make_legacy_wear_excel_bytes()

    resp = client.post(
        '/api/calculation/wear',
        files=[('files', ('legacy_wear_report.xlsx', excel_bytes,
                          'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))],
        data={'line': 'EAL'},
    )

    assert resp.status_code == 200
    body = resp.json()
    assert body['date'] == '2026-01-16'
    assert len(body['wear_results']) == 1
    assert body['wear_results'][0]['avg_wear_min'] == pytest.approx((10.4 + 10.1) / 2, abs=0.01)


def test_upload_wear_partitions_dedupe_and_metadata_lookup_by_identity(monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint

    chart_data = [
        ChartDataRecord(
            task_run_date='2026-01-01',
            line='TML',
            track='UP',
            section='Mainline',
            task_no='U1',
            station_start='WKS',
            station_end='TAW',
            tension_length='RAW',
            chainage=1000.0,
            wear_min=10.0,
            track_type='Tangent',
            overlap=None,
        ),
        ChartDataRecord(
            task_run_date='2026-02-01',
            line='TML',
            track='DN',
            section='Mainline',
            task_no='D1',
            station_start='TAW',
            station_end='WKS',
            tension_length='RAW',
            chainage=1000.0,
            wear_min=12.0,
            track_type='Tangent',
            overlap=None,
        ),
    ]

    def fake_parse_exception_report(_file_bytes):
        return [], chart_data

    class FakeMetadataManager:
        def __init__(self, config_path):
            self.config_path = config_path

        def get_tension_length_lookup(self, line, track, section):
            tension_length = 'UP-TL' if track == 'UP' else 'DN-TL'
            return pd.DataFrame([{
                'from_m': 990.0,
                'to_m': 1010.0,
                'tension_length': tension_length,
                'track_type': 'Tangent',
            }])

    monkeypatch.setattr(calculation_endpoint, 'parse_exception_report', fake_parse_exception_report)
    monkeypatch.setattr(calculation_endpoint, 'MetadataManager', FakeMetadataManager)

    response = client.post(
        '/api/calculation/wear',
        files=[('files', ('mock.xlsx', b'content', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))],
        data={'line': 'TML'},
    )

    assert response.status_code == 200
    results = response.json()['wear_results']
    identities = {(row['track'], row['section'], row['tension_length'], row['cycle_date']) for row in results}
    assert identities == {
        ('UP', 'Mainline', 'UP-TL', '2026-01-01'),
        ('DN', 'Mainline', 'DN-TL', '2026-02-01'),
    }


def test_upload_trend_returns_results_for_real_u3_files_without_repeated_report():
    base_dir = Path(__file__).resolve().parents[2] / 'docs' / 'test data' / 'EAL' / 'U3'
    report_paths = [
        base_dir / '20251206_EAL_U3_TAP-LOW_Exception_Report.xlsx',
        base_dir / '20251227_EAL_U3_TAP-LOW_Exception_Report.xlsx',
        base_dir / '20260116_EAL_U3_TAP-LOW_Exception_Report.xlsx',
    ]
    missing = [path for path in report_paths if not path.is_file()]
    if missing:
        pytest.skip(
            "Deleted historical EAL U3 fixtures unavailable: "
            + ", ".join(path.name for path in missing)
        )

    response = client.post(
        '/api/calculation/trend',
        files=[
            (
                'files',
                (
                    path.name,
                    path.read_bytes(),
                    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                ),
            )
            for path in report_paths
        ],
    )

    assert response.status_code == 200
    body = response.json()
    assert len(body['trend_results']) > 0


def test_parse_real_tml_exception_report_uses_streaming_chartdata_path(monkeypatch):
    report_path = (
        Path(__file__).resolve().parents[2]
        / 'docs' / 'test data' / 'TML' / 'Cycle 2'
        / '20260413_TML_U1_WKS-TAW_Exception_Report.xlsx'
    )
    import app.core.calculation.excel_parser as excel_parser

    def fail_pandas_fallback(*_args, **_kwargs):
        raise AssertionError('pandas fallback should not be used for a valid TML Exception Report')

    monkeypatch.setattr(excel_parser.pd, 'ExcelFile', fail_pandas_fallback)
    wire_wear, chart_data = parse_exception_report(report_path.read_bytes())

    assert len(wire_wear) > 0
    assert len(chart_data) > 40000
    assert {record.line for record in chart_data} == {'TML'}


@pytest.mark.parametrize(
    "relative_path",
    [
        Path("docs/test data/EAL/Cycle 8/20260608_EAL_U1_HUH-FOT_Exception_Report.xlsx"),
        Path("docs/test data/EAL/Cycle 8/20260611_EAL_UP_RAC_Exception_Report.xlsx"),
        Path("docs/test data/EAL/Cycle 8/20260613_EAL_UP_LOW_Exception_Report.xlsx"),
        Path("docs/test data/LMC/20260418_EAL_UP_LMC_Exception_Report.xlsx"),
    ],
    ids=("mainline", "rac", "low-s1", "lmc"),
)
def test_real_eal_wear_sources_resolve_mainline_only(relative_path):
    repository_root = Path(__file__).resolve().parents[2]
    report_path = repository_root / relative_path
    if not report_path.is_file():
        pytest.skip(f"Real-file fixture unavailable: {relative_path}")

    response = client.post(
        "/api/calculation/wear",
        files=[(
            "files",
            (
                report_path.name,
                report_path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        )],
        data={
            "line_group": "EAL",
            "cycle_date": "2026-06-17",
            "accepted_conflict_ids": "[]",
        },
    )

    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["records"]
    assert {
        classify_tension_length_scope("EAL", record["tension_length"])
        for record in payload["records"]
    } == {TensionLengthScope.MAINLINE}
    diagnostic_gaps = {
        gap
        for segment in payload["segments"]
        for gap in segment["diagnostic_gaps"]
        if gap not in {"segment_missing", "metadata_denominator_empty"}
    }
    assert all(
        classify_tension_length_scope("EAL", gap) is TensionLengthScope.MAINLINE
        for gap in diagnostic_gaps
    )


def test_real_tml_complete_cycle_matches_mainline_export_scope():
    repository_root = Path(__file__).resolve().parents[2]
    cycle_dir = repository_root / "docs/test data/TML/Cycle 2"
    report_paths = sorted(cycle_dir.glob("*Exception_Report.xlsx"))
    if len(report_paths) != 6:
        pytest.skip(
            f"Expected six TML Cycle 2 reports, found {len(report_paths)} in {cycle_dir}"
        )

    response = client.post(
        "/api/calculation/wear",
        files=[
            (
                "files",
                (
                    report_path.name,
                    report_path.read_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            )
            for report_path in report_paths
        ],
        data={
            "line_group": "TML",
            "cycle_date": "2026-05-12",
            "accepted_conflict_ids": "[]",
        },
    )

    assert response.status_code == 200, response.text
    payload = response.json()
    assert len(payload["records"]) == 164
    assert {
        classify_tension_length_scope("TML", record["tension_length"])
        for record in payload["records"]
    } == {TensionLengthScope.MAINLINE}
    assert all(
        classify_tension_length_scope("TML", gap) is TensionLengthScope.MAINLINE
        for segment in payload["segments"]
        for gap in segment["diagnostic_gaps"]
        if gap not in {"segment_missing", "metadata_denominator_empty"}
    )
