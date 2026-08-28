import os
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.core.database import DatabaseManager
from app.main import app
from app.core.calculation.wear_cycle_aggregation import preview_digest
from app.core.calculation.wear_cycle_types import (
    AggregatedWearRecord,
    BusinessKey,
    ConflictPreview,
    CyclePreview,
    MetadataInterval,
    SegmentCoverage,
)


@pytest.fixture()
def client(tmp_path):
    db_path = tmp_path / "api_wear_records.db"
    os.environ["TOV640_TEST_DB_PATH"] = str(db_path)
    DatabaseManager.reset_instance()
    manager = DatabaseManager(str(db_path))
    manager.close()
    with TestClient(app) as test_client:
        yield test_client
    DatabaseManager.reset_instance()
    os.environ.pop("TOV640_TEST_DB_PATH", None)


def _payload():
    return {
        "line_group": "EAL",
        "line_class": "LMC",
        "track": "UP",
        "section": "LMC",
        "cycle_date": "2026-02-01",
        "source_file_names": ["cycle.xlsx"],
        "saved_by": "api-test",
        "records": [
            {
                "tension_length": "H46",
                "from_m": 100.0,
                "to_m": 200.0,
                "avg_wear_min": 12.4,
                "sd": 0.2,
                "wear_percentage": 6.5,
            }
        ],
    }


def _cycle_preview():
    record = AggregatedWearRecord(
        BusinessKey("EAL", date(2026, 5, 28), "28"), "UP",
        Decimal("100"), Decimal("200"), 11.4, 10.0, None, False, (), ("EAL_U1.xlsx",),
    )
    segments = tuple(
        SegmentCoverage(name, True, 100.0, (), (f"EAL_{name}.xlsx",), (date(2026, 5, 28),))
        for name in ("U1", "U2", "U3", "D1", "D2", "D3", "LOW S1", "RAC UP", "RAC DN", "LMC UP", "LMC DN")
    )
    return CyclePreview(
        "EAL", date(2026, 5, 28), (record,), segments, (), (), (), True,
        datetime(2026, 5, 28, tzinfo=timezone.utc),
    )


def _metadata():
    return (MetadataInterval("28", "UP", Decimal("100"), Decimal("200"), "EAL UP"),)


def _historical_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "EAL"
    sheet.append(["Track", "Tension Length", "202401"])
    sheet.append(["UP", "28", 11.2])
    workbook.create_sheet("LRL")
    workbook.create_sheet("Notes")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_cycle_save_rebuilds_preview_and_rejects_duplicate(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    preview = _cycle_preview()
    rebuild_calls = []

    def rebuild_spy(**kwargs):
        rebuild_calls.append(kwargs)
        return preview

    monkeypatch.setattr(endpoint, "_build_complete_cycle_preview", rebuild_spy, raising=False)
    fields = {
        "line_group": "EAL",
        "cycle_date": "2026-05-28",
        "accepted_conflict_ids": "[]",
        "expected_preview_digest": preview_digest(preview),
        "expected_data_version": "0",
    }
    files = [("files", ("EAL_U1.xlsx", b"source", "application/octet-stream"))]

    saved = client.post("/api/calculation/wear-records/cycles", files=files, data=fields)
    duplicate = client.post("/api/calculation/wear-records/cycles", files=files, data={**fields, "expected_data_version": "1"})

    assert saved.status_code == 200
    assert saved.json()["cycle_date"] == "2026-05-28"
    assert saved.json()["wire_wear_data_version"] == 1
    assert duplicate.status_code == 409
    assert rebuild_calls == [
        {
            "file_payloads": [("EAL_U1.xlsx", b"source")],
            "line_group": "EAL",
            "cycle_date": "2026-05-28",
            "accepted_conflict_ids": set(),
        },
        {
            "file_payloads": [("EAL_U1.xlsx", b"source")],
            "line_group": "EAL",
            "cycle_date": "2026-05-28",
            "accepted_conflict_ids": set(),
        },
    ]


def test_cycle_save_rejects_upload_limits_before_rebuild(client, monkeypatch):
    import app.api.endpoints.calculation as calculation_endpoint
    import app.api.endpoints.wear_records as endpoint

    monkeypatch.setattr(calculation_endpoint, "MAX_WEAR_UPLOAD_FILES", 1, raising=False)
    monkeypatch.setattr(
        endpoint,
        "_build_complete_cycle_preview",
        lambda **_kwargs: pytest.fail("oversized upload must not rebuild preview"),
    )
    response = client.post(
        "/api/calculation/wear-records/cycles",
        files=[
            ("files", ("one.xlsx", b"a", "application/octet-stream")),
            ("files", ("two.xlsx", b"b", "application/octet-stream")),
        ],
        data={
            "line_group": "EAL",
            "cycle_date": "2026-05-28",
            "accepted_conflict_ids": "[]",
            "preview_digest": "digest",
            "expected_data_version": "0",
        },
    )

    assert response.status_code == 413
    assert "too many" in response.json()["detail"].lower()


def test_cycle_save_rejects_digest_mismatch(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    monkeypatch.setattr(endpoint, "_build_complete_cycle_preview", lambda **_kwargs: _cycle_preview(), raising=False)
    response = client.post(
        "/api/calculation/wear-records/cycles",
        files=[("files", ("EAL_U1.xlsx", b"source", "application/octet-stream"))],
        data={
            "line_group": "EAL", "cycle_date": "2026-05-28",
            "accepted_conflict_ids": "[]", "expected_preview_digest": "stale",
            "expected_data_version": "0",
        },
    )

    assert response.status_code == 409


def test_cycle_save_rejects_unaccepted_conflict(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    base = _cycle_preview()
    conflict = ConflictPreview(
        "conflict-1", "measurement-1", (("a.xlsx", 11.4), ("b.xlsx", 11.5)), 11.4, False,
    )
    preview = CyclePreview(
        base.line_group, base.cycle_date, base.records, base.segments, (conflict,), (),
        ("conflict_not_accepted",), False, base.generated_at,
    )
    monkeypatch.setattr(endpoint, "_build_complete_cycle_preview", lambda **_kwargs: preview)
    response = client.post(
        "/api/calculation/wear-records/cycles",
        files=[("files", ("EAL_U1.xlsx", b"source", "application/octet-stream"))],
        data={
            "line_group": "EAL", "cycle_date": "2026-05-28",
            "accepted_conflict_ids": "[]", "preview_digest": preview_digest(preview),
            "expected_data_version": "0",
        },
    )

    assert response.status_code == 422
    assert response.json()["detail"]["blocking_reasons"] == ["conflict_not_accepted"]


def test_changes_endpoint_returns_409_for_stale_timestamp(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    monkeypatch.setattr(endpoint, "_load_cycle_metadata", lambda _line: _metadata(), raising=False)
    added = client.post("/api/calculation/wear-records/changes", json={"operations": [{
        "kind": "add", "key": {"line_group": "EAL", "cycle_date": "2026-05-28", "tension_length": "28"},
        "avg_wear_min": 11.4,
    }]})
    assert added.status_code == 200

    response = client.post("/api/calculation/wear-records/changes", json={"operations": [{
        "kind": "edit", "key": {"line_group": "EAL", "cycle_date": "2026-05-28", "tension_length": "28"},
        "avg_wear_min": 11.2, "expected_updated_at": "2000-01-01T00:00:00Z",
    }]})

    assert response.status_code == 409
    assert response.json()["detail"]["operation_index"] == 0


def test_workbook_change_set_creates_backup_before_apply(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    monkeypatch.setattr(endpoint, "_load_identity_metadata", lambda _line, _line_class: _metadata())
    response = client.post("/api/calculation/wear-records/changes", json={
        "origin": "workbook",
        "expected_data_version": 0,
        "operations": [{
            "kind": "add",
            "key": {
                "line_group": "EAL", "line_class": "EAL",
                "cycle_date": "2026-05-28", "tension_length": "28",
            },
            "avg_wear_min": 11.4,
        }],
    })

    assert response.status_code == 200
    assert response.json()["added"] == 1
    backup_path = response.json()["backup_path"]
    assert backup_path
    assert Path(backup_path).is_file()


def test_workbook_change_set_applies_multiple_line_identities_atomically(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    metadata_by_identity = {
        ("EAL", "EAL"): (
            MetadataInterval("28", "UP", Decimal("100"), Decimal("200"), "EAL UP"),
        ),
        ("EAL", "LMC"): (
            MetadataInterval("H46", "UP", Decimal("300"), Decimal("400"), "LMC UP"),
        ),
        ("TML", "TML"): (
            MetadataInterval("99", "UP", Decimal("500"), Decimal("600"), "TML UP"),
        ),
    }
    monkeypatch.setattr(
        endpoint,
        "_load_identity_metadata",
        lambda line, line_class: metadata_by_identity[(line, line_class)],
    )
    response = client.post("/api/calculation/wear-records/changes", json={
        "origin": "workbook",
        "expected_data_version": 0,
        "operations": [
            {
                "kind": "add",
                "key": {
                    "line_group": line_group,
                    "line_class": line_class,
                    "cycle_date": "2026-05-28",
                    "tension_length": tension_length,
                },
                "avg_wear_min": avg_wear_min,
            }
            for line_group, line_class, tension_length, avg_wear_min in (
                ("EAL", "EAL", "28", 11.4),
                ("EAL", "LMC", "H46", 11.3),
                ("TML", "TML", "99", 11.2),
            )
        ],
    })

    assert response.status_code == 200
    assert response.json()["added"] == 3
    assert response.json()["wire_wear_data_version"] == 1
    backup_path = response.json()["backup_path"]
    assert backup_path
    assert Path(backup_path).is_file()

    for line_group, line_class in (("EAL", "EAL"), ("EAL", "LMC"), ("TML", "TML")):
        workbench = client.get(
            "/api/calculation/wear-records/workbench",
            params={"line_group": line_group, "line_class": line_class},
        )
        assert workbench.status_code == 200
        assert len(workbench.json()["records"]) == 1


def test_historical_workbook_discovery_and_preview_api(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    monkeypatch.setattr(endpoint, "_historical_metadata", lambda: _metadata())
    workbook_bytes = _historical_workbook_bytes()
    files = {
        "file": (
            "historical.xlsx",
            workbook_bytes,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }

    discovery = client.post(
        "/api/calculation/wear-records/historical/discover",
        files=files,
    )
    preview = client.post(
        "/api/calculation/wear-records/historical/preview",
        files=files,
        data={"selected_sheets": '["EAL"]'},
    )

    assert discovery.status_code == 200
    capabilities = {item["sheet_name"]: item for item in discovery.json()["sheets"]}
    assert capabilities["EAL"]["selected_by_default"] is True
    assert capabilities["LRL"]["support"] == "reserved"
    assert capabilities["Notes"]["support"] == "ignored"
    assert preview.status_code == 200
    body = preview.json()
    assert body["selected_sheets"] == ["EAL"]
    assert body["totals"]["new"] == 1
    assert body["candidates"][0]["key"] == {
        "line_group": "EAL",
        "line_class": "EAL",
        "cycle_date": "2024-01-01",
        "tension_length": "28",
    }
    assert body["candidates"][0]["source_cell"] == "C2"


def test_candidate_preview_classifies_new_update_no_change_duplicates_and_errors(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    metadata = tuple(
        MetadataInterval(name, "UP", Decimal(start), Decimal(end), "EAL UP")
        for name, start, end in (("28", "100", "200"), ("29", "200", "300"), ("30", "300", "400"))
    )
    monkeypatch.setattr(endpoint, "_load_identity_metadata", lambda _line, _line_class: metadata)
    added = client.post("/api/calculation/wear-records/changes", json={"operations": [
        {
            "kind": "add",
            "key": {
                "line_group": "EAL", "line_class": "EAL",
                "cycle_date": "2026-05-28", "tension_length": "28",
            },
            "avg_wear_min": 11.4,
        },
        {
            "kind": "add",
            "key": {
                "line_group": "EAL", "line_class": "EAL",
                "cycle_date": "2026-05-28", "tension_length": "29",
            },
            "avg_wear_min": 11.3,
        },
    ]})
    assert added.status_code == 200

    response = client.post("/api/calculation/wear-records/candidates/preview", json={
        "line_class": "EAL",
        "cycle_date": "2026-05-28",
        "rows": [
            {"row_id": "same", "tension_length": "28", "avg_wear_min": 11.4},
            {"row_id": "update", "tension_length": "29", "avg_wear_min": 11.2},
            {"row_id": "new", "tension_length": "30", "avg_wear_min": 10.9},
            {"row_id": "dup", "tension_length": "30", "avg_wear_min": 10.9},
            {"row_id": "bad", "tension_length": "28", "avg_wear_min": "not-a-number"},
        ],
    })

    assert response.status_code == 200
    body = response.json()
    assert body["wire_wear_data_version"] == 1
    assert [item["status"] for item in body["candidates"]] == [
        "no_change", "update", "new", "duplicate", "error",
    ]
    assert body["candidates"][0]["expected_updated_at"]
    assert body["candidates"][1]["existing_avg_wear_min"] == 11.3
    assert body["candidates"][4]["issues"][0] == {
        "code": "invalid_avg_wear_min",
        "message": "could not convert string to float: 'not-a-number'",
        "field": "avg_wear_min",
        "cell": "B5",
    }


def test_candidate_preview_maps_lmc_identity_and_does_not_mutate_database(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    lmc_metadata = (
        MetadataInterval("H46", "UP", Decimal("100"), Decimal("200"), "LMC UP"),
    )
    monkeypatch.setattr(endpoint, "_load_identity_metadata", lambda _line, _line_class: lmc_metadata)

    response = client.post("/api/calculation/wear-records/candidates/preview", json={
        "line_class": "LMC",
        "cycle_date": "2026-02-01",
        "rows": [{"row_id": "1", "tension_length": "H46", "avg_wear_min": 12.4}],
    })
    workbench = client.get(
        "/api/calculation/wear-records/workbench?line_group=EAL&line_class=LMC"
    )

    assert response.status_code == 200
    assert response.json()["candidates"][0]["key"] == {
        "line_group": "EAL", "line_class": "LMC",
        "cycle_date": "2026-02-01", "tension_length": "H46",
    }
    assert response.json()["candidates"][0]["status"] == "new"
    assert workbench.json()["records"] == []
    assert workbench.json()["wire_wear_data_version"] == 0


def test_changes_endpoint_rejects_stale_candidate_data_version(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    monkeypatch.setattr(endpoint, "_load_identity_metadata", lambda _line, _line_class: _metadata())
    added = client.post("/api/calculation/wear-records/changes", json={
        "expected_data_version": 0,
        "operations": [{
            "kind": "add",
            "key": {
                "line_group": "EAL", "line_class": "EAL",
                "cycle_date": "2026-05-28", "tension_length": "28",
            },
            "avg_wear_min": 11.4,
        }],
    })
    assert added.status_code == 200

    stale = client.post("/api/calculation/wear-records/changes", json={
        "expected_data_version": 0,
        "operations": [{
            "kind": "edit",
            "key": {
                "line_group": "EAL", "line_class": "EAL",
                "cycle_date": "2026-05-28", "tension_length": "28",
            },
            "avg_wear_min": 11.2,
            "expected_updated_at": "ignored-because-version-is-stale",
        }],
    })

    assert stale.status_code == 409
    assert "changed since the candidate preview" in stale.json()["detail"]


def test_changes_delete_row_and_workbench_contract(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    monkeypatch.setattr(endpoint, "_load_cycle_metadata", lambda _line: _metadata(), raising=False)
    assert client.post("/api/calculation/wear-records/changes", json={"operations": [{
        "kind": "add", "key": {"line_group": "EAL", "cycle_date": "2026-05-28", "tension_length": "28"},
        "avg_wear_min": 11.4,
    }]}).status_code == 200

    workbench = client.get("/api/calculation/wear-records/workbench?line_group=EAL")
    deleted = client.post("/api/calculation/wear-records/changes", json={"operations": [{
        "kind": "delete_row", "line_group": "EAL", "cycle_date": "2026-05-28",
    }]})

    assert workbench.status_code == 200
    assert set(workbench.json()) >= {
        "line_group", "columns", "matrix_rows", "latest_summary", "records",
        "catalog", "wire_wear_data_version",
    }
    assert deleted.status_code == 200
    assert deleted.json()["deleted"] == 1


def test_metadata_preview_resolves_manual_tension_length(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    monkeypatch.setattr(endpoint, "_load_cycle_metadata", lambda _line: _metadata(), raising=False)
    response = client.get("/api/calculation/wear-records/metadata-preview?line_group=EAL&tension_length=28")

    assert response.status_code == 200
    assert response.json()["catalog"] == [{
        "line_group": "EAL", "line_class": "EAL", "tension_length": "28",
        "track": "UP", "from_m": 100.0, "to_m": 200.0,
        "interval_count": 1,
        "intervals": [{"track": "UP", "from_m": 100.0, "to_m": 200.0}],
    }]


def test_sync_metadata_keeps_line_catalogs_isolated(monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    eal = (MetadataInterval("28", "UP", Decimal("0"), Decimal("10"), "EAL UP"),)
    lmc = (MetadataInterval("28", "UP", Decimal("20"), Decimal("30"), "LMC UP"),)
    tml = (MetadataInterval("28", "UP", Decimal("100"), Decimal("110"), "TML UP"),)
    monkeypatch.setattr(
        endpoint,
        "_load_cycle_metadata",
        lambda line, line_class=None: (
            lmc if line_class == "LMC" else eal if line == "EAL" else tml
        ),
    )

    assert endpoint._sync_metadata() == {
        ("EAL", "EAL"): eal,
        ("EAL", "LMC"): lmc,
        ("TML", "TML"): tml,
    }


def test_save_wire_wear_records_api(client):
    response = client.post("/api/calculation/wear-records", json=_payload())

    assert response.status_code == 200
    assert response.json()["saved_count"] == 1


def test_save_wire_wear_records_conflict_then_overwrite(client):
    assert client.post("/api/calculation/wear-records", json=_payload()).status_code == 200

    conflict = client.post("/api/calculation/wear-records", json=_payload())
    assert conflict.status_code == 409
    assert conflict.json()["detail"]["duplicate_count"] == 1

    overwrite = client.post("/api/calculation/wear-records?overwrite=true", json=_payload())
    assert overwrite.status_code == 200
    assert overwrite.json()["updated_count"] == 1


def test_save_wire_wear_records_check_only_does_not_insert(client):
    response = client.post("/api/calculation/wear-records?check_only=true", json=_payload())
    records = client.get("/api/calculation/wear-records?line_group=EAL")

    assert response.status_code == 200
    assert response.json()["saved_count"] == 0
    assert records.status_code == 200
    assert records.json()["records"] == []


def test_save_wire_wear_records_invalid_date_returns_422(client):
    payload = _payload()
    payload["cycle_date"] = "not-a-date"

    response = client.post("/api/calculation/wear-records", json=payload)

    assert response.status_code == 422


def test_list_wire_wear_records_invalid_date_filter_returns_422(client):
    response = client.get("/api/calculation/wear-records?date_from=not-a-date")

    assert response.status_code == 422


def test_list_dashboard_projection_api(client):
    client.post("/api/calculation/wear-records", json=_payload())

    records = client.get("/api/calculation/wear-records?line_group=EAL")
    dashboard = client.get("/api/calculation/wear-records/dashboard")
    projection = client.get("/api/calculation/wear-records/projection")

    assert records.status_code == 200
    assert records.json()["records"][0]["line_class"] == "LMC"
    assert records.json()["records"][0]["saved_by"] == "api-test"
    assert dashboard.status_code == 200
    assert "EAL" in dashboard.json()["line_groups"]
    assert projection.status_code == 200
    assert projection.json()["threshold_mm"] == 10.2


def test_cycle_dashboard_and_projection_use_new_query_contract(client):
    dashboard = client.get("/api/calculation/wear-records/dashboard?line_group=EAL")
    projection = client.get("/api/calculation/wear-records/projection?threshold_mm=10.2")
    invalid_projection = client.get("/api/calculation/wear-records/projection?threshold_mm=0")

    assert dashboard.status_code == 200
    assert set(dashboard.json()["line_groups"]) == {"EAL"}
    assert projection.status_code == 200
    assert projection.json()["threshold_mm"] == 10.2
    assert "threshold_percentage" in projection.json()
    assert invalid_projection.status_code == 422


def test_projection_endpoint_serializes_reference_like_tiny_positive_rate(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    monkeypatch.setattr(endpoint, "_load_cycle_metadata", lambda _line: _metadata(), raising=False)
    for cycle_date, thickness in (("2024-01-01", 12.0), ("2025-01-01", 11.999999)):
        response = client.post("/api/calculation/wear-records/changes", json={"operations": [{
            "kind": "add",
            "key": {
                "line_group": "EAL",
                "cycle_date": cycle_date,
                "tension_length": "28",
            },
            "avg_wear_min": thickness,
        }]})
        assert response.status_code == 200

    projection = client.get(
        "/api/calculation/wear-records/projection?threshold_mm=10.2"
    )

    assert projection.status_code == 200
    outside = projection.json()["line_groups"]["EAL"]["outside_horizon"]
    assert outside[0]["tension_length"] == "28"
    assert outside[0]["projected_crossing_date"] is None


def test_remaining_life_endpoint_returns_nullable_status_contract(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    monkeypatch.setattr(endpoint, "_load_cycle_metadata", lambda _line: _metadata(), raising=False)
    for cycle_date, thickness in (("2024-01-01", 12.0), ("2025-01-01", 11.0)):
        response = client.post("/api/calculation/wear-records/changes", json={"operations": [{
            "kind": "add",
            "key": {"line_group": "EAL", "cycle_date": cycle_date, "tension_length": "28"},
            "avg_wear_min": thickness,
        }]})
        assert response.status_code == 200

    remaining = client.get("/api/calculation/wear-records/remaining-life?threshold_mm=10.2")

    assert remaining.status_code == 200
    payload = remaining.json()
    assert payload["threshold_mm"] == 10.2
    assert payload["rows"][0]["line_group"] == "EAL"
    assert payload["rows"][0]["remaining_days"] is not None
    assert all(point["remaining_days"] >= 0 for point in payload["rows"][0]["curve"])


def test_workbench_summary_and_selected_tension_length_contract(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    monkeypatch.setattr(endpoint, "_load_cycle_metadata", lambda _line: _metadata(), raising=False)
    for cycle_date in ("2026-05-28", "2026-06-28"):
        response = client.post("/api/calculation/wear-records/changes", json={"operations": [{
            "kind": "add",
            "key": {
                "line_group": "EAL",
                "cycle_date": cycle_date,
                "tension_length": "28",
            },
            "avg_wear_min": 11.0,
        }]})
        assert response.status_code == 200

    summary = client.get(
        "/api/calculation/wear-records/workbench?line_group=EAL&summary_only=true"
    )
    selected = client.get(
        "/api/calculation/wear-records/workbench?line_group=EAL&selected_tension_length=28"
    )

    assert summary.status_code == 200
    assert summary.json()["summary_only"] is True
    assert summary.json()["matrix_rows"] == []
    assert summary.json()["records"] == []
    assert selected.status_code == 200
    assert {row["tension_length"] for row in selected.json()["records"]} == {"28"}


def test_cycle_workbench_summary_includes_all_history_trend_fields(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    monkeypatch.setattr(endpoint, "_load_cycle_metadata", lambda _line: _metadata(), raising=False)
    for cycle_date, thickness in (("2025-05-28", 12.0), ("2026-05-28", 11.0)):
        response = client.post("/api/calculation/wear-records/changes", json={"operations": [{
            "kind": "add",
            "key": {"line_group": "EAL", "cycle_date": cycle_date, "tension_length": "28"},
            "avg_wear_min": thickness,
        }]})
        assert response.status_code == 200

    response = client.get("/api/calculation/wear-records/workbench?line_group=EAL")

    assert response.status_code == 200
    summary = response.json()["latest_summary"][0]
    assert summary["trend_status"] == "eligible"
    assert summary["observation_count"] == 2
    assert summary["wear_rate_mm_per_year"] > 0
    assert summary["r_squared"] == pytest.approx(1.0)


def test_workbench_api_returns_history_summary_and_details(client):
    payload = _payload()
    payload["records"][0]["tension_length"] = "X1"
    payload["records"][0]["avg_wear_min"] = 12.0
    payload["records"][0]["wear_percentage"] = 10.0
    assert client.post("/api/calculation/wear-records", json=payload).status_code == 200

    second = _payload()
    second["line_class"] = "EAL"
    second["section"] = "Mainline"
    second["track"] = "DN"
    second["records"][0]["tension_length"] = "X1"
    second["records"][0]["avg_wear_min"] = 10.0
    second["records"][0]["wear_percentage"] = 14.0
    assert client.post("/api/calculation/wear-records", json=second).status_code == 200

    response = client.get("/api/calculation/wear-records/workbench?line_group=EAL")

    assert response.status_code == 200
    body = response.json()
    assert body["line_class"] == "EAL"
    assert body["history_rows"][0]["values"]["X1"] == 10.0
    assert len(body["detail_records"]["X1"]) == 1


def test_update_and_delete_wire_wear_record_api(client):
    assert client.post("/api/calculation/wear-records", json=_payload()).status_code == 200
    record_id = client.get("/api/calculation/wear-records?line_group=EAL").json()["records"][0]["record_id"]

    update = client.patch(
        f"/api/calculation/wear-records/{record_id}",
        json={"avg_wear_min": 11.1, "wear_percentage": 9.9},
    )
    records_after_update = client.get("/api/calculation/wear-records?line_group=EAL").json()["records"]
    delete = client.delete(f"/api/calculation/wear-records/{record_id}")
    records_after_delete = client.get("/api/calculation/wear-records?line_group=EAL").json()["records"]

    assert update.status_code == 200
    assert records_after_update[0]["avg_wear_min"] == 11.1
    assert records_after_update[0]["wear_percentage"] == 9.9
    assert delete.status_code == 200
    assert records_after_delete == []


def test_export_wire_wear_records_returns_excel(client):
    assert client.post("/api/calculation/wear-records", json=_payload()).status_code == 200

    response = client.get("/api/calculation/wear-records/export?line_group=EAL")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.content[:2] == b"PK"


def test_export_cycle_workbook_endpoint_uses_exact_analysis_filters(client, monkeypatch):
    import app.api.endpoints.wear_records as endpoint

    monkeypatch.setattr(endpoint, "_cycle_metadata_fingerprints", lambda: {"EAL": "hash"}, raising=False)
    monkeypatch.setattr(endpoint, "build_excel_report", lambda *args, **kwargs: b"PKreport", raising=False)
    response = client.get(
        "/api/calculation/wear-records/export.xlsx?line_group=EAL&cycle_date=2026-05-28"
    )

    assert response.status_code == 200
    assert response.content == b"PKreport"
    assert "wire-wear-cycle-EAL-2026-05-28.xlsx" in response.headers["content-disposition"]


def test_manual_add_wire_wear_record_api(client):
    payload = _payload()
    payload["records"][0]["tension_length"] = "X88"

    response = client.post("/api/calculation/wear-records/manual", json=payload)
    records = client.get("/api/calculation/wear-records?line_group=EAL").json()["records"]

    assert response.status_code == 200
    assert response.json()["saved_count"] == 1
    assert records[0]["tension_length"] == "X88"
