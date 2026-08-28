"""Load and seed the built-in one-year database record workbooks.

The packaged workbooks are deliberately read with openpyxl's read-only mode.
All validation happens before the first INSERT so a malformed asset can never
leave a partially populated database behind.
"""

from __future__ import annotations

import hashlib
import json
import logging
import math
import sqlite3
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence, Tuple

logger = logging.getLogger(__name__)

SEED_VERSION = "2025-2026-v2"
LEGACY_SEED_VERSION = "2025-2026-v1"
EXPECTED_RECORD_COUNTS: Mapping[str, int] = {"EAL": 1826, "TML": 974}
LEGACY_EXPECTED_RECORD_COUNTS: Mapping[str, int] = {"EAL": 1732, "TML": 974}
EXPECTED_SECTION_COUNTS: Mapping[str, Mapping[str, int]] = {
    "EAL": {"Mainline": 1584, "RAC": 132, "LOW": 16, "LMC": 94},
    "TML": {"Mainline": 974},
}
QUALITY_MANIFEST_NAME = "database-record-seed-quality.json"
WORKBOOKS: Mapping[str, str] = {
    "EAL": "EAL-1-year-database-record.xlsx",
    "TML": "TML-1-year-database-record.xlsx",
}
EXPECTED_WORKBOOK_SHA256: Mapping[str, str] = {
    "EAL-1-year-database-record.xlsx": (
        "ddc3d385c65e5a2475e8a1070fe8eff401f7967866cef11fc370c0a975f80fad"
    ),
    "TML-1-year-database-record.xlsx": (
        "1d9d989663f9ee67bd60c4bc7d9496a823d263fd3be67e341838d780d1a94165"
    ),
}

EXPECTED_HEADERS: Tuple[str, ...] = (
    "Run Date", "Line", "Track", "Section", "Task Number",
    "Station Start", "Station End", "ID", "FromM", "ToM", "Length",
    "Exception Type", "MaxValue", "MaxLocation", "Overlap",
    "Tension Length", "Track Type", "Level", "Previous 1", "Previous 2",
    "Reoccurrence ID", "Remarks", "ACTION", "CHECK DATE", "CHECKED BY",
    "CHECK RESULT", "VERIFY DEADLINE", "VERIFY DATE", "VERIFY RESULT",
    "VERIFIED BY", "ADJUST DEADLINE", "ADJUST DATE", "ADJUST RESULT",
    "ADJUSTED BY",
)

_HEADER_TO_FIELD: Mapping[str, str] = {
    "Run Date": "date_str", "Line": "line", "Track": "track",
    "Section": "section", "Task Number": "task_no",
    "Station Start": "station_start", "Station End": "station_end",
    "ID": "exception_id", "FromM": "from_m", "ToM": "to_m",
    "Length": "length", "Exception Type": "exception_type",
    "MaxValue": "max_value", "MaxLocation": "max_location",
    "Overlap": "overlap", "Tension Length": "tension_length",
    "Track Type": "track_type", "Level": "level",
    "Previous 1": "previous_1", "Previous 2": "previous_2",
    "Reoccurrence ID": "reoccurrence_id", "Remarks": "remarks",
    "ACTION": "action", "CHECK DATE": "check_date",
    "CHECKED BY": "checked_by", "CHECK RESULT": "check_result",
    "VERIFY DEADLINE": "verify_deadline", "VERIFY DATE": "verify_date",
    "VERIFY RESULT": "verify_result", "VERIFIED BY": "verified_by",
    "ADJUST DEADLINE": "adjust_deadline", "ADJUST DATE": "adjust_date",
    "ADJUST RESULT": "adjust_result", "ADJUSTED BY": "adjusted_by",
}

INSERT_COLUMNS: Tuple[str, ...] = (
    "exception_id", "exception_type", "level", "from_m", "to_m", "length",
    "max_value", "max_location", "track_type", "overlap", "tension_length",
    "landmark", "class", "threshold_value", "section", "previous_1",
    "previous_2", "repeat_count", "reoccurrence_id", "action", "check_date",
    "checked_by", "check_result", "remarks", "verify_deadline", "verify_date",
    "verify_result", "verified_by", "adjust_deadline", "adjust_date",
    "adjust_result", "adjusted_by", "line", "track", "date_str",
    "task_run_date", "task_no", "station_start", "station_end", "saved_by",
)


class DatabaseRecordSeedError(ValueError):
    """Raised when a built-in workbook does not satisfy the import contract."""


@dataclass(frozen=True)
class DatabaseRecordSeedBundle:
    records: Tuple[Dict[str, Any], ...]
    counts_by_line: Dict[str, int]
    counts_by_section: Dict[str, Dict[str, int]]
    output_sha256: Dict[str, str]
    source_sha256: Dict[str, str] = field(default_factory=dict)
    seed_version: str = SEED_VERSION


@dataclass(frozen=True)
class DatabaseRecordSeedResult:
    status: str
    inserted_count: int = 0
    counts_by_line: Dict[str, int] = field(default_factory=dict)
    output_sha256: Dict[str, str] = field(default_factory=dict)
    source_sha256: Dict[str, str] = field(default_factory=dict)
    error: Optional[str] = None


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _cell_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _number(value: Any, field_name: str, row_number: int) -> Optional[float]:
    value = _cell_value(value)
    if value is None:
        return None
    if isinstance(value, bool):
        raise DatabaseRecordSeedError(f"row {row_number}: {field_name} must be numeric")
    try:
        parsed = float(value)
    except (TypeError, ValueError) as exc:
        raise DatabaseRecordSeedError(
            f"row {row_number}: {field_name} must be numeric"
        ) from exc
    return int(parsed) if parsed.is_integer() else parsed


def _date(value: Any, row_number: int) -> str:
    value = _cell_value(value)
    if isinstance(value, datetime):
        compact = value.strftime("%Y%m%d")
    elif isinstance(value, str):
        compact = value.replace("-", "").replace("/", "")
    else:
        compact = str(value) if value is not None else ""
    if len(compact) != 8 or not compact.isdigit():
        raise DatabaseRecordSeedError(f"row {row_number}: Run Date must be YYYYMMDD")
    try:
        datetime.strptime(compact, "%Y%m%d")
    except ValueError as exc:
        raise DatabaseRecordSeedError(f"row {row_number}: invalid Run Date") from exc
    return compact


def _read_workbook(path: Path, expected_line: str) -> Tuple[List[Dict[str, Any]], str]:
    # Keep the normal non-empty database startup path free of the openpyxl
    # import and workbook parsing cost.
    from openpyxl import load_workbook

    if not path.is_file():
        raise FileNotFoundError(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise DatabaseRecordSeedError(f"cannot read workbook {path.name}: {exc}") from exc

    try:
        if workbook.sheetnames != ["Database Records"]:
            raise DatabaseRecordSeedError(
                f"{path.name}: expected only the 'Database Records' sheet"
            )
        worksheet = workbook["Database Records"]
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = tuple(str(value).strip() if value is not None else "" for value in next(rows))
        except StopIteration as exc:
            raise DatabaseRecordSeedError(f"{path.name}: workbook is empty") from exc
        if headers != EXPECTED_HEADERS:
            raise DatabaseRecordSeedError(
                f"{path.name}: unexpected header; expected {list(EXPECTED_HEADERS)}"
            )

        result: List[Dict[str, Any]] = []
        for row_number, values in enumerate(rows, start=2):
            if not any(_cell_value(value) is not None for value in values):
                continue
            values = list(values)
            values.extend([None] * (len(EXPECTED_HEADERS) - len(values)))
            raw = {header: _cell_value(values[index]) for index, header in enumerate(EXPECTED_HEADERS)}
            line = raw["Line"]
            if line != expected_line:
                raise DatabaseRecordSeedError(
                    f"{path.name} row {row_number}: Line must be {expected_line}"
                )
            required = ("ID", "Exception Type", "Track", "Section", "MaxLocation", "FromM", "ToM", "Level")
            missing = [name for name in required if raw[name] is None]
            if missing:
                raise DatabaseRecordSeedError(
                    f"{path.name} row {row_number}: missing {', '.join(missing)}"
                )
            date_str = _date(raw["Run Date"], row_number)
            level = str(raw["Level"])
            if level not in {"L1", "L2", "L3"}:
                raise DatabaseRecordSeedError(f"{path.name} row {row_number}: invalid Level")
            record: Dict[str, Any] = {
                _HEADER_TO_FIELD[header]: raw[header] for header in EXPECTED_HEADERS
            }
            record["date_str"] = date_str
            record["task_run_date"] = date_str
            for field_name in ("from_m", "to_m", "length", "max_value", "max_location"):
                record[field_name] = _number(record[field_name], field_name, row_number)
            record["repeat_count"] = 2
            record["saved_by"] = "built-in-seed"
            result.append(record)
        return result, _sha256(path)
    finally:
        workbook.close()


def _canonical_section(value: Any) -> str:
    normalized = str(value).strip().upper()
    if "LOW" in normalized:
        return "LOW"
    for section in ("MAINLINE", "RAC", "LMC"):
        if section in normalized:
            return section.title() if section == "MAINLINE" else section
    return normalized


def load_database_record_seed(config_dir: Path | str) -> DatabaseRecordSeedBundle:
    """Read and fully validate both standardised workbooks."""
    config_dir = Path(config_dir)
    records: List[Dict[str, Any]] = []
    output_hashes: Dict[str, str] = {}
    counts: Dict[str, int] = {}
    section_counts: Dict[str, Dict[str, int]] = {}
    for line, filename in WORKBOOKS.items():
        rows, output_hash = _read_workbook(config_dir / "database-records" / filename, line)
        records.extend(rows)
        output_hashes[filename] = output_hash
        counts[line] = len(rows)
        line_section_counts: Dict[str, int] = {}
        for row in rows:
            section = _canonical_section(row["section"])
            line_section_counts[section] = line_section_counts.get(section, 0) + 1
        section_counts[line] = line_section_counts

    keys = set()
    for record in records:
        key = (record["exception_id"], record["line"], record["track"], record["date_str"])
        if key in keys:
            raise DatabaseRecordSeedError(f"duplicate composite key: {'|'.join(key)}")
        keys.add(key)

    source_hashes: Dict[str, str] = {}
    manifest_path = config_dir / "database-records" / QUALITY_MANIFEST_NAME
    if not manifest_path.is_file():
        raise DatabaseRecordSeedError(f"missing quality manifest: {QUALITY_MANIFEST_NAME}")
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if manifest.get("schema") != "database-record-seed-quality-v2":
            raise DatabaseRecordSeedError("quality manifest schema mismatch")
        if manifest.get("seed_version") != SEED_VERSION:
            raise DatabaseRecordSeedError("quality manifest seed version mismatch")
        if counts != EXPECTED_RECORD_COUNTS:
            raise DatabaseRecordSeedError(f"unexpected workbook record counts: {counts}")
        if section_counts != EXPECTED_SECTION_COUNTS:
            raise DatabaseRecordSeedError(
                f"unexpected workbook section counts: {section_counts}"
            )
        if manifest.get("source_record_count") != sum(EXPECTED_RECORD_COUNTS.values()):
            raise DatabaseRecordSeedError("quality manifest source count mismatch")
        if manifest.get("retained_by_line") != EXPECTED_RECORD_COUNTS:
            raise DatabaseRecordSeedError("quality manifest line counts mismatch")
        if manifest.get("retained_by_section") != EXPECTED_SECTION_COUNTS:
            raise DatabaseRecordSeedError("quality manifest section counts mismatch")
        if manifest.get("retained_record_count") != sum(EXPECTED_RECORD_COUNTS.values()):
            raise DatabaseRecordSeedError("quality manifest total count mismatch")
        if manifest.get("excluded_record_count") != 0:
            raise DatabaseRecordSeedError("quality manifest excluded count mismatch")
        expected_hashes = manifest.get("output_sha256", {})
        if expected_hashes != EXPECTED_WORKBOOK_SHA256:
            raise DatabaseRecordSeedError("quality manifest output hashes mismatch")
        for filename, output_hash in output_hashes.items():
            if EXPECTED_WORKBOOK_SHA256.get(filename) != output_hash:
                raise DatabaseRecordSeedError(f"unexpected workbook hash: {filename}")
            if expected_hashes.get(filename) != output_hash:
                raise DatabaseRecordSeedError(f"output hash mismatch: {filename}")
        source_hashes = dict(manifest.get("source_sha256", {}))
        if source_hashes != EXPECTED_WORKBOOK_SHA256:
            raise DatabaseRecordSeedError("source hash mismatch")
        provenance = manifest.get("provenance", {})
        if provenance.get("base_seed_version") != LEGACY_SEED_VERSION:
            raise DatabaseRecordSeedError("quality manifest provenance mismatch")
        if not provenance.get("runtime_inputs_are_canonical"):
            raise DatabaseRecordSeedError("quality manifest canonical input mismatch")
        historical_rules = manifest.get("historical_rules")
        historical_exclusions = manifest.get("historical_exclusions")
        if not isinstance(historical_rules, list) or not historical_rules:
            raise DatabaseRecordSeedError("quality manifest historical rules missing")
        if not isinstance(historical_exclusions, list) or sum(
            int(exclusion.get("excluded_count", 0))
            for exclusion in historical_exclusions
        ) != 9:
            raise DatabaseRecordSeedError("quality manifest exclusion history mismatch")
    except DatabaseRecordSeedError:
        raise
    except Exception as exc:
        raise DatabaseRecordSeedError(f"invalid quality manifest: {exc}") from exc

    return DatabaseRecordSeedBundle(
        records=tuple(records), counts_by_line=counts,
        counts_by_section=section_counts,
        output_sha256=output_hashes, source_sha256=source_hashes,
    )


def _metadata(conn: sqlite3.Connection, key: str, value: str, description: str) -> None:
    conn.execute(
        """INSERT INTO system_metadata (key, value, description, updated_at)
           VALUES (?, ?, ?, CURRENT_TIMESTAMP)
           ON CONFLICT(key) DO UPDATE SET value=excluded.value,
             description=excluded.description, updated_at=CURRENT_TIMESTAMP""",
        (key, value, description),
    )


def _seed_metadata(conn: sqlite3.Connection) -> Dict[str, str]:
    return dict(conn.execute(
        """SELECT key, value FROM system_metadata
           WHERE key IN (
             'database_record_seed_version',
             'database_record_seed_record_count',
             'database_record_seed_counts_by_line'
           )"""
    ).fetchall())


def _record_counts_by_line(conn: sqlite3.Connection) -> Dict[str, int]:
    return {
        str(line): int(count)
        for line, count in conn.execute(
            """SELECT line, COUNT(*)
               FROM saved_repeated_exceptions
               GROUP BY line"""
        ).fetchall()
    }


def _record_counts_by_section(
    conn: sqlite3.Connection,
) -> Dict[str, Dict[str, int]]:
    counts: Dict[str, Dict[str, int]] = {}
    for line, section, count in conn.execute(
        """SELECT line, section, COUNT(*)
           FROM saved_repeated_exceptions
           GROUP BY line, section"""
    ).fetchall():
        line_counts = counts.setdefault(str(line), {})
        canonical_section = _canonical_section(section)
        line_counts[canonical_section] = (
            line_counts.get(canonical_section, 0) + int(count)
        )
    return counts


def _is_legacy_seed_eligible(
    metadata: Mapping[str, str],
    actual_counts: Mapping[str, int],
) -> bool:
    if metadata.get("database_record_seed_version") != LEGACY_SEED_VERSION:
        return False
    if metadata.get("database_record_seed_record_count") != str(
        sum(LEGACY_EXPECTED_RECORD_COUNTS.values())
    ):
        return False
    try:
        metadata_counts = json.loads(
            metadata.get("database_record_seed_counts_by_line", "")
        )
    except (TypeError, ValueError):
        return False
    return (
        metadata_counts == LEGACY_EXPECTED_RECORD_COUNTS
        and actual_counts == LEGACY_EXPECTED_RECORD_COUNTS
    )


def _write_seed_metadata(
    conn: sqlite3.Connection,
    bundle: DatabaseRecordSeedBundle,
    *,
    status: str,
    replacement_source_version: Optional[str],
) -> None:
    _metadata(
        conn,
        "database_record_seed_version",
        bundle.seed_version,
        "Built-in database record seed version",
    )
    _metadata(
        conn,
        "database_record_seed_record_count",
        str(len(bundle.records)),
        "Built-in database record seed row count",
    )
    _metadata(
        conn,
        "database_record_seed_counts_by_line",
        json.dumps(bundle.counts_by_line, sort_keys=True),
        "Built-in database record seed counts",
    )
    _metadata(
        conn,
        "database_record_seed_output_sha256",
        json.dumps(bundle.output_sha256, sort_keys=True),
        "Built-in database record workbook hashes",
    )
    _metadata(
        conn,
        "database_record_seed_source_sha256",
        json.dumps(bundle.source_sha256, sort_keys=True),
        "Canonical database record workbook hashes",
    )
    _metadata(
        conn,
        "database_record_seed_status",
        status,
        "Built-in database record seed completion status",
    )
    if replacement_source_version:
        _metadata(
            conn,
            "database_record_seed_replacement_source_version",
            replacement_source_version,
            "Database record seed version replaced during upgrade",
        )
    else:
        conn.execute(
            """DELETE FROM system_metadata
               WHERE key = 'database_record_seed_replacement_source_version'"""
        )


def _validate_applied_seed(
    conn: sqlite3.Connection,
    bundle: DatabaseRecordSeedBundle,
) -> None:
    actual_counts = _record_counts_by_line(conn)
    if actual_counts != bundle.counts_by_line:
        raise DatabaseRecordSeedError(
            f"applied database record counts mismatch: {actual_counts}"
        )
    actual_section_counts = _record_counts_by_section(conn)
    if actual_section_counts != bundle.counts_by_section:
        raise DatabaseRecordSeedError(
            f"applied database record section counts mismatch: {actual_section_counts}"
        )
    foreign_key_errors = conn.execute("PRAGMA foreign_key_check").fetchall()
    if foreign_key_errors:
        raise DatabaseRecordSeedError(
            f"foreign key check failed after seed: {len(foreign_key_errors)} violation(s)"
        )


def seed_database_records(conn: sqlite3.Connection, config_dir: Path | str) -> DatabaseRecordSeedResult:
    """Apply v2 to an empty DB or atomically replace the exact v1 baseline."""
    try:
        actual_counts = _record_counts_by_line(conn)
        existing = sum(actual_counts.values())
        metadata = _seed_metadata(conn)
        if existing and metadata.get("database_record_seed_version") == SEED_VERSION:
            return DatabaseRecordSeedResult(status="skipped_up_to_date", inserted_count=0)

        replacement_source_version: Optional[str] = None
        if existing and not _is_legacy_seed_eligible(metadata, actual_counts):
            return DatabaseRecordSeedResult(status="skipped_non_empty", inserted_count=0)
        if existing:
            replacement_source_version = LEGACY_SEED_VERSION

        # Load, hash, validate counts/sections and materialize rows before the
        # first destructive statement in a legacy replacement.
        bundle = load_database_record_seed(config_dir)
        placeholders = ", ".join("?" for _ in INSERT_COLUMNS)
        query = (
            f"INSERT INTO saved_repeated_exceptions ({', '.join(INSERT_COLUMNS)}) "
            f"VALUES ({placeholders})"
        )
        parameters = [
            tuple(record.get(column) for column in INSERT_COLUMNS)
            for record in bundle.records
        ]
        status = "replaced" if replacement_source_version else "seeded"
        conn.execute("SAVEPOINT database_record_seed")
        try:
            if replacement_source_version:
                conn.execute("DELETE FROM saved_repeated_exceptions")
            conn.executemany(query, parameters)
            _write_seed_metadata(
                conn,
                bundle,
                status=status,
                replacement_source_version=replacement_source_version,
            )
            _validate_applied_seed(conn, bundle)
            conn.execute("RELEASE SAVEPOINT database_record_seed")
        except Exception:
            conn.execute("ROLLBACK TO SAVEPOINT database_record_seed")
            conn.execute("RELEASE SAVEPOINT database_record_seed")
            raise
        logger.info(
            "Database record seed %s with %s records",
            status,
            len(bundle.records),
        )
        return DatabaseRecordSeedResult(
            status=status, inserted_count=len(bundle.records),
            counts_by_line=bundle.counts_by_line,
            output_sha256=bundle.output_sha256,
            source_sha256=bundle.source_sha256,
        )
    except Exception as exc:
        logger.error("Database record seed skipped: %s", exc)
        return DatabaseRecordSeedResult(status="failed", error=str(exc))
