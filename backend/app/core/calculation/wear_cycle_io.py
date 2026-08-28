"""Report export and two-phase JSON synchronization for normalized wear cycles."""
from __future__ import annotations

from copy import copy
from contextlib import contextmanager
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
import math
from pathlib import Path
import sqlite3
from typing import Iterable, Iterator, Mapping, Sequence
from uuid import uuid4

from openpyxl import Workbook

from .wear_calculator import calculate_wear_percentage
from .wear_cycle_application import create_sqlite_backup
from .wear_cycle_metadata import (
    MetadataResolutionError,
    MetadataValidationError,
    normalize_cycle_date,
    normalize_line_identity,
    normalize_line_group,
    normalize_tension_length,
    resolve_canonical_tl,
)
from .wear_cycle_repository import StaleDataVersionError, bump_data_version, current_data_version
from .wear_cycle_types import EXPECTED_SEGMENTS, MetadataInterval
from .wear_tl_scope import TensionLengthScope, classify_tension_length_scope


SYNC_SCHEMA = "wear-cycle-v1"
REPORT_SCHEMA = "wear-cycle-report-v1"
SyncMetadata = Sequence[MetadataInterval] | Mapping[object, Sequence[MetadataInterval]]
_PACKAGE_FIELDS = {
    "schema", "package_id", "exported_at", "source_workstation", "metadata_fingerprint",
    "cycles", "records", "segments", "conflict_decisions", "tombstones",
}


class SyncValidationError(ValueError):
    """Raised when a sync package cannot be interpreted safely."""


class SyncPreviewMismatchError(ValueError):
    """Raised when Apply is not using the preview that was approved."""


_ACTION_STATUS = {
    "create": "new",
    "update": "update",
    "unchanged": "no_change",
    "conflict": "conflict",
    "error": "error",
    "delete": "delete",
}


@dataclass(frozen=True)
class SyncAction:
    key: tuple[str, str, str, str]
    action: str
    detail: str = ""
    record: dict | None = None
    tombstone: dict | None = None
    local_record: dict | None = None
    local_tombstone: dict | None = None
    status: str = ""

    def __post_init__(self):
        if not self.status:
            object.__setattr__(self, "status", _ACTION_STATUS.get(self.action, self.action))


@dataclass(frozen=True)
class SyncPreview:
    source_package: dict
    actions: tuple[SyncAction, ...]
    expected_data_version: int
    preview_digest: str
    warnings: tuple[str, ...] = ()


def _timestamp(value: object) -> datetime:
    if not isinstance(value, str) or not value.strip():
        raise SyncValidationError("timestamp cannot be empty")
    try:
        parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
    except ValueError as exc:
        raise SyncValidationError(f"invalid timestamp: {value}") from exc
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.astimezone(timezone.utc)


def _timestamp_text(value: object) -> str:
    return _timestamp(value).isoformat(timespec="microseconds")


def _json(value: object) -> str:
    return json.dumps(value, ensure_ascii=True, sort_keys=True, separators=(",", ":"))


def _decode(value: object, default: object) -> object:
    if isinstance(value, str):
        try:
            return json.loads(value)
        except json.JSONDecodeError:
            return default
    return default if value is None else value


def _record_key(value: Mapping[str, object]) -> tuple[str, str, str, str]:
    if not isinstance(value.get("line_class"), str) or not str(value["line_class"]).strip():
        raise SyncValidationError("line_class is required")
    line_group, line_class = normalize_line_identity(
        value.get("line_group"), value.get("line_class")
    )
    return (
        line_group,
        line_class,
        normalize_cycle_date(value.get("cycle_date")).isoformat(),
        normalize_tension_length(value.get("tension_length")),
    )


def _row_record(row: sqlite3.Row) -> dict:
    item = dict(row)
    item.pop("record_id", None)
    item.pop("cycle_id", None)
    item["source_lineage"] = _decode(item.get("source_lineage"), [])
    intervals = _decode(item.get("physical_intervals"), [])
    if not intervals:
        intervals = [{
            "track": item.get("track"),
            "from_m": item.get("from_m"),
            "to_m": item.get("to_m"),
        }]
    item["physical_intervals"] = intervals
    item["interval_count"] = len(intervals)
    return item


def _rows(conn: sqlite3.Connection, sql: str, params: Iterable[object] = ()) -> list[dict]:
    return [dict(row) for row in conn.execute(sql, tuple(params)).fetchall()]


def _metadata_fingerprint_text(metadata_fingerprint: Mapping[str, str] | None) -> str:
    return _json(dict(sorted((metadata_fingerprint or {}).items())))


def build_excel_report(
    conn: sqlite3.Connection,
    *,
    line_group: str | None = None,
    line_class: str | None = None,
    cycle_date: str | None = None,
    metadata_fingerprint: Mapping[str, str] | None = None,
    app_version: str = "2.0.0",
    exported_at: datetime | None = None,
) -> bytes:
    """Build the report-only workbook from normalized committed cycle tables."""
    clauses: list[str] = []
    params: list[object] = []
    if line_group is not None:
        line, normalized_class = normalize_line_identity(line_group, line_class)
        clauses.extend(("records.line_group = ?", "records.line_class = ?"))
        params.extend((line, normalized_class))
    elif line_class is not None:
        raise MetadataValidationError("line_class requires line_group")
    if cycle_date is not None:
        clauses.append("records.cycle_date = ?")
        params.append(normalize_cycle_date(cycle_date).isoformat())
    where = " WHERE " + " AND ".join(clauses) if clauses else ""
    records = _rows(conn, f"""
        SELECT records.* FROM wire_wear_cycle_records AS records
        {where}
        ORDER BY records.cycle_date, records.from_m, records.to_m, records.tension_length
    """, params)
    records = [
        record
        for record in records
        if classify_tension_length_scope(
            record["line_group"], record["tension_length"]
        )
        is TensionLengthScope.MAINLINE
    ]
    cycle_ids = sorted({int(row["cycle_id"]) for row in records})
    coverage = _rows(conn, """
        SELECT cycles.line_group, cycles.line_class, cycles.cycle_date, segments.*
        FROM wire_wear_cycle_segments AS segments
        JOIN wire_wear_cycles AS cycles ON cycles.cycle_id = segments.cycle_id
        WHERE segments.cycle_id IN (%s)
        ORDER BY cycles.cycle_date, cycles.line_group, cycles.line_class, segments.segment_name
    """ % ",".join("?" for _ in cycle_ids), cycle_ids) if cycle_ids else []
    decisions = _rows(conn, """
        SELECT cycles.line_group, cycles.line_class, cycles.cycle_date, decisions.*
        FROM wire_wear_conflict_decisions AS decisions
        JOIN wire_wear_cycles AS cycles ON cycles.cycle_id = decisions.cycle_id
        WHERE decisions.cycle_id IN (%s)
        ORDER BY cycles.cycle_date, decisions.measurement_identity
    """ % ",".join("?" for _ in cycle_ids), cycle_ids) if cycle_ids else []

    workbook = Workbook()
    wear = workbook.active
    wear.title = "Wear Records"
    wear.append([
        "Cycle Date", "Line", "Class", "Track", "Tension Length", "From (m)", "To (m)",
        "Interval Count", "Physical Intervals", "Avg Wear Min", "Wear %", "Measurement SD",
    ])
    for record in records:
        intervals = _decode(record.get("physical_intervals"), []) or [{
            "track": record["track"], "from_m": record["from_m"], "to_m": record["to_m"],
        }]
        wear.append([
            record["cycle_date"], record["line_group"], record["line_class"],
            record["track"], record["tension_length"],
            record["from_m"], record["to_m"], len(intervals),
            "; ".join(
                f"{item['track']} {item['from_m']}-{item['to_m']}" for item in intervals
            ),
            record["avg_wear_min"], record["wear_percentage"],
            record["measurement_sd"],
        ])

    coverage_sheet = workbook.create_sheet("Cycle Coverage")
    coverage_sheet.append([
        "Cycle Date", "Line", "Class", "Segment", "Present", "Coverage %", "Diagnostic Gaps",
        "Source Files", "Acquisition From", "Acquisition To",
    ])
    seen_coverage = {
        (row["line_group"], row["line_class"], row["cycle_date"], row["segment_name"])
        for row in coverage
    }
    for row in coverage:
        diagnostic_gaps = [
            gap
            for gap in _decode(row["diagnostic_gaps"], [])
            if gap in {"segment_missing", "metadata_denominator_empty"}
            or classify_tension_length_scope(row["line_group"], gap)
            is TensionLengthScope.MAINLINE
        ]
        coverage_sheet.append([
            row["cycle_date"], row["line_group"], row["line_class"],
            row["segment_name"], bool(row["is_present"]),
            row["coverage_percentage"], "; ".join(diagnostic_gaps),
            "; ".join(_decode(row["source_file_names"], [])), row["acquisition_date_from"], row["acquisition_date_to"],
        ])
    for record in records:
        for segment in EXPECTED_SEGMENTS[record["line_group"]]:
            identity = (
                record["line_group"], record["line_class"], record["cycle_date"], segment
            )
            if identity not in seen_coverage:
                coverage_sheet.append([
                    record["cycle_date"], record["line_group"], record["line_class"],
                    segment, False, 0.0, "", "", None, None,
                ])
                seen_coverage.add(identity)

    audit = workbook.create_sheet("Conflict Audit")
    audit.append([
        "Cycle Date", "Line", "Class", "Source File", "Source Value",
        "Selected Lower Value", "Measurement Identity", "Accepted Time",
    ])
    for decision in decisions:
        for source_file, source_value in _decode(decision["source_values"], []):
            audit.append([
                decision["cycle_date"], decision["line_group"], decision["line_class"],
                source_file, source_value, decision["selected_wear_min"],
                decision["measurement_identity"], decision["accepted_at"],
            ])

    info = workbook.create_sheet("Workbook Info")
    info.append(["Field", "Value"])
    now = exported_at or datetime.now(timezone.utc)
    info.append(["Schema", REPORT_SCHEMA])
    info.append(["Exported At", now.astimezone(timezone.utc).isoformat(timespec="microseconds")])
    info.append(["App Version", app_version])
    info.append(["Metadata Fingerprint", _metadata_fingerprint_text(metadata_fingerprint)])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            font = copy(cell.font)
            font.bold = True
            cell.font = font
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_sync_package(
    conn: sqlite3.Connection,
    *,
    source_workstation: str = "TOV640 Analyzer",
    metadata_fingerprint: Mapping[str, str] | None = None,
    package_id: str | None = None,
    exported_at: datetime | None = None,
) -> dict:
    """Serialize all normalized cycle state into the stable wear-cycle-v1 schema."""
    cycles = _rows(
        conn,
        "SELECT * FROM wire_wear_cycles ORDER BY line_group, line_class, cycle_date",
    )
    cycle_lookup = {
        int(row["cycle_id"]): (row["line_group"], row["line_class"], row["cycle_date"])
        for row in cycles
    }
    for row in cycles:
        row.pop("cycle_id", None)
        row["source_lineage"] = _decode(row.get("source_lineage"), [])
    records = [_row_record(row) for row in conn.execute(
        """
        SELECT * FROM wire_wear_cycle_records
        ORDER BY line_group, line_class, cycle_date, tension_length
        """
    ).fetchall()]
    segments = _rows(conn, "SELECT * FROM wire_wear_cycle_segments ORDER BY cycle_id, segment_name")
    for row in segments:
        line, line_class, cycle = cycle_lookup.get(
            int(row.pop("cycle_id")), (None, None, None)
        )
        row.pop("segment_id", None)
        row["line_group"], row["line_class"], row["cycle_date"] = (
            line,
            line_class,
            cycle,
        )
        row["diagnostic_gaps"] = _decode(row.get("diagnostic_gaps"), [])
        row["source_file_names"] = _decode(row.get("source_file_names"), [])
    decisions = _rows(conn, "SELECT * FROM wire_wear_conflict_decisions ORDER BY cycle_id, measurement_identity")
    for row in decisions:
        line, line_class, cycle = cycle_lookup.get(
            int(row.pop("cycle_id")), (None, None, None)
        )
        row.pop("decision_id", None)
        row["line_group"], row["line_class"], row["cycle_date"] = (
            line,
            line_class,
            cycle,
        )
        row["source_values"] = _decode(row.get("source_values"), [])
    tombstones = _rows(
        conn,
        """
        SELECT * FROM wire_wear_deletion_tombstones
        ORDER BY line_group, line_class, cycle_date, tension_length
        """,
    )
    for row in tombstones:
        row.pop("tombstone_id", None)
    now = exported_at or datetime.now(timezone.utc)
    return {
        "schema": SYNC_SCHEMA,
        "package_id": package_id or str(uuid4()),
        "exported_at": now.astimezone(timezone.utc).isoformat(timespec="microseconds"),
        "source_workstation": str(source_workstation),
        "metadata_fingerprint": dict(sorted((metadata_fingerprint or {}).items())),
        "cycles": cycles,
        "records": records,
        "segments": segments,
        "conflict_decisions": decisions,
        "tombstones": tombstones,
    }


def _validate_package(package: object) -> dict:
    if not isinstance(package, dict):
        raise SyncValidationError("sync package must be an object")
    if package.get("schema") != SYNC_SCHEMA:
        raise SyncValidationError(f"unsupported sync schema; expected {SYNC_SCHEMA}")
    if set(package) != _PACKAGE_FIELDS:
        raise SyncValidationError("sync package does not match wear-cycle-v1")
    for name in ("package_id", "exported_at", "source_workstation"):
        if not isinstance(package.get(name), str) or not package[name].strip():
            raise SyncValidationError(f"{name} is required")
    _timestamp(package["exported_at"])
    if not isinstance(package["metadata_fingerprint"], dict):
        raise SyncValidationError("metadata_fingerprint must be an object")
    for name in ("cycles", "records", "segments", "conflict_decisions", "tombstones"):
        if not isinstance(package[name], list) or not all(isinstance(value, dict) for value in package[name]):
            raise SyncValidationError(f"{name} must be an array of objects")
    try:
        _validate_cycle_payload(package)
    except SyncValidationError:
        raise
    except (KeyError, MetadataValidationError, TypeError, ValueError) as exc:
        raise SyncValidationError(f"invalid cycle payload: {exc}") from exc
    return package


def _package_cycle_key(row: Mapping[str, object]) -> tuple[str, str, str]:
    try:
        if not isinstance(row.get("line_class"), str) or not str(row["line_class"]).strip():
            raise SyncValidationError("line_class is required")
        line_group, line_class = normalize_line_identity(
            row.get("line_group"), row.get("line_class")
        )
        cycle_date = normalize_cycle_date(row.get("cycle_date")).isoformat()
    except MetadataValidationError as exc:
        raise SyncValidationError(str(exc)) from exc
    if (
        row.get("line_group") != line_group
        or row.get("line_class") != line_class
        or row.get("cycle_date") != cycle_date
    ):
        raise SyncValidationError(
            "cycle identity must use canonical line_group, line_class, and cycle_date"
        )
    return line_group, line_class, cycle_date


def _validate_cycle_payload(package: Mapping[str, object]) -> None:
    cycle_keys: set[tuple[str, str, str]] = set()
    for cycle in package["cycles"]:
        key = _package_cycle_key(cycle)
        if key in cycle_keys:
            raise SyncValidationError("duplicate package cycle business key")
        cycle_keys.add(key)
        if cycle.get("source_type") not in {"analysis", "manual", "sync"}:
            raise SyncValidationError("cycle source_type is invalid")
        completeness = cycle.get("completeness_state")
        if not isinstance(completeness, str) or completeness not in {"complete", "incomplete"}:
            raise SyncValidationError("cycle completeness_state must be complete or incomplete")
        if not isinstance(cycle.get("source_lineage"), list):
            raise SyncValidationError("cycle source_lineage must be an array")
        _timestamp(cycle.get("created_at"))
        _timestamp(cycle.get("updated_at"))
        for field in ("acquisition_date_from", "acquisition_date_to"):
            if cycle.get(field) is not None:
                normalize_cycle_date(cycle[field])

    for record in package["records"]:
        if _package_cycle_key(record) not in cycle_keys:
            raise SyncValidationError("record must reference a package cycle")

    segment_keys: set[tuple[str, str, str, str]] = set()
    for segment in package["segments"]:
        cycle_key = _package_cycle_key(segment)
        name = segment.get("segment_name")
        if cycle_key not in cycle_keys or not isinstance(name, str) or not name.strip():
            raise SyncValidationError("segment must reference a package cycle and have a name")
        if name.strip() not in EXPECTED_SEGMENTS[cycle_key[0]]:
            raise SyncValidationError("segment name is not expected for its line")
        key = (*cycle_key, name.strip())
        if key in segment_keys:
            raise SyncValidationError("duplicate package segment business key")
        segment_keys.add(key)
        is_present = segment.get("is_present")
        if not isinstance(is_present, (bool, int)) or is_present not in {True, False, 0, 1}:
            raise SyncValidationError("segment is_present must be boolean")
        coverage = float(segment.get("coverage_percentage"))
        if not math.isfinite(coverage) or not 0.0 <= coverage <= 100.0:
            raise SyncValidationError("segment coverage_percentage must be between 0 and 100")
        for field in ("diagnostic_gaps", "source_file_names"):
            if not isinstance(segment.get(field), list):
                raise SyncValidationError(f"segment {field} must be an array")
        for field in ("acquisition_date_from", "acquisition_date_to"):
            if segment.get(field) is not None:
                normalize_cycle_date(segment[field])

    decision_keys: set[tuple[str, str, str, str]] = set()
    for decision in package["conflict_decisions"]:
        cycle_key = _package_cycle_key(decision)
        identity = decision.get("measurement_identity")
        if cycle_key not in cycle_keys or not isinstance(identity, str) or not identity.strip():
            raise SyncValidationError("conflict decision must reference a package cycle and have an identity")
        key = (*cycle_key, identity.strip())
        if key in decision_keys:
            raise SyncValidationError("duplicate package conflict decision business key")
        decision_keys.add(key)
        source_values = decision.get("source_values")
        if not isinstance(source_values, list) or not all(
            isinstance(value, (list, tuple)) and len(value) == 2
            and isinstance(value[0], str) and value[0].strip()
            and math.isfinite(float(value[1]))
            for value in source_values
        ):
            raise SyncValidationError("conflict decision source_values must contain file/value pairs")
        selected = float(decision.get("selected_wear_min"))
        if not math.isfinite(selected):
            raise SyncValidationError("conflict selected_wear_min must be finite")
        _timestamp(decision.get("accepted_at"))

    tombstone_keys: set[tuple[str, str, str, str]] = set()
    for tombstone in package["tombstones"]:
        key = _record_key(tombstone)
        if key in tombstone_keys:
            raise SyncValidationError("duplicate package tombstone business key")
        tombstone_keys.add(key)
        _timestamp(tombstone.get("deleted_at"))
        source_package_id = tombstone.get("source_package_id")
        if source_package_id is not None and (
            not isinstance(source_package_id, str) or not source_package_id.strip()
        ):
            raise SyncValidationError("tombstone source_package_id must be text")


def _canonical_package(package: dict) -> dict:
    result = json.loads(_json(package))
    for collection in ("cycles", "records", "segments", "conflict_decisions", "tombstones"):
        result[collection] = sorted(result[collection], key=lambda item: _json(item))
    return result


def _metadata_geometry(record: dict, canonical: object) -> bool:
    record_intervals = record.get("physical_intervals") or [{
        "track": record.get("track"),
        "from_m": record.get("from_m"),
        "to_m": record.get("to_m"),
    }]
    canonical_intervals = [
        {
            "track": interval.track,
            "from_m": float(interval.from_m),
            "to_m": float(interval.to_m),
        }
        for interval in canonical.intervals
    ] or [{
        "track": canonical.track,
        "from_m": float(canonical.from_m),
        "to_m": float(canonical.to_m),
    }]
    return (
        str(record.get("track")) == canonical.track
        and float(record.get("from_m")) == float(canonical.from_m)
        and float(record.get("to_m")) == float(canonical.to_m)
        and [
            {
                "track": str(interval.get("track")),
                "from_m": float(interval.get("from_m")),
                "to_m": float(interval.get("to_m")),
            }
            for interval in record_intervals
        ] == canonical_intervals
    )


def _action_digest(source_package: dict, actions: Sequence[SyncAction], version: int) -> str:
    payload = {
        "source_package": source_package,
        "actions": [asdict(action) for action in actions],
        "expected_data_version": version,
    }
    return sha256(_json(payload).encode("utf-8")).hexdigest()


def _metadata_for_identity(
    metadata: SyncMetadata,
    line_group: str,
    line_class: str,
) -> Sequence[MetadataInterval]:
    line, normalized_class = normalize_line_identity(line_group, line_class)
    if not isinstance(metadata, Mapping):
        return metadata
    intervals = (
        metadata.get((line, normalized_class))
        or metadata.get(f"{line}/{normalized_class}")
        or metadata.get(normalized_class)
        or (metadata.get(line) if line == normalized_class else None)
    )
    if not intervals:
        raise MetadataResolutionError(
            f"no canonical metadata for line identity {line}/{normalized_class}"
        )
    return intervals


def preview_sync_import(
    conn: sqlite3.Connection,
    package: object,
    metadata: SyncMetadata,
    *,
    metadata_fingerprint: Mapping[str, str] | None = None,
) -> SyncPreview:
    """Validate an entire sync package and determine deterministic per-key actions."""
    source = _canonical_package(_validate_package(package))
    local_fingerprint = dict(metadata_fingerprint or {})
    warnings: list[str] = []
    if source["metadata_fingerprint"] != local_fingerprint:
        warnings.append("metadata fingerprint differs; local canonical metadata is authoritative")
    records: dict[tuple[str, str, str, str], dict] = {}
    tombstones: dict[tuple[str, str, str, str], dict] = {}
    errors: dict[tuple[str, str, str, str], str] = {}
    # A historical seed contains many rows sharing the same canonical
    # tension-length geometry. Resolve each identity once instead of scanning
    # the full metadata workbook for every record.
    canonical_cache: dict[tuple[str, str, str], object] = {}
    for record in source["records"]:
        try:
            key = _record_key(record)
            _timestamp(record.get("updated_at"))
            average = float(record.get("avg_wear_min"))
            if not math.isfinite(average):
                raise SyncValidationError("average wear must be finite")
            if key in records:
                errors[key] = "duplicate package record business key"
                continue
            canonical_key = (key[0], key[1], key[3])
            canonical = canonical_cache.get(canonical_key)
            if canonical is None:
                canonical = resolve_canonical_tl(
                    key[0],
                    key[3],
                    _metadata_for_identity(metadata, key[0], key[1]),
                    line_class=key[1],
                )
                canonical_cache[canonical_key] = canonical
            if source["metadata_fingerprint"] != local_fingerprint and not _metadata_geometry(record, canonical):
                errors[key] = "metadata fingerprint differs and canonical geometry changed"
                continue
            physical_intervals = [
                {"track": item.track, "from_m": float(item.from_m), "to_m": float(item.to_m)}
                for item in canonical.intervals
            ]
            normalized = dict(record)
            normalized.update({
                "line_group": key[0], "line_class": key[1],
                "cycle_date": key[2], "tension_length": key[3],
                "track": canonical.track, "from_m": float(canonical.from_m), "to_m": float(canonical.to_m),
                "avg_wear_min": average, "wear_percentage": calculate_wear_percentage(average),
                "updated_at": _timestamp_text(record["updated_at"]),
                "created_at": _timestamp_text(record.get("created_at") or record["updated_at"]),
                "source_lineage": list(record.get("source_lineage") or []),
                "physical_intervals": physical_intervals,
                "interval_count": len(physical_intervals),
            })
            records[key] = normalized
        except (MetadataResolutionError, MetadataValidationError, SyncValidationError, TypeError, ValueError) as exc:
            key = (
                str(record.get("line_group", "?")),
                str(record.get("line_class", "?")),
                str(record.get("cycle_date", "?")),
                str(record.get("tension_length", "?")),
            )
            errors[key] = f"unknown or invalid tension length: {exc}"
    for tombstone in source["tombstones"]:
        try:
            key = _record_key(tombstone)
            _timestamp(tombstone.get("deleted_at"))
            if key in tombstones:
                errors[key] = "duplicate package tombstone business key"
                continue
            normalized = dict(tombstone)
            normalized.update({
                "line_group": key[0],
                "line_class": key[1],
                "cycle_date": key[2],
                "tension_length": key[3],
                "deleted_at": _timestamp_text(tombstone["deleted_at"]),
            })
            tombstones[key] = normalized
        except (MetadataValidationError, SyncValidationError) as exc:
            errors[("?", "?", "?", str(tombstone.get("tension_length", "?")))] = str(exc)

    local_records = {
        _record_key(dict(row)): _row_record(row)
        for row in conn.execute("SELECT * FROM wire_wear_cycle_records").fetchall()
    }
    local_tombstones = {
        _record_key(dict(row)): dict(row)
        for row in conn.execute("SELECT * FROM wire_wear_deletion_tombstones").fetchall()
    }
    actions: list[SyncAction] = []
    for key in sorted(set(records) | set(tombstones) | set(errors)):
        if key in errors:
            actions.append(SyncAction(key, "error", errors[key]))
            continue
        remote_record, remote_tombstone = records.get(key), tombstones.get(key)
        remote_delete = remote_tombstone is not None and (
            remote_record is None or _timestamp(remote_tombstone["deleted_at"]) >= _timestamp(remote_record["updated_at"])
        )
        remote_time = _timestamp(remote_tombstone["deleted_at"] if remote_delete else remote_record["updated_at"])
        local_record, local_tombstone = local_records.get(key), local_tombstones.get(key)
        local_delete = local_tombstone is not None and (
            local_record is None or _timestamp(local_tombstone["deleted_at"]) >= _timestamp(local_record["updated_at"])
        )
        local_time = None if local_record is None and local_tombstone is None else _timestamp(
            local_tombstone["deleted_at"] if local_delete else local_record["updated_at"]
        )
        if local_time is not None and remote_time < local_time:
            actions.append(SyncAction(
                key, "unchanged", "local change is newer", remote_record, remote_tombstone,
                local_record, local_tombstone, "keep_local",
            ))
        elif local_time is not None and remote_time == local_time and remote_delete != local_delete:
            actions.append(SyncAction(
                key, "delete" if remote_delete else "unchanged",
                "tombstone wins equal timestamp", remote_record, remote_tombstone,
                local_record, local_tombstone,
                "delete" if remote_delete else "keep_local",
            ))
        elif remote_delete:
            actions.append(SyncAction(
                key, "delete" if not local_delete else "unchanged", "",
                remote_record, remote_tombstone, local_record, local_tombstone,
                "delete" if not local_delete else "no_change",
            ))
        elif local_record is None or local_delete:
            actions.append(SyncAction(
                key, "create", "", remote_record, remote_tombstone,
                local_record, local_tombstone,
            ))
        elif remote_time == local_time and _json(remote_record) != _json(local_record):
            actions.append(SyncAction(
                key, "conflict", "records differ at the same timestamp",
                remote_record, remote_tombstone, local_record, local_tombstone,
            ))
        elif remote_time > local_time:
            actions.append(SyncAction(
                key, "update", "", remote_record, remote_tombstone,
                local_record, local_tombstone,
            ))
        else:
            actions.append(SyncAction(
                key, "unchanged", "", remote_record, remote_tombstone,
                local_record, local_tombstone,
            ))
    version = current_data_version(conn)
    digest = _action_digest(source, actions, version)
    return SyncPreview(source, tuple(actions), version, digest, tuple(warnings))


@contextmanager
def _transaction(conn: sqlite3.Connection) -> Iterator[None]:
    if conn.in_transaction:
        raise SyncValidationError("sync requires an idle database connection")
    conn.execute("BEGIN IMMEDIATE")
    try:
        yield
    except Exception:
        conn.rollback()
        raise
    else:
        conn.commit()


def _ensure_parent(
    conn: sqlite3.Connection,
    record: dict,
    cycles: Mapping[tuple[str, str, str], dict],
) -> int:
    key = (record["line_group"], record["line_class"], record["cycle_date"])
    incoming = cycles.get(key, {})
    now = record["updated_at"]
    conn.execute("""
        INSERT INTO wire_wear_cycles (
            line_group, line_class, cycle_date, source_type,
            acquisition_date_from, acquisition_date_to,
            completeness_state, source_lineage, created_at, updated_at
        ) VALUES (?, ?, ?, 'sync', ?, ?, ?, ?, ?, ?)
        ON CONFLICT(line_group, line_class, cycle_date) DO UPDATE SET
            source_type = excluded.source_type,
            acquisition_date_from = excluded.acquisition_date_from,
            acquisition_date_to = excluded.acquisition_date_to,
            completeness_state = excluded.completeness_state,
            source_lineage = excluded.source_lineage,
            updated_at = excluded.updated_at
    """, (
        *key, incoming.get("acquisition_date_from"), incoming.get("acquisition_date_to"),
        incoming.get("completeness_state", "incomplete"), _json(incoming.get("source_lineage", [])),
        incoming.get("created_at", now), now,
    ))
    return int(conn.execute(
        """
        SELECT cycle_id FROM wire_wear_cycles
        WHERE line_group = ? AND line_class = ? AND cycle_date = ?
        """,
        key,
    ).fetchone()[0])


def _apply_action(
    conn: sqlite3.Connection,
    action: SyncAction,
    cycles: Mapping[tuple[str, str, str], dict],
    package_id: str,
) -> None:
    if action.action in {"unchanged", "error", "conflict"}:
        return
    key = action.key
    if action.action == "delete":
        conn.execute(
            """
            DELETE FROM wire_wear_cycle_records
            WHERE line_group = ? AND line_class = ?
              AND cycle_date = ? AND tension_length = ?
            """,
            key,
        )
        deleted_at = action.tombstone["deleted_at"] if action.tombstone else action.record["updated_at"]
        conn.execute("""
            INSERT INTO wire_wear_deletion_tombstones (
                line_group, line_class, cycle_date, tension_length,
                deleted_at, source_package_id
            ) VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(line_group, line_class, cycle_date, tension_length) DO UPDATE SET
                deleted_at = excluded.deleted_at, source_package_id = excluded.source_package_id
        """, (*key, deleted_at, package_id))
        conn.execute(
            """
            DELETE FROM wire_wear_cycles
            WHERE line_group = ? AND line_class = ? AND cycle_date = ?
              AND NOT EXISTS (
                  SELECT 1 FROM wire_wear_cycle_records
                  WHERE line_group = ? AND line_class = ? AND cycle_date = ?
              )
            """,
            (*key[:3], *key[:3]),
        )
        return
    record = action.record
    cycle_id = _ensure_parent(conn, record, cycles)
    conn.execute("""
        INSERT INTO wire_wear_cycle_records (
            cycle_id, line_group, line_class, cycle_date, tension_length,
            track, from_m, to_m,
            avg_wear_min, wear_percentage, measurement_sd, physical_intervals,
            source_lineage, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(line_group, line_class, cycle_date, tension_length) DO UPDATE SET
            cycle_id = excluded.cycle_id, track = excluded.track, from_m = excluded.from_m,
            to_m = excluded.to_m, avg_wear_min = excluded.avg_wear_min,
            wear_percentage = excluded.wear_percentage, measurement_sd = excluded.measurement_sd,
            physical_intervals = excluded.physical_intervals,
            source_lineage = excluded.source_lineage, updated_at = excluded.updated_at
    """, (
        cycle_id, *key, record["track"], record["from_m"], record["to_m"], record["avg_wear_min"],
        record["wear_percentage"], record.get("measurement_sd"),
        _json(record.get("physical_intervals", [])), _json(record.get("source_lineage", [])),
        record["created_at"], record["updated_at"],
    ))
    conn.execute(
        """
        DELETE FROM wire_wear_deletion_tombstones
        WHERE line_group = ? AND line_class = ?
          AND cycle_date = ? AND tension_length = ?
        """,
        key,
    )


def _replace_cycle_audit(
    conn: sqlite3.Connection,
    source: dict,
    cycle_keys: set[tuple[str, str, str]],
) -> None:
    for line_group, line_class, cycle_date in cycle_keys:
        cycle_row = conn.execute(
            """
            SELECT cycle_id FROM wire_wear_cycles
            WHERE line_group = ? AND line_class = ? AND cycle_date = ?
            """,
            (line_group, line_class, cycle_date),
        ).fetchone()
        if cycle_row is None:
            continue
        cycle_id = int(cycle_row[0])
        conn.execute("DELETE FROM wire_wear_cycle_segments WHERE cycle_id = ?", (cycle_id,))
        conn.execute("DELETE FROM wire_wear_conflict_decisions WHERE cycle_id = ?", (cycle_id,))
        for segment in source["segments"]:
            if (
                segment.get("line_group"),
                segment.get("line_class"),
                segment.get("cycle_date"),
            ) != (line_group, line_class, cycle_date):
                continue
            conn.execute("""
                INSERT INTO wire_wear_cycle_segments (
                    cycle_id, segment_name, is_present, coverage_percentage,
                    diagnostic_gaps, source_file_names, acquisition_date_from, acquisition_date_to
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                cycle_id, segment["segment_name"], int(bool(segment["is_present"])),
                float(segment["coverage_percentage"]), _json(segment.get("diagnostic_gaps", [])),
                _json(segment.get("source_file_names", [])), segment.get("acquisition_date_from"),
                segment.get("acquisition_date_to"),
            ))
        for decision in source["conflict_decisions"]:
            if (
                decision.get("line_group"),
                decision.get("line_class"),
                decision.get("cycle_date"),
            ) != (line_group, line_class, cycle_date):
                continue
            conn.execute("""
                INSERT INTO wire_wear_conflict_decisions (
                    cycle_id, measurement_identity, source_values, selected_wear_min, accepted_at
                ) VALUES (?, ?, ?, ?, ?)
            """, (
                cycle_id, decision["measurement_identity"], _json(decision.get("source_values", [])),
                float(decision["selected_wear_min"]), _timestamp_text(decision["accepted_at"]),
            ))


def _connection_database_path(conn: sqlite3.Connection) -> Path:
    for row in conn.execute("PRAGMA database_list").fetchall():
        if row[1] == "main" and row[2]:
            return Path(str(row[2]))
    raise SyncValidationError("sync backup requires a file-backed SQLite database")


def _decision_key(value: object) -> tuple[str, str, str, str]:
    if isinstance(value, Mapping):
        return _record_key(value)
    if isinstance(value, (list, tuple)) and len(value) == 4:
        return _record_key({
            "line_group": value[0],
            "line_class": value[1],
            "cycle_date": value[2],
            "tension_length": value[3],
        })
    if isinstance(value, str):
        for separator in ("|", "::"):
            parts = value.split(separator)
            if len(parts) == 4:
                return _decision_key(parts)
    raise SyncValidationError("conflict decision key must contain the full business key")


def _decision_value(value: object) -> str:
    normalized = str(value).strip().lower().replace("-", "_").replace(" ", "_")
    if normalized in {"incoming", "imported", "remote", "use_incoming", "use_remote"}:
        return "incoming"
    if normalized in {"local", "keep_local", "use_local"}:
        return "local"
    raise SyncValidationError("conflict decision must choose incoming or local")


def _conflict_decision_map(
    value: Mapping[object, object] | Sequence[Mapping[str, object]] | None,
) -> dict[tuple[str, str, str, str], str]:
    if value is None:
        return {}
    decisions: dict[tuple[str, str, str, str], str] = {}
    if isinstance(value, Mapping):
        items = value.items()
    elif isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
        parsed_items: list[tuple[object, object]] = []
        for item in value:
            if not isinstance(item, Mapping):
                raise SyncValidationError("conflict decisions must be objects")
            key = item.get("key", item)
            decision = item.get("decision", item.get("resolution"))
            parsed_items.append((key, decision))
        items = parsed_items
    else:
        raise SyncValidationError("conflict decisions must be a mapping or array")
    for raw_key, raw_decision in items:
        key = _decision_key(raw_key)
        if key in decisions:
            raise SyncValidationError("duplicate conflict decision business key")
        decisions[key] = _decision_value(raw_decision)
    return decisions


def _resolve_conflicts(
    actions: Sequence[SyncAction],
    decisions: Mapping[tuple[str, str, str, str], str],
) -> tuple[SyncAction, ...]:
    conflicts = {action.key: action for action in actions if action.action == "conflict"}
    unknown = set(decisions) - set(conflicts)
    if unknown:
        raise SyncValidationError("conflict decision does not match the approved preview")
    unresolved = set(conflicts) - set(decisions)
    if unresolved:
        raise SyncValidationError("sync preview contains unresolved conflicts")
    resolved: list[SyncAction] = []
    for action in actions:
        if action.action != "conflict":
            resolved.append(action)
            continue
        if decisions[action.key] == "incoming":
            resolved.append(SyncAction(
                action.key, "update", "conflict resolved with incoming record",
                action.record, action.tombstone, action.local_record,
                action.local_tombstone, "update",
            ))
        else:
            resolved.append(SyncAction(
                action.key, "unchanged", "conflict resolved by keeping local record",
                action.record, action.tombstone, action.local_record,
                action.local_tombstone, "keep_local",
            ))
    return tuple(resolved)


def apply_sync_import(
    conn: sqlite3.Connection,
    *,
    source_package: object,
    preview_digest: str,
    expected_data_version: int,
    metadata: SyncMetadata,
    metadata_fingerprint: Mapping[str, str] | None = None,
    conflict_decisions: Mapping[object, object] | Sequence[Mapping[str, object]] | None = None,
    db_path: Path | str | None = None,
) -> dict:
    """Revalidate and atomically apply a preview without retaining server-side state."""
    if current_data_version(conn) != expected_data_version:
        raise StaleDataVersionError("wear data version changed; preview again")
    preview = preview_sync_import(conn, source_package, metadata, metadata_fingerprint=metadata_fingerprint)
    if preview.preview_digest != preview_digest:
        raise SyncPreviewMismatchError("sync preview changed; preview again")
    blockers = [action for action in preview.actions if action.action == "error"]
    if blockers:
        raise SyncValidationError("sync preview contains blocking actions")
    decisions = _conflict_decision_map(conflict_decisions)
    resolved_actions = _resolve_conflicts(preview.actions, decisions)
    cycles = {
        (row["line_group"], row["line_class"], row["cycle_date"]): row
        for row in preview.source_package["cycles"]
    }
    changed = [
        action for action in resolved_actions
        if action.action in {"create", "update", "delete"}
    ]
    if current_data_version(conn) != expected_data_version:
        raise StaleDataVersionError("wear data version changed; preview again")
    backup_path = None
    if changed:
        backup_path = create_sqlite_backup(
            conn,
            db_path or _connection_database_path(conn),
        )
    with _transaction(conn):
        if current_data_version(conn) != expected_data_version:
            raise StaleDataVersionError("wear data version changed; preview again")
        for action in changed:
            _apply_action(conn, action, cycles, preview.source_package["package_id"])
        _replace_cycle_audit(
            conn,
            preview.source_package,
            {action.key[:3] for action in changed if action.action in {"create", "update"}},
        )
        if changed:
            data_version = bump_data_version(conn)
        else:
            data_version = current_data_version(conn)
    return {
        "created": sum(action.action == "create" for action in changed),
        "updated": sum(action.action == "update" for action in changed),
        "deleted": sum(action.action == "delete" for action in changed),
        "keep_local": sum(action.status == "keep_local" for action in resolved_actions),
        "no_change": sum(action.status == "no_change" for action in resolved_actions),
        "unchanged": sum(action.action == "unchanged" for action in resolved_actions),
        "conflicts_resolved": len(decisions),
        "backup_path": backup_path,
        "data_version": data_version,
    }
