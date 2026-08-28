"""Deterministic historical workbook discovery and preview orchestration."""

from __future__ import annotations

from collections import defaultdict
from dataclasses import asdict, dataclass
from io import BytesIO
from typing import Any, Iterable, Sequence
import sqlite3
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.calculation.wear_cycle_metadata import (
    MetadataValidationError,
    normalize_tension_length,
)
from app.core.calculation.wear_cycle_types import MetadataInterval
from app.core.calculation.wear_historical_adapters import (
    ADAPTERS,
    HistoricalImportDiagnostic,
    HistoricalSheetCapability,
    ParsedHistoricalSheet,
    sheet_capability,
)
from app.core.calculation.wear_record_candidates import (
    ClassifiedWearRecordCandidate,
    WearRecordCandidate,
    classify_wear_record_candidates,
)


class HistoricalWorkbookError(ValueError):
    """Raised when workbook bytes or sheet selections are invalid."""


class StaleHistoricalPreviewError(HistoricalWorkbookError):
    """Raised when the supplied committed data version is no longer current."""


_CLASSIFIER_TL_BATCH_SIZE = 16


@dataclass(frozen=True)
class HistoricalWorkbookDiscovery:
    sheets: tuple[HistoricalSheetCapability, ...]

    def to_dict(self) -> dict[str, Any]:
        return {"sheets": [asdict(sheet) for sheet in self.sheets]}


@dataclass(frozen=True)
class HistoricalSheetSummary:
    sheet_name: str
    skipped: int
    new: int
    update: int
    no_change: int
    duplicate: int
    error: int

    @property
    def total(self) -> int:
        return self.new + self.update + self.no_change + self.duplicate + self.error

    def to_dict(self) -> dict[str, Any]:
        return {**asdict(self), "total": self.total}


@dataclass(frozen=True)
class HistoricalWorkbookPreview:
    discovery: HistoricalWorkbookDiscovery
    selected_sheets: tuple[str, ...]
    data_version: int
    candidates: tuple[ClassifiedWearRecordCandidate, ...]
    diagnostics: tuple[HistoricalImportDiagnostic, ...]
    sheet_summaries: tuple[HistoricalSheetSummary, ...]

    def to_dict(self) -> dict[str, Any]:
        totals = {
            "skipped": sum(item.skipped for item in self.sheet_summaries),
            "new": sum(item.new for item in self.sheet_summaries),
            "update": sum(item.update for item in self.sheet_summaries),
            "no_change": sum(item.no_change for item in self.sheet_summaries),
            "duplicate": sum(item.duplicate for item in self.sheet_summaries),
            "error": sum(item.error for item in self.sheet_summaries),
        }
        totals["total"] = sum(
            totals[key] for key in ("new", "update", "no_change", "duplicate", "error")
        )
        return {
            **self.discovery.to_dict(),
            "selected_sheets": list(self.selected_sheets),
            "wire_wear_data_version": self.data_version,
            "sheet_summaries": [item.to_dict() for item in self.sheet_summaries],
            "totals": totals,
            "candidates": [candidate.to_dict() for candidate in self.candidates],
            "diagnostics": [diagnostic.to_dict() for diagnostic in self.diagnostics],
        }


def _workbook(workbook_bytes: bytes):
    if not isinstance(workbook_bytes, bytes) or not workbook_bytes:
        raise HistoricalWorkbookError("workbook must contain XLSX bytes")
    try:
        return load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise HistoricalWorkbookError("workbook is not a readable XLSX file") from exc


def _discovery(sheet_names: Iterable[str]) -> HistoricalWorkbookDiscovery:
    return HistoricalWorkbookDiscovery(
        tuple(sheet_capability(sheet_name) for sheet_name in sheet_names)
    )


def discover_historical_workbook(workbook_bytes: bytes) -> HistoricalWorkbookDiscovery:
    workbook = _workbook(workbook_bytes)
    try:
        return _discovery(workbook.sheetnames)
    finally:
        workbook.close()


def _classify_groups(
    conn: sqlite3.Connection,
    candidates: Sequence[WearRecordCandidate],
    metadata: Sequence[MetadataInterval],
    committed_data_version: int,
) -> tuple[ClassifiedWearRecordCandidate, ...]:
    metadata_by_tl: dict[str, list[MetadataInterval]] = defaultdict(list)
    for interval in metadata:
        metadata_by_tl[normalize_tension_length(interval.tension_length)].append(interval)

    groups: dict[
        tuple[str, str, str], dict[str, list[WearRecordCandidate]]
    ] = defaultdict(lambda: defaultdict(list))
    candidate_order: list[str] = []
    for candidate in candidates:
        group_key = (
            str(candidate.line_group),
            str(candidate.line_class),
            candidate.cycle_date.isoformat(),
        )
        try:
            tension_length_key = normalize_tension_length(candidate.tension_length)
        except (MetadataValidationError, TypeError, ValueError):
            tension_length_key = "!invalid"
        groups[group_key][tension_length_key].append(candidate)
        if candidate.row_id is not None:
            candidate_order.append(candidate.row_id)

    classified_by_id: dict[str, ClassifiedWearRecordCandidate] = {}
    for group_key in sorted(groups):
        candidates_by_tl = groups[group_key]
        tension_lengths = sorted(candidates_by_tl)
        for offset in range(0, len(tension_lengths), _CLASSIFIER_TL_BATCH_SIZE):
            batch_tension_lengths = tension_lengths[
                offset:offset + _CLASSIFIER_TL_BATCH_SIZE
            ]
            batch_candidates = tuple(
                candidate
                for tension_length in batch_tension_lengths
                for candidate in candidates_by_tl[tension_length]
            )
            batch_metadata = tuple(
                interval
                for tension_length in batch_tension_lengths
                for interval in metadata_by_tl.get(tension_length, ())
            )
            preview = classify_wear_record_candidates(
                conn, batch_candidates, batch_metadata
            )
            if preview.data_version != committed_data_version:
                raise StaleHistoricalPreviewError(
                    "wire wear data version changed; reload and preview again"
                )
            for candidate in preview.candidates:
                if candidate.row_id is None:
                    raise HistoricalWorkbookError(
                        "classified workbook candidate has no row id"
                    )
                classified_by_id[candidate.row_id] = candidate
    return tuple(classified_by_id[row_id] for row_id in candidate_order)


def _summary(
    parsed: ParsedHistoricalSheet,
    candidates: Sequence[ClassifiedWearRecordCandidate],
) -> HistoricalSheetSummary:
    counts = {status: 0 for status in ("new", "update", "no_change", "duplicate", "error")}
    for candidate in candidates:
        if candidate.source_sheet == parsed.sheet_name:
            counts[candidate.status] += 1
    counts["error"] += len(parsed.diagnostics)
    return HistoricalSheetSummary(
        sheet_name=parsed.sheet_name,
        skipped=parsed.skipped_cells,
        **counts,
    )


def preview_historical_workbook(
    conn: sqlite3.Connection,
    workbook_bytes: bytes,
    selected_sheets: Sequence[str],
    metadata: Sequence[MetadataInterval],
    committed_data_version: int,
) -> HistoricalWorkbookPreview:
    """Parse and classify selected workbook sheets without writing SQLite."""
    workbook = _workbook(workbook_bytes)
    try:
        discovery = _discovery(workbook.sheetnames)
        selected = tuple(dict.fromkeys(selected_sheets))
        if not selected:
            raise HistoricalWorkbookError("at least one supported sheet must be selected")
        unavailable = [
            sheet_name
            for sheet_name in selected
            if sheet_name not in workbook.sheetnames or sheet_name not in ADAPTERS
        ]
        if unavailable:
            raise HistoricalWorkbookError(
                "selected sheets are not supported: " + ", ".join(unavailable)
            )

        parsed_sheets = tuple(
            ADAPTERS[sheet_name].parse(workbook[sheet_name])
            for sheet_name in workbook.sheetnames
            if sheet_name in selected
        )
        raw_candidates = tuple(
            candidate
            for parsed in parsed_sheets
            for candidate in parsed.candidates
        )
        classified = _classify_groups(
            conn, raw_candidates, metadata, committed_data_version
        ) if raw_candidates else ()
        diagnostics = tuple(
            diagnostic
            for parsed in parsed_sheets
            for diagnostic in parsed.diagnostics
        )
        summaries = tuple(_summary(parsed, classified) for parsed in parsed_sheets)
        return HistoricalWorkbookPreview(
            discovery=discovery,
            selected_sheets=selected,
            data_version=committed_data_version,
            candidates=classified,
            diagnostics=diagnostics,
            sheet_summaries=summaries,
        )
    finally:
        workbook.close()
