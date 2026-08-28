"""Source adapters for historical Wire Wear monthly workbooks."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from typing import Any, Literal, Protocol
import math
import re

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.calculation.wear_record_candidates import WearRecordCandidate


SheetSupport = Literal["supported", "reserved", "ignored"]


@dataclass(frozen=True)
class HistoricalSheetCapability:
    sheet_name: str
    support: SheetSupport
    selected_by_default: bool
    enabled: bool
    reason_code: str


@dataclass(frozen=True)
class HistoricalImportDiagnostic:
    code: str
    message: str
    sheet: str
    cell: str | None
    original_value: Any = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "code": self.code,
            "message": self.message,
            "sheet": self.sheet,
            "cell": self.cell,
            "original_value": self.original_value,
        }


@dataclass(frozen=True)
class ParsedHistoricalSheet:
    sheet_name: str
    candidates: tuple[WearRecordCandidate, ...]
    diagnostics: tuple[HistoricalImportDiagnostic, ...]
    skipped_cells: int


class HistoricalSheetAdapter(Protocol):
    sheet_name: str
    line_group: str
    line_class: str

    def parse(self, worksheet: Worksheet) -> ParsedHistoricalSheet: ...


def _normalized_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _month_heading(value: Any) -> date | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float) and math.isfinite(value) and value.is_integer():
        text = str(int(value))
    elif isinstance(value, str):
        text = value.strip()
    else:
        return None
    if re.fullmatch(r"\d{6}", text) is None:
        return None
    year = int(text[:4])
    month = int(text[4:])
    try:
        return date(year, month, 1)
    except ValueError:
        return None


def _track(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    normalized = value.strip().upper()
    return {"UP": "UP", "DOWN": "DOWN", "DN": "DOWN"}.get(normalized)


def _has_source_value(value: Any) -> bool:
    if value is None:
        return False
    return not (isinstance(value, str) and not value.strip())


@dataclass(frozen=True)
class MonthlyMatrixAdapter:
    sheet_name: str
    line_group: str
    line_class: str

    def parse(self, worksheet: Worksheet) -> ParsedHistoricalSheet:
        diagnostics: list[HistoricalImportDiagnostic] = []
        candidates: list[WearRecordCandidate] = []
        skipped_cells = 0

        required = ((1, "track"), (2, "tension length"))
        missing_required = False
        for column, expected in required:
            cell = worksheet.cell(1, column)
            if _normalized_header(cell.value) != expected:
                diagnostics.append(HistoricalImportDiagnostic(
                    code="missing_required_column",
                    message=f"column {get_column_letter(column)} must be {expected.title()}",
                    sheet=worksheet.title,
                    cell=cell.coordinate,
                    original_value=cell.value,
                ))
                missing_required = True
        if missing_required:
            return ParsedHistoricalSheet(
                worksheet.title, (), tuple(diagnostics), skipped_cells
            )

        months: list[tuple[int, date]] = []
        month_cells: dict[date, str] = {}
        for column in range(3, worksheet.max_column + 1):
            cell = worksheet.cell(1, column)
            column_has_data = any(
                _has_source_value(worksheet.cell(row, column).value)
                for row in range(2, worksheet.max_row + 1)
            )
            if not _has_source_value(cell.value) and not column_has_data:
                continue
            month = _month_heading(cell.value)
            if month is None:
                diagnostics.append(HistoricalImportDiagnostic(
                    code="invalid_month_heading",
                    message="month heading must be a valid YYYYMM value",
                    sheet=worksheet.title,
                    cell=cell.coordinate,
                    original_value=cell.value,
                ))
                continue
            first_cell = month_cells.get(month)
            if first_cell is not None:
                diagnostics.append(HistoricalImportDiagnostic(
                    code="duplicate_month_heading",
                    message=f"month heading duplicates {first_cell}",
                    sheet=worksheet.title,
                    cell=cell.coordinate,
                    original_value=cell.value,
                ))
                continue
            month_cells[month] = cell.coordinate
            months.append((column, month))

        for row in range(2, worksheet.max_row + 1):
            track_cell = worksheet.cell(row, 1)
            tension_cell = worksheet.cell(row, 2)
            measurement_cells = [worksheet.cell(row, column) for column, _ in months]
            populated = [cell for cell in measurement_cells if _has_source_value(cell.value)]
            if not populated:
                skipped_cells += len(measurement_cells)
                continue

            normalized_track = _track(track_cell.value)
            if normalized_track is None:
                diagnostics.append(HistoricalImportDiagnostic(
                    code="invalid_track",
                    message="Track must normalize to UP or DOWN",
                    sheet=worksheet.title,
                    cell=track_cell.coordinate,
                    original_value=track_cell.value,
                ))
                continue
            if not _has_source_value(tension_cell.value):
                diagnostics.append(HistoricalImportDiagnostic(
                    code="empty_tension_length",
                    message="Tension Length cannot be empty on a populated row",
                    sheet=worksheet.title,
                    cell=tension_cell.coordinate,
                    original_value=tension_cell.value,
                ))
                continue

            for column, month in months:
                cell = worksheet.cell(row, column)
                raw_value = cell.value
                if raw_value is None or (
                    isinstance(raw_value, str) and raw_value.strip() in {"", "."}
                ):
                    skipped_cells += 1
                    continue
                if isinstance(raw_value, bool) or not isinstance(raw_value, (int, float)):
                    diagnostics.append(HistoricalImportDiagnostic(
                        code="non_numeric_wear",
                        message="wear value must be numeric, blank, or a period",
                        sheet=worksheet.title,
                        cell=cell.coordinate,
                        original_value=raw_value,
                    ))
                    continue
                candidates.append(WearRecordCandidate(
                    line_group=self.line_group,
                    line_class=self.line_class,
                    cycle_date=month,
                    tension_length=tension_cell.value,
                    avg_wear_min=raw_value,
                    track=normalized_track,
                    source_type="historical_excel",
                    source_sheet=worksheet.title,
                    source_row=row,
                    tension_length_cell=tension_cell.coordinate,
                    avg_wear_min_cell=cell.coordinate,
                    track_cell=track_cell.coordinate,
                    original_value=raw_value,
                    row_id=f"{worksheet.title}!{cell.coordinate}",
                ))

        return ParsedHistoricalSheet(
            worksheet.title,
            tuple(candidates),
            tuple(diagnostics),
            skipped_cells,
        )


class EALAdapter(MonthlyMatrixAdapter):
    def __init__(self) -> None:
        super().__init__("EAL", "EAL", "EAL")


class TMLAdapter(MonthlyMatrixAdapter):
    def __init__(self) -> None:
        super().__init__("TML", "TML", "TML")


class LMCAdapter(MonthlyMatrixAdapter):
    def __init__(self) -> None:
        super().__init__("LMC", "EAL", "LMC")


ADAPTERS: dict[str, HistoricalSheetAdapter] = {
    adapter.sheet_name: adapter
    for adapter in (EALAdapter(), TMLAdapter(), LMCAdapter())
}


def sheet_capability(sheet_name: str) -> HistoricalSheetCapability:
    if sheet_name in ADAPTERS:
        return HistoricalSheetCapability(
            sheet_name, "supported", True, True, "supported_monthly_matrix"
        )
    if sheet_name == "LRL":
        return HistoricalSheetCapability(
            sheet_name, "reserved", False, False, "reserved_future_adapter"
        )
    return HistoricalSheetCapability(
        sheet_name, "ignored", False, False, "unknown_sheet"
    )
