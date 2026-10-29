"""Bounded-memory ChartData writer; legacy exception/report semantics unchanged."""
from copy import copy
import io
import math
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
import pandas as pd

from .exporter import ExcelExporter, report_number_format
from .tov1050_contract import LINE_WORKBOOKS

REPORT_NUMERIC_FORMAT = "0.####"
REPORT_NUMERIC_COLUMNS = {
    "Chainage", "stg_max", "stg_min", "height_min", "height_max", "wear_min", "wear_max",
    *[f"{kind}{index}" for kind in ("height", "wear", "stagger") for index in range(1, 5)],
}


def _excel_value(value):
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, float):
        return None if math.isnan(value) else str(value) if math.isinf(value) else value
    return value


def export_tov1050_report(results, chart_df=None, metadata=None, *, chart_statistics_df=None,
                          output_path=None, progress=None, _numeric_format=REPORT_NUMERIC_FORMAT):
    """Delegate Summary/exception generation, stream full ChartData row by row.

    This function does not map metadata, aggregate exceptions, or downsample.
    The small legacy workbook provides exact headers/order and formatting.
    """
    # The legacy exporter adds report columns in place; published detector frames
    # are immutable snapshots shared by background exports.
    report_results = {name: value.copy(deep=True) if isinstance(value, pd.DataFrame) else value
                      for name, value in results.items()}
    small = ExcelExporter.export_report(report_results, chart_df=None, numeric_format=_numeric_format)
    source = load_workbook(small, read_only=True)
    target = Workbook(write_only=True)
    original_header = source["Summary"]["A1"]
    prototype = WriteOnlyCell(target.create_sheet("_style_template"))
    for attribute in ("font", "border", "fill", "alignment"):
        setattr(prototype, attribute, copy(getattr(original_header, attribute)))
    header_style = prototype._style
    target.remove(prototype.parent)
    numeric_styles = {}

    def style(sheet, number_format):
        if number_format not in numeric_styles:
            prototype = WriteOnlyCell(sheet)
            prototype.number_format = number_format
            numeric_styles[number_format] = prototype._style
        return numeric_styles[number_format]

    def header(sheet, values):
        cells=[]
        for value in values:
            cell=WriteOnlyCell(sheet,value=value)
            cell._style = header_style
            cells.append(cell)
        sheet.append(cells)

    def write_frame(title, frame, columns):
        sheet = target.create_sheet(title)
        header(sheet, columns)
        formatted = [bool(_numeric_format and (column in REPORT_NUMERIC_COLUMNS
                     or column.endswith(("_min", "_max")))) for column in columns]
        total = len(frame)
        if progress:
            progress(f"exporting_{title}", 0, total, "rows")
        for index, row in enumerate(frame[columns].itertuples(index=False, name=None), 1):
            cells = []
            for value, should_format in zip(row, formatted):
                value = _excel_value(value)
                if not should_format:
                    cells.append(value)
                else:
                    cell = WriteOnlyCell(sheet, value=value)
                    cell._style = style(sheet, report_number_format(value, _numeric_format))
                    cells.append(cell)
            sheet.append(cells)
            if progress and index % 1000 == 0:
                progress(f"exporting_{title}", index, total, "rows")
        if progress:
            progress(f"exporting_{title}", total, total, "rows")

    try:
        if progress:
            progress("exporting", 0, None, "rows")
        for original in source:
            sheet=target.create_sheet(original.title)
            for index,row in enumerate(original.iter_rows()):
                if index == 0:
                    header(sheet,[cell.value for cell in row])
                else:
                    cells = []
                    for original_cell in row:
                        cell = WriteOnlyCell(sheet, value=original_cell.value)
                        cell._style = style(sheet, original_cell.number_format)
                        cells.append(cell)
                    sheet.append(cells)
                if progress and index % 1000 == 0:
                    progress("exporting", index, None, "rows")
        if chart_df is not None and not chart_df.empty:
            write_frame("ChartData", chart_df, ExcelExporter.chart_report_columns(chart_df))
        if chart_statistics_df is not None and not chart_statistics_df.empty:
            write_frame("ChainageStatistics", chart_statistics_df, list(chart_statistics_df.columns))
        if metadata:
            sheet=target.create_sheet("_Metadata")
            sheet.sheet_state="hidden"
            header(sheet,["key","value"])
            for key,value in metadata.items():
                sheet.append([key,str(value)])
        if progress:
            progress("finalizing", None, None, None)
        output = Path(output_path) if output_path is not None else io.BytesIO()
        target.save(output)
        if progress:
            progress("finalizing", None, None, None)
        if output_path is None:
            output.seek(0)
        return output
    finally:
        source.close()
        # Write-only worksheets own temporary XML files, including on cancellation.
        for sheet in target:
            if sheet._writer is not None:
                if not sheet.closed:
                    sheet.close()
                if Path(sheet._writer.out).exists():
                    sheet._writer.cleanup()
        target.close()


def export_analysis_report(results, chart_df=None, metadata=None, line="", *, chart_statistics_df=None,
                           output_path=None, progress=None):
    if chart_statistics_df is not None or output_path is not None or progress is not None:
        return export_tov1050_report(results, chart_df, metadata, chart_statistics_df=chart_statistics_df,
                                     output_path=output_path, progress=progress,
                                     _numeric_format=REPORT_NUMERIC_FORMAT if str(line).strip().upper() in LINE_WORKBOOKS else None)
    if str(line).strip().upper() in LINE_WORKBOOKS and chart_df is not None and len(chart_df)>50_000:
        return export_tov1050_report(results,chart_df,metadata)
    if str(line).strip().upper() in LINE_WORKBOOKS:
        return ExcelExporter.export_report(results, chart_df=chart_df, metadata=metadata,
                                           numeric_format=REPORT_NUMERIC_FORMAT)
    return ExcelExporter.export_report(results,chart_df=chart_df,metadata=metadata)
