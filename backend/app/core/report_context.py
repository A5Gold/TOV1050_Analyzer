"""Shared, read-only identity contract for legacy and current report workbooks.

Dates/tasks describe an acquisition. Comparison scope deliberately excludes them.
Source identities are retained, including the legacy LMC/TKS line aliases.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
import json
import io
import math
import posixpath
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import zipfile
from typing import Any, Iterable, Mapping

from app.core.line_registry import normalize_context, normalize_direction
from openpyxl import load_workbook


def load_report_preflight(source):
    """Open only small context sheets without touching any ChartData XML.

    openpyxl's read-only loader scans dimensionless write-only worksheets at
    initialization. Filtering workbook sheet references first avoids that scan
    for engineering capability checks. The source workbook is never modified.
    """
    source = io.BytesIO(source) if isinstance(source, bytes) else source
    sheet_namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    relation_namespace = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    with zipfile.ZipFile(source) as archive:
        root = ET.fromstring(archive.read('xl/workbook.xml'))
        relations = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets = {item.attrib['Id']: item.attrib.get('Target', '') for item in relations}
        sheets = root.find(f'{{{sheet_namespace}}}sheets')
        selected_paths = set()
        if sheets is not None:
            for sheet in list(sheets):
                if sheet.attrib.get('name') not in {'Summary', 'Wire Wear', '_Metadata'}:
                    sheets.remove(sheet)
                    continue
                target = targets.get(sheet.attrib.get(f'{{{relation_namespace}}}id'), '')
                selected_paths.add(posixpath.normpath(target.lstrip('/') if target.startswith('/') else 'xl/' + target))
        essentials = {'[Content_Types].xml', '_rels/.rels', 'xl/workbook.xml',
                      'xl/_rels/workbook.xml.rels', 'xl/styles.xml', 'xl/sharedStrings.xml'}
        filtered = io.BytesIO()
        with zipfile.ZipFile(filtered, 'w', compression=zipfile.ZIP_STORED) as target_archive:
            for name in essentials | selected_paths:
                if name not in archive.namelist():
                    continue
                target_archive.writestr(name, ET.tostring(root) if name == 'xl/workbook.xml' else archive.read(name))
    filtered.seek(0)
    return load_workbook(filtered, read_only=True, data_only=True, keep_links=False)


CONTEXT_ALIASES = {
    "line": ("line", "parent_line", "logical_line"),
    "branch": ("branch", "session", "legacy_session"),
    "track": ("track", "direction"),
    "section": ("section",),
    "task_no": ("task_no", "task no"),
    "station_start": ("station_start", "st. start", "stn start"),
    "station_end": ("station_end", "st. end", "stn end"),
    "run_date": ("task_run_date", "task run date", "run_date", "run date"),
}


class ReportContextError(ValueError):
    def __init__(self, message: str, code: str = "INVALID_REPORT_CONTEXT"):
        super().__init__(message)
        self.code = code

    def detail(self) -> dict:
        return {"code": self.code, "message": str(self)}


def text_value(value: Any) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    return str(value).strip()


def context_values(row: Mapping[str, Any]) -> dict[str, str]:
    columns = {str(key).strip().lower(): value for key, value in row.items()}
    return {
        key: next((text_value(columns[name]) for name in aliases if text_value(columns.get(name))), "")
        for key, aliases in CONTEXT_ALIASES.items()
    }


@dataclass(frozen=True)
class ReportContext:
    line: str = ""
    branch: str = ""
    direction: str = ""
    track: str = ""
    section: str = ""
    task_no: str = ""
    station_start: str = ""
    station_end: str = ""
    run_date: str = ""
    legacy_session: str = ""
    source_line: str = ""
    source_track: str = ""
    algorithm: str = ""
    algorithm_mode: str = ""

    @property
    def scope(self) -> tuple[str, str, str]:
        return self.line, self.branch, self.direction


def normalize_report_row(row: Mapping[str, Any], fallback: ReportContext | None = None,
                         *, source: str = "report") -> ReportContext:
    values = context_values(row)
    base = fallback or ReportContext()
    raw_line = values["line"] or base.line
    raw_branch = values["branch"]
    section = values["section"] or base.section
    source_line = text_value(row.get('source_line')) or values["line"] or base.source_line
    source_track = text_value(row.get('source_track')) or values["track"] or base.source_track
    # These aliases belong to report/legacy-wear sources, not the parent catalog.
    if raw_line.upper() in {"LMC", "TKS"}:
        alias_branch = raw_line.upper()
        if raw_branch and raw_branch.upper() != alias_branch:
            raise ReportContextError(f"{source}: line {raw_line} conflicts with branch {raw_branch}")
        raw_line = "EAL" if alias_branch == "LMC" else "TKL"
        raw_branch = alias_branch
    if not raw_line:
        try:
            unknown_line_direction = normalize_direction(values['track'] or base.direction)
        except ValueError as exc:
            raise ReportContextError(f'{source}: {exc}') from exc
        return ReportContext(
            direction=unknown_line_direction,
            track=values["track"] or base.track, section=section,
            task_no=values["task_no"] or base.task_no,
            station_start=values["station_start"] or base.station_start,
            station_end=values["station_end"] or base.station_end,
            run_date=values["run_date"] or base.run_date, algorithm=base.algorithm,
            algorithm_mode=base.algorithm_mode,
            source_line=source_line, source_track=source_track,
        )
    # A row's explicit legacy Section may identify a branch; it must take
    # precedence over an absent branch inherited from another row.
    branch = raw_branch or base.branch or None
    try:
        normalized = normalize_context(raw_line, branch=branch,
                                       direction=values["track"] or base.direction or None,
                                       section=section or None)
    except ValueError as exc:
        raise ReportContextError(f"{source}: {exc}") from exc
    return ReportContext(
        line=normalized.line, branch=normalized.branch,
        direction=normalized.direction or "", track=values["track"] or base.track,
        section=section, task_no=values["task_no"] or base.task_no,
        station_start=values["station_start"] or base.station_start,
        station_end=values["station_end"] or base.station_end,
        run_date=values["run_date"] or base.run_date,
        legacy_session=normalized.legacy_session, source_line=source_line,
        source_track=source_track, algorithm=base.algorithm, algorithm_mode=base.algorithm_mode,
    )


def ensure_compatible_scopes(contexts: Iterable[ReportContext], *, source: str = "reports") -> ReportContext:
    """Reject known disagreements; wholly unidentified legacy reports stay unknown."""
    entries = list(contexts)
    for field in ("line", "branch", "direction"):
        values = {getattr(item, field) for item in entries if getattr(item, field)}
        if len(values) > 1:
            raise ReportContextError(
                f"{source}: incompatible {field}: {', '.join(sorted(values))}. "
                "Compare each line/branch/direction separately.", "INCOMPATIBLE_REPORT_SCOPE")
    fields = ReportContext.__dataclass_fields__
    merged = {}
    for field in fields:
        values = {getattr(item, field) for item in entries if getattr(item, field)}
        merged[field] = next(iter(values)) if len(values) == 1 else ""
    return ReportContext(**merged)


def sheet_records(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = [text_value(value) for value in next(rows, ())]
    for row in rows:
        if any(value is not None for value in row):
            yield {key: row[index] if index < len(row) else None
                   for index, key in enumerate(headers) if key}


def metadata_values(workbook) -> dict[str, Any]:
    if "_Metadata" not in workbook.sheetnames:
        return {}
    values = {}
    for row in workbook["_Metadata"].iter_rows(values_only=True):
        if len(row) < 2 or not text_value(row[0]) or text_value(row[0]).lower() == "key":
            continue
        value = row[1]
        if isinstance(value, str) and value[:1] in {"{", "["}:
            try:
                value = json.loads(value)
            except json.JSONDecodeError:
                pass
        values[text_value(row[0])] = value
    return values


def filename_context(filename: str) -> ReportContext:
    """Only recognized report filename layouts supply missing identity fields."""
    parts = Path(filename).stem.split("_")
    if len(parts) < 4 or not re.fullmatch(r"\d{8}", parts[0]):
        return ReportContext()
    if parts[1].upper() not in {"EAL", "TML", "LMC", "AEL", "TCL", "KTL", "ISL", "TWL", "TKL", "TKS", "DRL"}:
        return ReportContext()
    row = {"line": parts[1], "run_date": parts[0]}
    if parts[2].upper() in {"UP", "DN", "DOWN", "UT", "DT"}:
        row.update(track=parts[2], section=parts[3])
        if len(parts) >= 6 and parts[4] != "Exception":
            row.update(station_start=parts[4], station_end=parts[5])
    else:
        row["task_no"] = parts[2]
        if "-" in parts[3]:
            row["station_start"], row["station_end"] = parts[3].split("-", 1)
    return normalize_report_row(row, source="filename")


def read_report_context(workbook, filename: str = "", *, chart_rows=None,
                        include_chart: bool = False) -> ReportContext:
    """Preflight small sheets without decoding ChartData unless explicitly needed."""
    meta = metadata_values(workbook)
    capabilities = meta.get("metadata_capabilities")
    values = dict(capabilities) if isinstance(capabilities, dict) else {}
    values.update({key: value for key, value in meta.items() if not isinstance(value, (dict, list))})
    algorithm = text_value(meta.get('aggregation_version') or meta.get('aggregation_algorithm') or meta.get('algorithm'))
    mode = text_value(meta.get('aggregation') or meta.get('chart_data_mode') or meta.get('analysis_mode'))
    mode = {'median': 'chainage_median', 'raw_overview': 'raw'}.get(mode, mode)
    meta_context = normalize_report_row(values, source="_Metadata")
    contexts = [meta_context]
    seen = set()

    def append_contexts(rows, source):
        for row in rows:
            values = context_values(row)
            values['source_line'] = text_value(row.get('source_line'))
            key = tuple(values.values())
            if key not in seen:
                seen.add(key)
                contexts.append(normalize_report_row(values, source=source))

    for name in ("Summary", "Wire Wear"):
        if name in workbook.sheetnames:
            append_contexts(sheet_records(workbook[name]), name)
    if chart_rows is not None:
        append_contexts(chart_rows, 'ChartData')
    elif include_chart and "ChartData" in workbook.sheetnames:
        append_contexts(sheet_records(workbook['ChartData']), 'ChartData')
    context = ensure_compatible_scopes(contexts)
    # Filename is a fallback, never an authority over workbook cells.
    fallback = filename_context(filename)
    if not context.line or fallback.line == context.line:
        context = normalize_report_row(asdict(context), fallback, source='filename fallback')
    return ReportContext(**{**asdict(context), "algorithm": algorithm, 'algorithm_mode': mode})


def validate_algorithm_compatibility(contexts: Iterable[ReportContext]) -> bool:
    entries = list(contexts)
    for field in ('algorithm', 'algorithm_mode'):
        known = {getattr(item, field) for item in entries if getattr(item, field)}
        if len(known) > 1:
            raise ReportContextError("Reports use different analysis algorithms. Regenerate them with the same mode/version.",
                                     "INCOMPATIBLE_REPORT_ALGORITHM")
    return bool(entries) and all(item.algorithm for item in entries)
