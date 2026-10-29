"""Dedicated ChartData alignment API for the Version Difference module."""

import logging
from dataclasses import asdict
from pathlib import Path
from typing import Any, Dict, List, Optional

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from fastapi import APIRouter, File, HTTPException, UploadFile

from app.core.calculation.chainage_alignment import build_aligned_comparison, prepare_chart
from app.core.report_context import (
    CONTEXT_ALIASES, ReportContext, ReportContextError,
    ensure_compatible_scopes, filename_context, normalize_report_row,
    read_report_context, validate_algorithm_compatibility,
)

logger = logging.getLogger(__name__)
router = APIRouter()

MAX_ROWS = 320_000
MAX_FILE_BYTES = 64 * 1024 * 1024
MAX_TOTAL_BYTES = 320 * 1024 * 1024
# Marker is retained as a lightweight compatibility field for existing API
# fixtures and diagnostics; alignment ignores all columns outside this set.
RECOGNIZED_COLUMNS = {"Chainage", "Marker", *(f"{metric}{channel}" for metric in ("height", "stagger", "wear") for channel in range(1, 5))}


class ReportChart(dict):
    """Compact numeric columns with identity kept outside alignment input."""
    def __init__(self, columns, context=None, sections=()):
        super().__init__(columns)
        self.context = context or ReportContext()
        self.sections = frozenset(sections)


def _validate_comparison_context(charts):
    contexts = [getattr(chart, 'context', ReportContext()) for chart in charts if chart is not None]
    shared = ensure_compatible_scopes(contexts, source='Version Difference')
    sections = {section.strip().casefold() for chart in charts if chart is not None
                for section in getattr(chart, 'sections', ()) if section}
    if len(sections) > 1:
        raise ReportContextError('Version Difference requires one compatible Section. Compare sections separately.',
                                 'INCOMPATIBLE_REPORT_SECTION')
    return shared, validate_algorithm_compatibility(contexts)


async def _load_chart_data(upload: UploadFile, role: str) -> Optional[Dict[str, List[Any]]]:
    """Read one workbook once and return only its ChartData columns."""
    filename = upload.filename or ""
    workbook = None
    if Path(filename).suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
        raise HTTPException(status_code=400, detail=f"The {role} file must be a supported Excel workbook.")

    try:
        upload.file.seek(0, 2)
        size = upload.file.tell()
        upload.file.seek(0)
        if size > MAX_FILE_BYTES:
            raise HTTPException(status_code=413, detail=f"The {role} file exceeds the 64 MiB limit.")
        if size == 0:
            raise ValueError("empty upload")
        if Path(filename).suffix.lower() in {".xlsx", ".xlsm"}:
            workbook = load_workbook(upload.file, read_only=True, data_only=True)
            if "ChartData" not in workbook.sheetnames:
                workbook.close()
                return None
            sheet = workbook["ChartData"]
            context = read_report_context(workbook, filename)
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
            selected = [index for index, header in enumerate(headers) if header in RECOGNIZED_COLUMNS]
            if "Chainage" not in headers:
                workbook.close()
                raise HTTPException(status_code=422, detail=f"The {role} ChartData sheet is missing required Chainage.")
            columns: Dict[str, List[Any]] = {headers[index]: [] for index in selected}
            identity_indices = [index for index, header in enumerate(headers)
                                if header.lower() in {alias for aliases in CONTEXT_ALIASES.values() for alias in aliases}]
            unique_contexts = set()
            for row_count, row in enumerate(rows, start=1):
                if row_count > MAX_ROWS:
                    workbook.close()
                    raise HTTPException(status_code=422, detail=f"The {role} file exceeds the 320,000 ChartData row limit.")
                for index in selected:
                    columns[headers[index]].append(row[index] if index < len(row) else None)
                unique_contexts.add(tuple(row[index] if index < len(row) else None for index in identity_indices))
            row_contexts = [normalize_report_row(dict(zip((headers[index] for index in identity_indices), values)),
                                                context, source=f'{role} ChartData') for values in unique_contexts]
            shared = ensure_compatible_scopes([context, *row_contexts], source=f'{role} report')
            sections = {item.section for item in [context, *row_contexts] if item.section}
            workbook.close()
            return ReportChart(columns, shared, sections)
        # Legacy .xls compatibility path. The byte guard still applies and the
        # dataframe is immediately converted to the compact recognized columns.
        chart_df = pd.read_excel(upload.file, sheet_name="ChartData", nrows=MAX_ROWS + 1)
        if len(chart_df.index) > MAX_ROWS:
            raise HTTPException(status_code=422, detail=f"The {role} file exceeds the 320,000 ChartData row limit.")
        if "Chainage" not in chart_df.columns:
            raise HTTPException(status_code=422, detail=f"The {role} ChartData sheet is missing required Chainage.")
        contexts = [normalize_report_row(row, filename_context(filename), source=f'{role} ChartData')
                    for row in chart_df[[column for column in chart_df.columns
                                       if str(column).strip().lower() in {alias for aliases in CONTEXT_ALIASES.values() for alias in aliases}]].drop_duplicates().to_dict('records')]
        shared = ensure_compatible_scopes(contexts)
        columns = chart_df[[column for column in chart_df.columns if str(column).strip() in RECOGNIZED_COLUMNS]].replace({np.nan: None}).to_dict(orient="list")
        return ReportChart(columns, shared, {item.section for item in contexts if item.section})
    except ReportContextError as exc:
        raise HTTPException(status_code=422, detail=exc.detail()) from exc
    except HTTPException:
        raise
    except Exception:
        logger.exception("Could not read Version Difference %s workbook", role)
        raise HTTPException(status_code=400, detail=f"The {role} file could not be read as an Excel workbook.")
    finally:
        if workbook is not None:
            workbook.close()

@router.post("/analyze/version-difference")
async def analyze_version_difference(
    latest: UploadFile = File(...),
    previous_1: UploadFile = File(...),
    previous_2: Optional[UploadFile] = File(None),
    previous_3: Optional[UploadFile] = File(None),
    previous_4: Optional[UploadFile] = File(None),
):
    """Align each provided previous cycle independently against Latest."""
    ordered_uploads = (
        ("latest", latest),
        ("previous_1", previous_1),
        ("previous_2", previous_2),
        ("previous_3", previous_3),
        ("previous_4", previous_4),
    )
    uploads = [
        (role, upload) for role, upload in ordered_uploads if upload is not None
    ]

    total_bytes = 0
    for role, upload in uploads:
        upload.file.seek(0, 2)
        total_bytes += upload.file.tell()
        upload.file.seek(0)
    if total_bytes > MAX_TOTAL_BYTES:
        raise HTTPException(status_code=413, detail="The Version Difference request exceeds the 320 MiB total upload limit.")

    charts: Dict[str, Optional[Dict[str, List[Any]]]] = {}
    charts["latest"] = await _load_chart_data(latest, "latest")

    if charts["latest"] is None:
        raise HTTPException(status_code=422, detail="The latest report must contain a ChartData sheet.")
    prepared_latest = prepare_chart(charts["latest"])

    comparisons = []
    observed_contexts = [ReportChart({}, getattr(charts['latest'], 'context', None), getattr(charts['latest'], 'sections', ()))]
    for key, upload in uploads[1:]:
        previous_chart = await _load_chart_data(upload, key)
        try:
            if previous_chart is not None:
                observed_contexts.append(ReportChart({}, getattr(previous_chart, 'context', None), getattr(previous_chart, 'sections', ())))
            context, algorithm_known = _validate_comparison_context(observed_contexts)
        except ReportContextError as exc:
            raise HTTPException(status_code=422, detail=exc.detail()) from exc
        try:
            comparison = build_aligned_comparison(charts["latest"], previous_chart, prepared_latest=prepared_latest)
        except TypeError as exc:
            # Preserve compatibility with lightweight test doubles and older
            # integrations that still expose the two-argument helper.
            if "prepared_latest" not in str(exc):
                raise
            comparison = build_aligned_comparison(charts["latest"], previous_chart)
        if previous_chart is None:
            reason = f"The {key} report does not contain a ChartData sheet."
            comparison["reason"] = reason
            for metric in comparison.get("metrics", {}).values():
                metric["reason"] = reason
        comparison.update({"key": key, "previous_file": upload.filename or key,
                           'context': asdict(context), 'algorithm_status': 'verified' if algorithm_known else 'unknown'})
        comparisons.append(comparison)
        previous_chart = None

    return {
        "status": "ready",
        "latest_file": latest.filename or "latest",
        "comparisons": comparisons,
    }
