"""Explicitly run with venv Python from repository root; never overwrites a run.

Real localhost HTTP routes/server thread, not TestClient. RSS includes the client
and validation in the same process. Snapshot access supplies the exact immutable
raw oracle; this is not an independent historical detector implementation.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
import hashlib
import io
import itertools
import json
import math
import os
from pathlib import Path
import platform
import socket
import sys
import threading
import time
import traceback
import zipfile

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "backend"))

import httpx
import numpy as np
import openpyxl
from openpyxl import load_workbook
import pandas as pd
import uvicorn

from app.api.endpoints import analysis
from app.core.chainage_aggregation import aggregate_chainage, CHANNELS
from app.core.exporter import ExcelExporter
from app.core.tov1050_report_export import REPORT_NUMERIC_FORMAT
from app.main import app

HERE = Path(__file__).parent
LOG = None
STOP = threading.Event()
RSS = []
LATENCIES = []


def emit(event, **values):
    record = {"event": event, "wall_time": time.time(), **values}
    line = json.dumps(record, ensure_ascii=False, default=str)
    print(line, flush=True)
    if LOG:
        LOG.write(line + "\n")
        LOG.flush()


def digest(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def raw_digest(frame):
    hashed = pd.util.hash_pandas_object(frame, index=True).to_numpy()
    return hashlib.sha256(hashed.tobytes()).hexdigest()


def rss_sample():
    while not STOP.wait(0.25):
        RSS.append((time.perf_counter(), rss_bytes()))


def rss_bytes():
    if os.name == "nt":
        import ctypes
        from ctypes import wintypes
        class MemoryCounters(ctypes.Structure):
            _fields_ = [("cb", wintypes.DWORD), ("PageFaultCount", wintypes.DWORD),
                        *[(name, ctypes.c_size_t) for name in (
                            "PeakWorkingSetSize", "WorkingSetSize", "QuotaPeakPagedPoolUsage", "QuotaPagedPoolUsage",
                            "QuotaPeakNonPagedPoolUsage", "QuotaNonPagedPoolUsage", "PagefileUsage", "PeakPagefileUsage", "PrivateUsage")]]
        counters = MemoryCounters()
        counters.cb = ctypes.sizeof(counters)
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.GetCurrentProcess.restype = wintypes.HANDLE
        psapi = ctypes.WinDLL("psapi", use_last_error=True)
        psapi.GetProcessMemoryInfo.argtypes = (wintypes.HANDLE, ctypes.c_void_p, wintypes.DWORD)
        psapi.GetProcessMemoryInfo.restype = wintypes.BOOL
        if not psapi.GetProcessMemoryInfo(kernel.GetCurrentProcess(), ctypes.byref(counters), counters.cb):
            raise ctypes.WinError(ctypes.get_last_error())
        return counters.WorkingSetSize
    if Path("/proc/self/statm").exists():
        return int(Path("/proc/self/statm").read_text().split()[1]) * os.sysconf("SC_PAGE_SIZE")
    import resource
    return resource.getrusage(resource.RUSAGE_SELF).ru_maxrss * (1 if sys.platform == "darwin" else 1024)


def request(client, method, path, **kwargs):
    start = time.perf_counter()
    response = client.request(method, path, **kwargs)
    latency = time.perf_counter() - start
    if path.startswith("/api/jobs/") and not path.endswith(("/result", "/download")):
        LATENCIES.append(latency)
    if response.status_code >= 400:
        raise RuntimeError(f"{method} {path}: {response.status_code} {response.text[:2000]}")
    return response


def wait_job(client, created, session, limit):
    started = time.perf_counter()
    job_id = created["job_id"]
    transitions = []
    previous = None
    stage_started = started
    while True:
        state = request(client, "GET", f"/api/jobs/{job_id}", params={"client_session_id": session}).json()
        stage = state["stage"]
        if stage != previous:
            if previous is not None:
                transitions.append({"stage": previous, "observed_seconds": time.perf_counter() - stage_started})
            previous, stage_started = stage, time.perf_counter()
        emit("poll", **state)
        if state["status"] == "ready":
            transitions.append({"stage": previous, "observed_seconds": time.perf_counter() - stage_started})
            return job_id, time.perf_counter() - started, transitions
        if state["status"] in {"failed", "cancelled", "timed_out"}:
            raise RuntimeError(f"Job failed: {state}")
        if time.perf_counter() - started > limit:
            request(client, "POST", f"/api/jobs/{job_id}/cancel", params={"client_session_id": session})
            raise TimeoutError(f"Job exceeded benchmark limit {limit}s; cooperative cancel requested")
        time.sleep(1)


def equal(actual, expected, location):
    if expected is None or expected is pd.NA or expected is pd.NaT:
        assert actual is None, (location, actual, expected)
    elif isinstance(expected, (float, np.floating)):
        if math.isnan(expected):
            assert actual is None, (location, actual, expected)
        elif math.isinf(expected):
            assert actual == str(expected), (location, actual, expected)
        else:
            # XLSX numeric serialization has about 15 significant digits.
            assert isinstance(actual, (int, float)) and math.isclose(actual, expected, rel_tol=1e-14, abs_tol=1e-9), (location, actual, expected)
    else:
        assert actual == expected, (location, actual, expected)


def verify_frame(sheet, expected):
    rows = sheet.iter_rows(values_only=True, max_col=len(expected.columns))
    columns = next(rows)
    assert list(columns) == list(expected.columns), (sheet.title, columns, list(expected.columns))
    missing = object()
    count = 0
    for count, (actual, wanted) in enumerate(itertools.zip_longest(rows, expected.itertuples(index=False, name=None), fillvalue=missing), 1):
        assert actual is not missing and wanted is not missing, (sheet.title, count, "row count mismatch")
        for column, value, expected_value in zip(columns, actual, wanted):
            equal(value, expected_value, (sheet.title, count, column))
        if count % 100000 == 0:
            emit("verified_rows", sheet=sheet.title, rows=count)
    assert count == len(expected)
    return count


def verify_mean_oracle(raw, mean):
    """Check every mean/min/max/count against independently grouped raw arrays."""
    context = [name for name in raw if name.lower() in {
        "line", "track", "section", "session", "session_id", "task_no", "task_run_date",
        "station_start", "station_end", "class", "track type", "overlap", "tension length",
        "landmark", "source_file", "source_filename", "metadata_segment_id", "segment_id"}]
    keys = raw[["Chainage", *context]].copy().reset_index(drop=True)
    counter = pd.to_numeric(raw.get("SampleCount", pd.Series(np.nan, index=raw.index)), errors="coerce")
    keys["_pass"] = counter.ffill().diff().lt(0).cumsum().to_numpy()
    keys["_invalid"] = np.where(np.isfinite(pd.to_numeric(keys.Chainage, errors="coerce")), -1, np.arange(len(raw)))
    groups = keys.groupby(list(keys), sort=False, dropna=False).ngroup()
    assert len(mean) == groups.nunique()
    for column in CHANNELS:
        if column not in raw:
            continue
        values = pd.to_numeric(raw[column], errors="coerce").to_numpy(dtype=float, na_value=np.nan)
        finite = np.isfinite(values)
        group_ids = groups.to_numpy()
        counts = np.bincount(group_ids, weights=finite.astype(int), minlength=len(mean))
        sums = np.bincount(group_ids, weights=np.where(finite, values, 0), minlength=len(mean))
        wanted_mean = np.divide(sums, counts, out=np.full(len(mean), np.nan), where=counts != 0)
        minimum, maximum = np.full(len(mean), np.inf), np.full(len(mean), -np.inf)
        np.minimum.at(minimum, group_ids[finite], values[finite])
        np.maximum.at(maximum, group_ids[finite], values[finite])
        minimum[counts == 0] = np.nan
        maximum[counts == 0] = np.nan
        for name, wanted in ((column, wanted_mean), (column + "_min", minimum), (column + "_max", maximum), (column + "_count", counts)):
            np.testing.assert_allclose(mean[name].to_numpy(), wanted, rtol=1e-13, atol=1e-9, equal_nan=True, err_msg=name)
    emit("mean_oracle_verified", raw_rows=len(raw), grouped_rows=len(mean))


def validate_workbook(path, raw, mean, mode, exceptions, revision):
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
    with_copy = {name: frame.copy(deep=True) for name, frame in exceptions.items()}
    legacy = load_workbook(ExcelExporter.export_report(with_copy, numeric_format=REPORT_NUMERIC_FORMAT), read_only=True)
    book = load_workbook(path, read_only=True)
    try:
        for name in legacy.sheetnames:
            actual = book[name].iter_rows(values_only=True, max_col=legacy[name].max_column)
            wanted = legacy[name].iter_rows(values_only=True)
            sentinel = object()
            for row, (a, b) in enumerate(itertools.zip_longest(actual, wanted, fillvalue=sentinel), 1):
                assert a is not sentinel and b is not sentinel
                for column, value, expected in zip(range(len(b)), a, b):
                    equal(value, expected, (name, row, column))
        frame = raw if mode == "raw" else mean
        chart_columns = [name for name in ExcelExporter.CATENARY_COLUMNS if name in frame]
        assert len(chart_columns) == 31, chart_columns
        verify_frame(book["ChartData"], frame[chart_columns])
        if mode == "chainage_mean":
            verify_frame(book["ChainageStatistics"], mean)
        else:
            assert "ChainageStatistics" not in book.sheetnames
        assert book["_Metadata"].sheet_state == "hidden"
        metadata = dict(itertools.islice(book["_Metadata"].values, 1, None))
        assert metadata["analysis_revision"] == revision
        assert metadata["chart_data_mode"] == mode
        assert metadata["raw_row_count"] == str(len(raw))
        if mode == "chainage_mean":
            assert metadata["grouped_row_count"] == str(len(mean))
            assert metadata["aggregation"] == "chainage_mean"
            assert metadata["aggregation_version"] == "exact_chainage_channel_mean_v1"
        return {"bytes": path.stat().st_size, "sha256": digest(path), "chart_rows": len(frame), "metadata": metadata}
    finally:
        book.close()
        legacy.close()


def run_case(client, label, size, limit):
    source = ROOT / f"00 Reference Document/Raw data/TWL/test file {size}.csv"
    session = f"http-benchmark-{label}-{time.time_ns()}"
    payload = dict(file_path=str(source), line="TWL", section="Mainline", session="Mainline", track="UT",
                   date_str="20260907", station_start="CEN", station_end="TSW", task_no="U1",
                   client_session_id=session, chart_data_mode="chainage_mean")
    rss_start = len(RSS)
    create_started = time.perf_counter()
    created = request(client, "POST", "/api/analysis/jobs", json=payload).json()
    create_seconds = time.perf_counter() - create_started
    job_id, seconds, stages = wait_job(client, created, session, limit)
    result = request(client, "GET", f"/api/jobs/{job_id}/result", params={"client_session_id": session}).json()
    with analysis.STATE_LOCK:
        snapshot = analysis.ANALYSIS_SESSIONS[session]
    raw = snapshot["raw_df"]
    before = raw_digest(raw)
    mean_started = time.perf_counter()
    mean, aggregation = aggregate_chainage(raw)
    verify_mean_oracle(raw, mean)
    emit("analysis_complete", label=label, seconds=seconds, create_seconds=create_seconds,
         stages=stages, raw_rows=len(raw), aggregation=aggregation,
         validation_aggregation_seconds=time.perf_counter() - mean_started,
         exceptions={name: len(frame) for name, frame in snapshot["results"].items()})
    reports = {}
    for mode in ("raw", "chainage_mean"):
        created = request(client, "POST", "/api/export/report/jobs", json={
            "client_session_id": session, "analysis_revision": result["analysis_revision"], "chart_data_mode": mode}).json()
        export_id, export_seconds, export_stages = wait_job(client, created, session, limit)
        destination = HERE / f"{label}-{mode}.xlsx"
        download_started = time.perf_counter()
        with client.stream("GET", f"/api/jobs/{export_id}/download", params={"client_session_id": session}) as response:
            response.raise_for_status()
            assert ("_ChainageMean" in response.headers["content-disposition"]) == (mode == "chainage_mean")
            with destination.open("xb") as output:
                for chunk in response.iter_bytes(65536):
                    output.write(chunk)
        download_seconds = time.perf_counter() - download_started
        validation_started = time.perf_counter()
        reports[mode] = validate_workbook(destination, raw, mean, mode, snapshot["results"], result["analysis_revision"])
        reports[mode].update(export_seconds=export_seconds, download_seconds=download_seconds,
                             validation_seconds=time.perf_counter() - validation_started, observed_stages=export_stages)
        emit("report_verified", label=label, mode=mode, **reports[mode])
    assert raw_digest(raw) == before, "Raw frame changed during export"
    case = dict(label=label, source=str(source), input_bytes=source.stat().st_size,
                source_sha256=digest(source), analysis_seconds=seconds, create_seconds=create_seconds,
                raw_rows=len(raw), grouped_rows=len(mean), reduction_ratio=len(raw) / max(1, len(mean)),
                reports=reports, sampled_peak_rss_bytes=max(value for _, value in RSS[rss_start:]))
    emit("case_passed", **case)
    with analysis.STATE_LOCK:
        analysis.ANALYSIS_SESSIONS.pop(session, None)
    return case


def main():
    global LOG
    parser = argparse.ArgumentParser()
    parser.add_argument("--partial-only", action="store_true")
    parser.add_argument("--skip-warmup", action="store_true")
    parser.add_argument("--job-limit-seconds", type=float, default=3600)
    args = parser.parse_args()
    LOG = (HERE / "events.jsonl").open("x", encoding="utf-8")
    server = None
    thread = None
    try:
        import openpyxl.xml
        emit("environment", python=platform.python_version(), pandas=pd.__version__, openpyxl=openpyxl.__version__,
             lxml=openpyxl.xml.LXML, cpu_count=os.cpu_count(), platform=platform.platform(),
             metadata_sha256=digest(ROOT / "config/TWL metadata.xlsx"),
             rss_scope="same process: localhost server, HTTP client and verification; sampled each 250ms",
             latency_scope="HTTP status request duration, includes local scheduling/GIL",
             float_tolerance="xlsx rel=1e-14 abs=1e-9; independent mean oracle rel=1e-13 abs=1e-9",
             limitations="No independent historical full detector baseline; snapshot is complete raw correctness oracle. Background load uncontrolled.")
        sock = socket.socket()
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
        sock.close()
        server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning"))
        thread = threading.Thread(target=server.run, daemon=True)
        thread.start()
        sampler = threading.Thread(target=rss_sample, daemon=True)
        sampler.start()
        deadline = time.perf_counter() + 30
        while not server.started:
            if not thread.is_alive() or time.perf_counter() > deadline:
                raise RuntimeError("HTTP server did not start")
            time.sleep(0.1)
        cases = []
        with httpx.Client(base_url=f"http://127.0.0.1:{port}", timeout=30, trust_env=False) as client:
            if not args.skip_warmup:
                cases.append(run_case(client, "partial-warmup", "partial", args.job_limit_seconds))
            cases.append(run_case(client, "partial-normal", "partial", args.job_limit_seconds))
            if not args.partial_only:
                cases.append(run_case(client, "full-normal", "full", args.job_limit_seconds))
        stats = {"count": len(LATENCIES), "median_seconds": float(np.median(LATENCIES)),
                 "p95_seconds": float(np.percentile(LATENCIES, 95)), "max_seconds": max(LATENCIES)}
        summary = {"status": "passed", "cases": cases, "poll_latency": stats,
                   "sampled_peak_rss_bytes": max(value for _, value in RSS)}
        (HERE / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        emit("all_passed", poll_latency=stats)
    except BaseException:
        emit("failed", traceback=traceback.format_exc())
        raise
    finally:
        STOP.set()
        if RSS:
            with (HERE / "rss-samples.csv").open("x", encoding="utf-8") as output:
                output.write("perf_counter,rss_bytes\n")
                output.writelines(f"{moment},{value}\n" for moment, value in RSS)
        if server:
            server.should_exit = True
        if thread:
            thread.join(10)
        LOG.close()


if __name__ == "__main__":
    main()
