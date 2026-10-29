"""Isolated XML-engine A/B; production writer and active jobs remain untouched."""
from __future__ import annotations

import argparse
from datetime import date, datetime, timedelta
import hashlib
import json
import math
import os
from pathlib import Path
import platform
import statistics
import subprocess
import sys
import time
from zipfile import ZipFile

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
sys.path.insert(0, str(ROOT / "backend"))


def fixture():
    import numpy as np
    import pandas as pd
    from app.core.exporter import ExcelExporter

    n = 10000
    i = np.arange(n)
    frame = pd.DataFrame(index=i)
    for column in ExcelExporter.CATENARY_COLUMNS:
        if column in ("Chainage", "height_min", "height_max", "wear_min", "wear_max", "stg_max", "stg_min") or column.startswith(("height", "stagger", "wear")):
            frame[column] = 98150.1234567 + i * 0.0123456789
        else:
            frame[column] = ["TWL", "中文 <tag> & \"quoted\"", None, "", "=1+2", "00123", " line\nnew\t tab ", "emoji 🚆"][len(frame.columns) % 8]
    frame["task_run_date"] = [datetime(2026, 9, 7) + timedelta(seconds=int(v)) for v in i]
    # Exercise numeric/null conversion, both infinity signs, and precision.
    frame.loc[0:6, "height1"] = [None, float("nan"), float("inf"), float("-inf"), 0., -0.123456789123456, 3500.1234567]
    frame.loc[0:6, "Overlap"] = [None, "", "Y", "N", "<>&", "中文", "=SUM(1,2)"]
    frame.loc[7, "task_run_date"] = pd.NaT
    stats = pd.DataFrame({
        "group_id": range(20), "height1_min": [None, float("nan"), float("inf"), float("-inf"), 1.123456789] * 4,
        "height1_max": range(20), "valid_count": [0, 1, 2, 3, 4] * 4,
        "sample_count": [10] * 20, "source_row_first": range(20),
        "source_row_last": range(9, 29), "flag": [True, False] * 10,
        "text": ["中文", "<>&", "", None, "=1+2"] * 4,
        "day": [date(2026, 9, 7), None, pd.NaT, datetime(2026, 9, 7, 8, 9, 10), date(1900, 1, 1)] * 4,
    })
    results = {}
    for name in ("Low Height", "High Height", "Wire Wear", "Stagger Left", "Stagger Right"):
        results[name] = pd.DataFrame([{
            "id": "例外-" + name, "FromM": 98150.1234567, "ToM": 98151.7654321,
            "maxValue": -0.123456789123456, "maxLocation": 98150.3333333,
            "level": "L1", "Tension Length": "A, B <&>", "Landmark": None,
            "task_run_date": datetime(2026, 9, 7), "line": "TWL", "Overlap": "Y",
            "check_date": "2026-09-07T08:09:10Z",
        }])
    metadata = {"aggregation": "chainage_mean", "raw_points": n, "grouped_points": 1000,
                "sample_counter_column": "SampleCount [n.]", "empty": None,
                "unicode": "中文<&>🚆", "timestamp": "2026-10-26T00:00:00Z"}
    return results, frame, stats, metadata


def atom(value):
    if isinstance(value, (datetime, date)):
        return [type(value).__name__, value.isoformat()]
    if isinstance(value, float):
        return ["float", value.hex()]
    return [type(value).__name__, value]


def fingerprint(path):
    from openpyxl import load_workbook
    with ZipFile(path) as archive:
        assert archive.testzip() is None
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    try:
        for sheet in workbook:
            digest = hashlib.sha256()
            style_cache = {}
            count = rows = width = 0
            for rows, row in enumerate(sheet.iter_rows(), 1):
                width = max(width, len(row))
                values = []
                for cell in row:
                    sid = getattr(cell, "_style_id", 0)
                    if sid not in style_cache:
                        style_cache[sid] = [cell.number_format, str(cell.font), str(cell.fill), str(cell.border),
                                            str(cell.alignment), str(cell.protection)]
                    values.append([atom(cell.value), cell.data_type, style_cache[sid]])
                    count += 1
                digest.update(json.dumps(values, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
                digest.update(b"\n")
            sheets.append({"name": sheet.title, "state": sheet.sheet_state, "rows": rows,
                           "columns": width, "cells": count, "sha256": digest.hexdigest()})
    finally:
        workbook.close()
    return sheets


def run_worker(engine):
    import openpyxl
    from app.core.tov1050_report_export import export_tov1050_report
    expected = engine == "lxml"
    assert openpyxl.LXML is expected, (engine, openpyxl.LXML)
    result, frame, stats, metadata = fixture()
    records = []
    for iteration in range(4):
        destination = HERE / f"{engine}-{iteration}.xlsx"
        assert not destination.exists(), destination
        events = []
        started = time.perf_counter()
        def progress(stage, processed=None, total=None, unit=None):
            if not events or events[-1][0] != stage:
                events.append([stage, time.perf_counter() - started])
        export_tov1050_report(result, frame, metadata, chart_statistics_df=stats,
                              output_path=destination, progress=progress)
        elapsed = time.perf_counter() - started
        record = {"iteration": iteration, "warmup": iteration == 0, "seconds": elapsed,
                  "bytes": destination.stat().st_size, "stages": events}
        records.append(record)
        print(json.dumps({"engine": engine, **record}), flush=True)
    # Validation is explicitly outside writer timing.
    fingerprints = [fingerprint(HERE / f"{engine}-{iteration}.xlsx") for iteration in range(4)]
    assert all(value == fingerprints[0] for value in fingerprints), "repeat workbook mismatch"
    details = {"engine": engine, "LXML": openpyxl.LXML, "openpyxl": openpyxl.__version__,
               "python": sys.version, "platform": platform.platform(), "records": records,
               "median_seconds": statistics.median(row["seconds"] for row in records if not row["warmup"]),
               "fingerprint": fingerprints[0]}
    if expected:
        from lxml import etree
        details["lxml"] = etree.LXML_VERSION
        details["libxml"] = etree.LIBXML_VERSION
        details["lxml_file"] = etree.__file__
    with (HERE / f"{engine}-result.json").open("x", encoding="utf-8") as output:
        json.dump(details, output, ensure_ascii=False, indent=2)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--worker", choices=("stdlib", "lxml"))
    args = parser.parse_args()
    if args.worker:
        run_worker(args.worker)
        return
    assert not (HERE / "summary.json").exists()
    for engine in ("stdlib", "lxml"):
        env = os.environ.copy()
        env["OPENPYXL_LXML"] = "False" if engine == "stdlib" else "True"
        # Both engines have the same target import path; only the documented
        # openpyxl selection flag differs. Two fresh processes, sequential load.
        env["PYTHONPATH"] = str(HERE / "lxml-target") + os.pathsep + str(ROOT / "backend")
        with (HERE / f"{engine}-process.log").open("x", encoding="utf-8") as log:
            process = subprocess.run([sys.executable, str(Path(__file__).resolve()), "--worker", engine],
                                     cwd=ROOT, env=env, stdout=log, stderr=subprocess.STDOUT)
        assert process.returncode == 0, f"{engine} failed: see process log"
    baseline = json.loads((HERE / "stdlib-result.json").read_text(encoding="utf-8"))
    candidate = json.loads((HERE / "lxml-result.json").read_text(encoding="utf-8"))
    assert baseline["fingerprint"] == candidate["fingerprint"], "cross-engine workbook mismatch"
    summary = {"equivalent_all_cells_and_styles": True,
               "fixture": {"chart_rows": 10000, "chart_columns": 31, "statistics_rows": 20,
                           "warmup_per_engine": 1, "measured_repeats_per_engine": 3},
               "stdlib_median_seconds": baseline["median_seconds"],
               "lxml_median_seconds": candidate["median_seconds"],
               "speedup": baseline["median_seconds"] / candidate["median_seconds"],
               "limitation": "Lightweight synthetic A/B while full benchmark runs; not a full-data or packaging validation.",
               "writer_sha256": hashlib.sha256((ROOT / "backend/app/core/tov1050_report_export.py").read_bytes()).hexdigest()}
    with (HERE / "summary.json").open("x", encoding="utf-8") as output:
        json.dump(summary, output, ensure_ascii=False, indent=2)
    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
