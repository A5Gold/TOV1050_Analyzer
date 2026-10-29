"""Final counter-alias/full-retained acceptance; independent batch, opt-in XLSX.

Run from repo root with venv Python. Does not modify batch-jobs-01 artifacts.
All output files use exclusive creation. Default verifies full retained data;
--write-compact also generates and reads every cell of the final compact report.
"""
from __future__ import annotations

import argparse
import hashlib
import itertools
import json
import math
from pathlib import Path
import platform
import socket
import sys
import threading
import time
import traceback
import zipfile

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "backend"))

import numpy as np
import openpyxl
from openpyxl import load_workbook
import pandas as pd
import httpx
import uvicorn

from app.api.endpoints import analysis
from app.core.chainage_aggregation import aggregate_chainage
from app.core.exporter import ExcelExporter
from app.core.tov1050_report_export import export_analysis_report, REPORT_NUMERIC_FORMAT
from app.main import app

HERE = Path(__file__).parent
LOG = None
IDENTITY_NAMES = {
    "line", "track", "section", "session", "session_id", "task_no", "task_run_date",
    "station_start", "station_end", "class", "track type", "overlap", "tension length",
    "landmark", "source_file", "source_filename", "metadata_segment_id", "segment_id",
}


def emit(event, **values):
    line = json.dumps({"event": event, "wall_time": time.time(), **values}, ensure_ascii=False, default=str)
    print(line, flush=True)
    LOG.write(line + "\n")
    LOG.flush()


def file_digest(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def frame_digest(frame):
    return hashlib.sha256(pd.util.hash_pandas_object(frame, index=True).to_numpy().tobytes()).hexdigest()


def verify_groups(raw, derived, metadata):
    """Independent finite-counter pass construction and NumPy group reductions."""
    counter_name = next(name for name in ("SampleCount", "SampleCount [n.]") if name in raw)
    assert metadata["sample_counter_column"] == counter_name
    counts = pd.to_numeric(raw[counter_name], errors="coerce").to_numpy(dtype=float, na_value=np.nan)
    finite_positions = np.flatnonzero(np.isfinite(counts))
    reset_positions = finite_positions[1:][np.diff(counts[finite_positions]) < 0]
    steps = np.zeros(len(raw), dtype=np.int64)
    steps[reset_positions] = 1
    passes = np.cumsum(steps)
    identities = [name for name in raw if name.lower() in IDENTITY_NAMES]
    keys = raw[["Chainage", *identities]].copy().reset_index(drop=True)
    keys["_pass"] = passes
    chainage = pd.to_numeric(keys.Chainage, errors="coerce").to_numpy(dtype=float, na_value=np.nan)
    keys["_invalid"] = np.where(np.isfinite(chainage), -1, np.arange(len(raw)))
    group_ids = keys.groupby(list(keys), sort=False, dropna=False).ngroup().to_numpy(dtype=np.int64)
    group_count = int(group_ids.max()) + 1 if len(group_ids) else 0
    assert len(derived) == group_count
    first = np.full(group_count, len(raw), dtype=np.int64)
    last = np.full(group_count, -1, dtype=np.int64)
    np.minimum.at(first, group_ids, np.arange(len(raw)))
    np.maximum.at(last, group_ids, np.arange(len(raw)))
    np.testing.assert_array_equal(derived.group_id, np.arange(group_count))
    np.testing.assert_array_equal(derived.pass_id, passes[first])
    np.testing.assert_array_equal(derived.sample_count, np.bincount(group_ids, minlength=group_count))
    np.testing.assert_array_equal(derived.SampleCount_first, raw[counter_name].to_numpy()[first])
    np.testing.assert_array_equal(derived.SampleCount_last, raw[counter_name].to_numpy()[last])
    np.testing.assert_array_equal(derived.source_row_first, raw.source_row_number.to_numpy()[first])
    np.testing.assert_array_equal(derived.source_row_last, raw.source_row_number.to_numpy()[last])
    for column in ["Chainage", *identities]:
        pd.testing.assert_series_equal(derived[column].reset_index(drop=True), raw[column].iloc[first].reset_index(drop=True), check_names=False)
    channels = [f"{kind}{index}" for kind in ("height", "stagger", "wear") for index in range(1, 5)]
    for column in channels:
        if column not in raw:
            continue
        values = pd.to_numeric(raw[column], errors="coerce").to_numpy(dtype=float, na_value=np.nan)
        finite = np.isfinite(values)
        valid_counts = np.bincount(group_ids, weights=finite.astype(int), minlength=group_count)
        sums = np.bincount(group_ids, weights=np.where(finite, values, 0), minlength=group_count)
        means = np.divide(sums, valid_counts, out=np.full(group_count, np.nan), where=valid_counts != 0)
        minimum, maximum = np.full(group_count, np.inf), np.full(group_count, -np.inf)
        np.minimum.at(minimum, group_ids[finite], values[finite])
        np.maximum.at(maximum, group_ids[finite], values[finite])
        minimum[valid_counts == 0] = np.nan
        maximum[valid_counts == 0] = np.nan
        for target, wanted in ((column, means), (column + "_min", minimum),
                               (column + "_max", maximum), (column + "_count", valid_counts)):
            np.testing.assert_allclose(derived[target], wanted, rtol=1e-13, atol=1e-9, equal_nan=True, err_msg=target)
        emit("channel_verified", channel=column, groups=group_count)
    for kind, prefix in (("height", "height"), ("stagger", "stg"), ("wear", "wear")):
        columns = [name for name in channels if name in derived and name.startswith(kind)]
        for statistic in ("min", "max"):
            expected = getattr(derived[columns], statistic)(axis=1)
            np.testing.assert_allclose(derived[f"{prefix}_{statistic}"], expected, rtol=1e-13, atol=1e-9, equal_nan=True)
    return {"raw_rows": len(raw), "grouped_rows": group_count, "counter_column": counter_name,
            "counter_decreases": len(reset_positions), "identity_columns": identities,
            "counter_first": raw[counter_name].iloc[0], "counter_last": raw[counter_name].iloc[-1]}


def equal(actual, expected, location):
    # XLSX inline empty strings are read by openpyxl as blank/None. This
    # conversion applies only to the exact empty string, never whitespace/0.
    if isinstance(expected, str) and expected == "":
        assert actual is None or actual == "", (location, actual, expected)
    elif expected is None or expected is pd.NA or expected is pd.NaT:
        assert actual is None, (location, actual, expected)
    elif isinstance(expected, (float, np.floating)):
        if math.isnan(expected):
            assert actual is None, (location, actual, expected)
        elif math.isinf(expected):
            assert actual == str(expected), (location, actual, expected)
        else:
            assert isinstance(actual, (int, float)) and math.isclose(actual, expected, rel_tol=1e-14, abs_tol=1e-9), (location, actual, expected)
    else:
        assert actual == expected, (location, actual, expected)


def verify_sheet(sheet, expected):
    rows = sheet.iter_rows(values_only=True, max_col=len(expected.columns))
    assert list(next(rows)) == list(expected.columns)
    sentinel = object()
    count = 0
    for count, (actual, wanted) in enumerate(itertools.zip_longest(rows, expected.itertuples(index=False, name=None), fillvalue=sentinel), 1):
        assert actual is not sentinel and wanted is not sentinel, (sheet.title, count, "row count")
        for column, value, expected_value in zip(expected.columns, actual, wanted):
            equal(value, expected_value, (sheet.title, count, column))
        if count % 20000 == 0:
            emit("sheet_rows_verified", sheet=sheet.title, rows=count)
    assert count == len(expected)


def verify_workbook(path, derived, metadata, results, *, legacy_alias_schema=False):
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
    legacy_results = {name: frame.copy(deep=True) for name, frame in results.items()}
    legacy = load_workbook(ExcelExporter.export_report(legacy_results, numeric_format=REPORT_NUMERIC_FORMAT), read_only=True)
    book = load_workbook(path, read_only=True)
    try:
        assert book.sheetnames == [*legacy.sheetnames, "ChartData", "ChainageStatistics", "_Metadata"]
        chart_columns = [name for name in ExcelExporter.CATENARY_COLUMNS if name in derived]
        assert len(chart_columns) == 31
        verify_sheet(book["ChartData"], derived[chart_columns])
        missing_legacy_columns = ["SampleCount_first", "SampleCount_last"] if legacy_alias_schema else []
        expected_statistics = derived.drop(columns=missing_legacy_columns)
        verify_sheet(book["ChainageStatistics"], expected_statistics)
        assert book["_Metadata"].sheet_state == "hidden"
        exported_metadata = dict(itertools.islice(book["_Metadata"].values, 1, None))
        expected_metadata = dict(metadata)
        if legacy_alias_schema:
            # Only the known pre-native-counter-alias schema is accepted. Its
            # revision belongs to the preserved old analysis, not this rerun.
            assert "sample_counter_column" not in exported_metadata
            expected_metadata.pop("sample_counter_column")
            expected_metadata["pass_rule"] = "SampleCount decrease"
            assert isinstance(exported_metadata.get("analysis_revision"), str) and exported_metadata["analysis_revision"]
            expected_metadata["analysis_revision"] = exported_metadata["analysis_revision"]
        for key, value in expected_metadata.items():
            assert exported_metadata[key] == str(value), (key, exported_metadata[key], value)
        if not legacy_alias_schema:
            assert exported_metadata["sample_counter_column"] in ("SampleCount", "SampleCount [n.]")
        sentinel = object()
        for name in legacy.sheetnames:
            for row, (actual, wanted) in enumerate(itertools.zip_longest(
                book[name].iter_rows(values_only=True, max_col=legacy[name].max_column),
                legacy[name].iter_rows(values_only=True), fillvalue=sentinel), 1):
                assert actual is not sentinel and wanted is not sentinel, (name, row)
                for column, value, expected in zip(range(len(wanted)), actual, wanted):
                    equal(value, expected, (name, row, column))
        return {"bytes": path.stat().st_size, "sha256": file_digest(path), "sheets": book.sheetnames,
                "chart_columns": len(chart_columns), "statistics_columns": len(expected_statistics.columns),
                "statistics_schema": list(expected_statistics.columns),
                "legacy_alias_schema": legacy_alias_schema,
                "legacy_missing_columns": missing_legacy_columns,
                "legacy_metadata_difference": {"sample_counter_column": "absent", "pass_rule": "SampleCount decrease",
                                               "analysis_revision": "preserved old revision"} if legacy_alias_schema else None,
                "metadata": exported_metadata}
    finally:
        book.close()
        legacy.close()


def export_over_http(session, revision, destination):
    """Final manager/API integration; persist every HTTP latency as measured."""
    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    port = sock.getsockname()[1]
    sock.close()
    server = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning"))
    thread = threading.Thread(target=server.run, daemon=True)
    thread.start()
    latency = []
    try:
        deadline = time.perf_counter() + 30
        while not server.started:
            if not thread.is_alive() or time.perf_counter() > deadline:
                raise RuntimeError("Final export HTTP server did not start")
            time.sleep(0.1)
        with httpx.Client(base_url=f"http://127.0.0.1:{port}", timeout=30, trust_env=False) as client:
            def request(method, path, **kwargs):
                started = time.perf_counter()
                response = client.request(method, path, **kwargs)
                duration = time.perf_counter() - started
                if method == "GET" and path.startswith("/api/jobs/"):
                    latency.append(duration)
                emit("http_duration", method=method, path=path, seconds=duration, status_code=response.status_code)
                response.raise_for_status()
                return response
            started = time.perf_counter()
            response = request("POST", "/api/export/report/jobs", json={"client_session_id": session,
                               "analysis_revision": revision, "chart_data_mode": "chainage_mean"})
            create_seconds = time.perf_counter() - started
            assert response.status_code == 202
            identity = response.json()["job_id"]
            emit("final_http_job_created", job_id=identity, create_seconds=create_seconds, session=session, revision=revision)
            started = time.perf_counter()
            while True:
                state = request("GET", f"/api/jobs/{identity}", params={"client_session_id": session}).json()
                emit("final_http_poll", **state)
                if state["status"] == "ready":
                    break
                if state["status"] in ("failed", "cancelled", "timed_out"):
                    raise RuntimeError(f"Final export job failed: {state}")
                if time.perf_counter() - started > 3600:
                    request("POST", f"/api/jobs/{identity}/cancel", params={"client_session_id": session})
                    raise TimeoutError("Final benchmark 3600s guard exceeded; cooperative cancellation requested")
                time.sleep(1)
            export_seconds = time.perf_counter() - started
            download_started = time.perf_counter()
            with client.stream("GET", f"/api/jobs/{identity}/download", params={"client_session_id": session}) as download:
                download.raise_for_status()
                assert "_ChainageMean" in download.headers["content-disposition"]
                with destination.open("wb") as output:
                    for chunk in download.iter_bytes(65536):
                        output.write(chunk)
            details = {"job_id": identity, "create_seconds": create_seconds, "export_seconds": export_seconds,
                       "download_seconds": time.perf_counter() - download_started,
                       "poll_latency": {"count": len(latency), "median_seconds": float(np.median(latency)),
                                        "p95_seconds": float(np.percentile(latency, 95)), "max_seconds": max(latency)},
                       "latency_scope": "status client.request wall time includes localhost scheduling/GIL; each duration persisted in http_duration events"}
            emit("final_http_export_complete", **details)
            return details
    finally:
        server.should_exit = True
        thread.join(10)
        emit("final_http_server_stopped", stopped=not thread.is_alive())


def main():
    global LOG
    parser = argparse.ArgumentParser()
    parser.add_argument("--size", choices=("partial", "full"), default="full")
    parser.add_argument("--write-compact", action="store_true")
    parser.add_argument("--existing-compact", type=Path,
                        help="Also verify a preserved pre-counter-alias compact XLSX against the same loaded raw frame")
    args = parser.parse_args()
    LOG = (HERE / f"{args.size}-events.jsonl").open("x", encoding="utf-8")
    try:
        source = ROOT / f"00 Reference Document/Raw data/TWL/test file {args.size}.csv"
        emit("environment", python=platform.python_version(), pandas=pd.__version__, openpyxl=openpyxl.__version__,
             source=str(source), source_sha256=file_digest(source), metadata_sha256=file_digest(ROOT / "config/TWL metadata.xlsx"),
             source_files_sha256={name: file_digest(ROOT / name) for name in (
                 "backend/app/core/chainage_aggregation.py", "backend/app/core/tov1050_report_export.py",
                 "backend/app/core/tov1050_data_ingestion.py", "backend/app/api/endpoints/analysis.py")},
             validation="final revision complete retained raw, independent NumPy groups/statistics/counter ranges; optional final XLSX all cells",
             tolerance="oracle rtol=1e-13 atol=1e-9; XLSX rtol=1e-14 atol=1e-9")
        session = f"final-compact-alias-{args.size}-{time.time_ns()}"
        started = time.perf_counter()
        last_stage = [None]
        def progress(stage, processed=None, total=None, unit=None):
            if stage != last_stage[0] or processed is not None:
                emit("progress", stage=stage, processed=processed, total=total, unit=unit)
            last_stage[0] = stage
        response = analysis._analyze_sync(analysis.AnalyzeRequest(
            file_path=str(source), line="TWL", section="Mainline", session="Mainline", track="UT",
            date_str="20260907", station_start="CEN", station_end="TSW", task_no="U1",
            client_session_id=session, chart_data_mode="raw"), progress=progress)
        snapshot = analysis.ANALYSIS_SESSIONS[session]
        raw = snapshot["raw_df"]
        raw_before = frame_digest(raw)
        exceptions_before = {name: frame_digest(frame) for name, frame in snapshot["results"].items()}
        emit("analysis_complete", seconds=time.perf_counter() - started, rows=len(raw), cleaning=response["cleaning_summary"],
             exceptions={name: len(frame) for name, frame in snapshot["results"].items()})
        started = time.perf_counter()
        derived, aggregation = aggregate_chainage(raw, progress=progress)
        oracle = verify_groups(raw, derived, aggregation)
        assert frame_digest(raw) == raw_before
        report = {"status": "passed", "scope": args.size, "oracle": oracle, "aggregation": aggregation,
                  "aggregation_and_oracle_seconds": time.perf_counter() - started,
                  "cleaning_summary": response["cleaning_summary"], "raw_digest": raw_before,
                  "analysis_revision": snapshot["analysis_revision"]}
        emit("final_alias_groups_verified", **report)
        if args.write_compact or args.existing_compact:
            metadata = {"analysis_revision": snapshot["analysis_revision"], "chart_data_mode": "chainage_mean",
                        "raw_row_count": len(raw), "grouped_row_count": len(derived), **aggregation}
            capabilities = snapshot.get("metadata_capabilities")
            if capabilities:
                metadata.update(metadata_capabilities=json.dumps(capabilities, ensure_ascii=False),
                                enrichment_version="tov640-first-nonempty-v1")
        if args.existing_compact:
            assert oracle["counter_decreases"] == 0, "Old alias-blind grouping only equivalent when native counter never resets"
            existing = args.existing_compact.resolve(strict=True)
            started = time.perf_counter()
            report["existing_workbook"] = verify_workbook(existing, derived, metadata, snapshot["results"], legacy_alias_schema=True)
            report["existing_validation_seconds"] = time.perf_counter() - started
            emit("existing_compact_verified", path=str(existing), **report["existing_workbook"])
        if args.write_compact:
            filename = analysis.generate_filename(date_str="20260907", line="TWL", track="UT", section="Mainline",
                                                  file_type="report", task_no="U1", station_start="CEN", station_end="TSW", session="Mainline")
            destination = HERE / f"{args.size}_{Path(filename).stem}_ChainageMean.xlsx"
            # Reserve the filename exclusively, so reruns cannot replace artifacts.
            with destination.open("xb"):
                pass
            started = time.perf_counter()
            report["http_export"] = export_over_http(session, snapshot["analysis_revision"], destination)
            report["export_seconds"] = report["http_export"]["export_seconds"]
            emit("final_compact_written", path=str(destination), seconds=report["export_seconds"], bytes=destination.stat().st_size)
            started = time.perf_counter()
            report["workbook"] = verify_workbook(destination, derived, metadata, snapshot["results"])
            report["workbook_validation_seconds"] = time.perf_counter() - started
        assert frame_digest(raw) == raw_before
        for name, frame in snapshot["results"].items():
            assert frame_digest(frame) == exceptions_before[name], name
        with (HERE / f"{args.size}-summary.json").open("x", encoding="utf-8") as output:
            json.dump(report, output, ensure_ascii=False, indent=2, default=str)
        emit("all_passed", **report)
    except BaseException:
        emit("failed", traceback=traceback.format_exc())
        raise
    finally:
        LOG.close()


if __name__ == "__main__":
    main()
