"""Disposable constant-memory proof of concept; not a production writer."""
from __future__ import annotations
from copy import copy
from datetime import date, datetime
import importlib.util
import hashlib
import json
from pathlib import Path
import statistics
import sys
import time
from zipfile import ZipFile

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
sys.path.insert(0, str(HERE / "target"))
sys.path.insert(0, str(ROOT / "backend"))
import openpyxl
import pandas as pd
import xlsxwriter
from app.core.exporter import ExcelExporter
from app.core.tov1050_report_export import export_tov1050_report, _excel_value, REPORT_NUMERIC_COLUMNS

spec = importlib.util.spec_from_file_location("xml_fixture", HERE.parent / "xml-engine-01/benchmark_xml_engine.py")
fixture_module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fixture_module)


def prototype(results, frame, metadata, stats, output):
    # Keep legacy Summary/exception preparation identical. Only the final
    # row-by-row XML output implementation is swapped in this experiment.
    source = openpyxl.load_workbook(ExcelExporter.export_report(
        {key: value.copy(deep=True) for key, value in results.items()}, numeric_format="0.####"), read_only=True)
    target = xlsxwriter.Workbook(output, {"constant_memory": True, "strings_to_urls": False})
    formats = {}
    def fmt(number_format):
        if number_format not in formats:
            formats[number_format] = target.add_format({"num_format": number_format})
        return formats[number_format]
    header = target.add_format({"bold": True, "font_name": "Calibri", "font_size": 11,
                                "border": 1, "align": "center", "valign": "top"})
    def write(sheet, row, col, value, number_format="General", is_header=False):
        value = _excel_value(value)
        cell_format = header if is_header else fmt(number_format)
        if isinstance(value, datetime):
            sheet.write_datetime(row, col, value, fmt("yyyy-mm-dd h:mm:ss") if number_format == "General" else cell_format)
        elif isinstance(value, date):
            sheet.write_datetime(row, col, value, fmt("yyyy-mm-dd") if number_format == "General" else cell_format)
        elif value is None or value == "":
            # Both become blank cells on openpyxl read-back. XML blank type
            # differs between engines; this experiment records that limitation.
            sheet.write_blank(row, col, None, cell_format)
        else:
            sheet.write(row, col, value, cell_format)
    try:
        for old_sheet in source:
            sheet = target.add_worksheet(old_sheet.title)
            for row_index, row in enumerate(old_sheet):
                for column, cell in enumerate(row):
                    write(sheet, row_index, column, cell.value, cell.number_format, row_index == 0)
        for name, data, columns in (("ChartData", frame, ExcelExporter.CATENARY_COLUMNS),
                                    ("ChainageStatistics", stats, list(stats.columns))):
            sheet = target.add_worksheet(name)
            for column, label in enumerate(columns):
                write(sheet, 0, column, label, is_header=True)
            for row_index, row in enumerate(data[columns].itertuples(index=False, name=None), 1):
                for column, (label, value) in enumerate(zip(columns, row)):
                    number_format = "0.####" if label in REPORT_NUMERIC_COLUMNS or label.endswith(("_min", "_max")) else "General"
                    write(sheet, row_index, column, value, number_format)
        sheet = target.add_worksheet("_Metadata")
        sheet.hide()
        write(sheet, 0, 0, "key", is_header=True)
        write(sheet, 0, 1, "value", is_header=True)
        for row_index, (key, value) in enumerate(metadata.items(), 1):
            write(sheet, row_index, 0, key)
            write(sheet, row_index, 1, str(value))
    finally:
        source.close()
        target.close()


def digest_workbook(path):
    with ZipFile(path) as archive:
        assert archive.testzip() is None
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    result = []
    try:
        for sheet in book:
            rows = sheet.iter_rows()
            first = next(rows)
            width = len(first)
            digest = hashlib.sha256()
            count = 0
            def consume(row):
                nonlocal count
                row = list(row) + [None] * (width - len(row))
                values = []
                for cell in row:
                    value = cell.value if cell is not None else None
                    # openpyxl EmptyCell.number_format is None for omitted XML
                    # cells; Excel applies General just as for explicit blanks.
                    number_format = (cell.number_format or "General") if cell is not None else "General"
                    data_type = cell.data_type if cell is not None and value is not None else "blank"
                    if isinstance(value, (datetime, date)):
                        value = [type(value).__name__, value.isoformat()]
                    # JSON numeric equality, not byte-level XML identity.
                    values.append([value, data_type, number_format])
                digest.update(json.dumps(values, ensure_ascii=False, separators=(",", ":")).encode())
                digest.update(b"\n")
                count += 1
            consume(first)
            for row in rows:
                consume(row)
            result.append({"name": sheet.title, "state": sheet.sheet_state, "rows": count,
                           "columns": width, "sha256": digest.hexdigest()})
    finally:
        book.close()
    return result


def main():
    assert openpyxl.LXML is False, "baseline must use installed stdlib engine"
    results, frame, stats, metadata = fixture_module.fixture()
    assert frame.shape == (10000, 31)
    assert sum(col in REPORT_NUMERIC_COLUMNS for col in frame) == 19
    records = []
    # One warm-up each; measured order ABBAAB limits sequential-order bias.
    engines = ["openpyxl", "xlsxwriter", "openpyxl", "xlsxwriter", "xlsxwriter", "openpyxl", "openpyxl", "xlsxwriter"]
    seen = {"openpyxl": 0, "xlsxwriter": 0}
    for engine in engines:
        iteration = seen[engine]
        seen[engine] += 1
        path = HERE / f"{engine}-{iteration}.xlsx"
        assert not path.exists()
        start = time.perf_counter()
        if engine == "openpyxl":
            export_tov1050_report(results, frame, metadata, chart_statistics_df=stats, output_path=path)
        else:
            prototype(results, frame, metadata, stats, path)
        record = {"engine": engine, "iteration": iteration, "warmup": iteration == 0,
                  "seconds": time.perf_counter() - start, "bytes": path.stat().st_size}
        records.append(record)
        print(json.dumps(record), flush=True)
    fingerprints = {}
    for record in records:
        path = HERE / f"{record['engine']}-{record['iteration']}.xlsx"
        fingerprints[path.name] = digest_workbook(path)
    baseline = next(iter(fingerprints.values()))
    matches = {name: value == baseline for name, value in fingerprints.items()}
    summary = {"python": sys.version, "openpyxl": openpyxl.__version__, "LXML": openpyxl.LXML,
               "xlsxwriter": xlsxwriter.__version__, "xlsxwriter_file": xlsxwriter.__file__,
               "fixture": {"rows": 10000, "columns": 31, "numeric_columns": 19},
               "records": records, "fingerprints": fingerprints, "all_equal": all(matches.values()), "matches": matches,
               "limitations": ["Concurrent full benchmark creates background load; synthetic fixture only.",
                               "Compares values, nonblank data types, number formats, row/column order and hidden state; blank XML type is normalized.",
                               "Full header styles/layout contract, cancellation/temp cleanup, ZIP64/large-file, packaging remain unvalidated; prototype is disposable."]}
    for engine in seen:
        summary[engine + "_median_seconds"] = statistics.median(record["seconds"] for record in records if record["engine"] == engine and not record["warmup"])
    summary["speedup"] = summary["openpyxl_median_seconds"] / summary["xlsxwriter_median_seconds"]
    with (HERE / "summary.json").open("x", encoding="utf-8") as output:
        json.dump(summary, output, ensure_ascii=False, indent=2)
    print(json.dumps({key: value for key, value in summary.items() if key not in ("records", "fingerprints")}, ensure_ascii=False, indent=2), flush=True)
    assert summary["all_equal"], "read-back mismatch; inspect summary.json"


if __name__ == "__main__":
    main()
