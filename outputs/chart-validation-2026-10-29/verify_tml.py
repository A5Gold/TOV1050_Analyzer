"""Real DATAC analysis/export verification in an isolated output directory."""
from pathlib import Path
import gc
import json
import shutil
import sys
import time

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))
from app.api.endpoints import analysis as api

line = sys.argv[1] if len(sys.argv) > 1 else "TML"
source = ROOT / ("00 Reference Document/Raw data/TWL/20260907_TWL_UT.csv" if line == "TWL"
                 else "00 Reference Document/Raw data/TML/TML_D1_260410.datac")
metadata = ROOT / f"config/{line} metadata.xlsx"
date = "20260907" if line == "TWL" else "20260410"
track = "UT" if line == "TWL" else "DN"
task = "U1" if line == "TWL" else "D1"
session = f"real-{line.lower()}"
before = {str(path): api._file_digest(path) for path in (source, metadata)}
started = time.perf_counter()
output = api._analyze_sync(api.AnalyzeRequest(file_path=str(source), line=line, track=track, section="Mainline",
    date_str=date, task_no=task, client_session_id=session,
    station_start="CEN" if line == "TWL" else None, station_end="TSW" if line == "TWL" else None))
analysis_seconds = time.perf_counter() - started
state = api.ANALYSIS_SESSIONS[session]
print(line, "analysis", round(analysis_seconds, 3), "seconds", flush=True)
started = time.perf_counter()
job_info = api.JOBS.submit("export", session, lambda: api._export_job_work(state, "chainage_median"))
job = api.JOBS.jobs[job_info["job_id"]]
job.future.result(timeout=300)
assert job.status == "ready", job.error
destination = Path(__file__).parent / job.filename
shutil.copy2(job.path, destination)
export_seconds = time.perf_counter() - started
workbook = load_workbook(destination, read_only=True, data_only=True)
metadata_values = {row[0]: row[1] for row in workbook["_Metadata"].iter_rows(values_only=True) if row[0] != "key"}
chart_rows = workbook["ChartData"].iter_rows(values_only=True)
chart_header = next(chart_rows)
first_row = next(chart_rows)
chart_count = 1 + sum(1 for _ in chart_rows)
assert chart_count == len(state["analysis_df"])
assert metadata_values["algorithm"] == "same_chainage_median_sample_v1"
assert metadata_values["source_sha256"] == before[str(source)]
assert metadata_values["metadata_sha256"] == before[str(metadata)]
context = dict(zip(chart_header, first_row))
assert context["line"] == line and context["track"] == track
assert context["Section"] == "Mainline" and context["task_no"] == task
assert all(column in chart_header for column in ("Overlap", "Tension Length", "Landmark"))
if line == "TWL":
    assert context["station_start"] == "CEN" and context["station_end"] == "TSW"
assert "MedianSample" not in destination.name
summary = list(workbook["Summary"].iter_rows(min_row=1, max_row=3, values_only=True))
sheets = workbook.sheetnames
workbook.close()
after = {str(path): api._file_digest(path) for path in (source, metadata)}
assert before == after
result = {"source": str(source), "source_system": "TOV1050" if line == "TWL" else "TOV640", "selected_run_date": date,
          "source_date_note": "DATAC source Date=11.04.2026; selected form context remains authoritative" if line == "TML" else "form context authoritative",
          "analysis_seconds": analysis_seconds, "export_seconds": export_seconds,
          "exception_counts": {key: len(rows) for key, rows in output["exceptions"].items()},
          "aggregation": output["aggregation_metadata"], "chart_rows": chart_count,
          "chart_header": list(chart_header), "first_chart_context": {key: context.get(key) for key in
              ("line", "track", "Section", "task_no", "station_start", "station_end", "task_run_date", "Chainage")},
          "threshold_segment_count": len(output["threshold_segments"]), "output": str(destination),
          "output_bytes": destination.stat().st_size, "sheets": sheets, "summary_first_rows": summary,
          "metadata": metadata_values, "sha256_before": before, "sha256_after": after,
          "formal_database_written": False, "excel_native_visual_qa": "not performed"}
(Path(__file__).parent / f"{line.lower()}-verification.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(line, "export", round(export_seconds, 3), "seconds", chart_count, "ChartData rows", destination.name, flush=True)
api.close_chart_caches()
api.JOBS.close()
api.ANALYSIS_SESSIONS.clear()
gc.collect()
